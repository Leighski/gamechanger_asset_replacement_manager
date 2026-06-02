#!/usr/bin/env python3
"""
Inspired historical delivery audio audit.

Recursively scans Inspired delivery folders for *_COM.wav and *_CFX.wav,
validates against the official Inspired WAV specification, and writes
Inspired_Audio_Audit.xlsx (Summary / Failures / All Files tabs).

Default S3 prefix: s3://gcsuploads-mu/Inspired/DELIVERIES/

Requires: ffprobe on PATH, openpyxl, AWS CLI (``aws``) for S3 mode.
"""

from __future__ import annotations

import argparse
import subprocess
import sys
import tempfile
from collections import Counter
from pathlib import Path

_SCRIPT_DIR = Path(__file__).resolve().parent
if str(_SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(_SCRIPT_DIR))

from services.inspired_audio_validator import (  # noqa: E402
    INSPIRED_WAV_BITRATE_KBPS,
    INSPIRED_WAV_CHANNELS,
    INSPIRED_WAV_CODEC,
    INSPIRED_WAV_SAMPLE_RATE,
    audit_wav_pass_fail,
    validate_inspired_wav,
)

DEFAULT_S3_URI = "s3://gcsuploads-mu/Inspired/DELIVERIES/"


def _openpyxl_ok() -> bool:
    try:
        import openpyxl  # noqa: F401

        return True
    except ImportError:
        return False


def _is_target_wav(name: str) -> bool:
    lower = name.casefold()
    return lower.endswith("_com.wav") or lower.endswith("_cfx.wav")


def _list_s3_wavs(s3_uri: str) -> list[str]:
    if not s3_uri.endswith("/"):
        s3_uri += "/"
    proc = subprocess.run(
        ["aws", "s3", "ls", s3_uri, "--recursive"],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
    )
    if proc.returncode != 0:
        raise RuntimeError((proc.stderr or proc.stdout or "aws s3 ls failed").strip())
    out: list[str] = []
    for line in proc.stdout.splitlines():
        parts = line.split()
        if len(parts) < 4:
            continue
        key = parts[-1]
        if _is_target_wav(Path(key).name):
            bucket = s3_uri.replace("s3://", "").split("/", 1)[0]
            out.append(f"s3://{bucket}/{key}")
    return out


def _list_local_wavs(root: Path) -> list[Path]:
    return [p for p in root.rglob("*.wav") if _is_target_wav(p.name)]


def _download_s3(uri: str, dest: Path) -> None:
    dest.parent.mkdir(parents=True, exist_ok=True)
    proc = subprocess.run(
        ["aws", "s3", "cp", uri, str(dest)],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
    )
    if proc.returncode != 0:
        raise RuntimeError(f"aws s3 cp failed for {uri}: {(proc.stderr or proc.stdout).strip()}")


def _audit_one(path_label: str, local_path: Path) -> dict[str, str]:
    fields = validate_inspired_wav(local_path)
    pf = audit_wav_pass_fail(fields)
    ch = fields.get("channels", "")
    sr = fields.get("sample_rate", "")
    codec = fields.get("codec", "")
    return {
        "File Path": path_label,
        "Filename": Path(path_label).name,
        "Sample Rate": sr,
        "Channels": ch,
        "Codec": codec,
        "Bit Depth": fields.get("bit_depth", ""),
        "Bitrate": fields.get("bitrate", ""),
        "PASS / FAIL": pf,
        "Notes": fields.get("validation_notes", fields.get("failure_reason", "")),
    }


def _failure_category(row: dict[str, str]) -> str:
    if row.get("PASS / FAIL") == "PASS":
        return ""
    notes = (row.get("Notes") or "").casefold()
    ch = row.get("Channels", "")
    if ch == "2" or "stereo" in notes:
        return "stereo"
    if "sample_rate" in notes:
        return "sample_rate"
    if "codec" in notes:
        return "codec"
    if "channel" in notes:
        return "channels"
    return "other"


def write_audit_xlsx(path: Path, rows: list[dict[str, str]], *, deliveries_scanned: int) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws_all = wb.active
    ws_all.title = "All Files"
    headers = ["File Path", "Filename", "Sample Rate", "Channels", "Codec", "Bit Depth", "Bitrate", "PASS / FAIL", "Notes"]
    ws_all.append(headers)
    for cell in ws_all[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws_all.append([row.get(h, "") for h in headers])

    failures = [r for r in rows if r.get("PASS / FAIL") == "FAIL"]
    ws_fail = wb.create_sheet("Failures")
    ws_fail.append(headers)
    for cell in ws_fail[1]:
        cell.font = Font(bold=True)
    for row in failures:
        ws_fail.append([row.get(h, "") for h in headers])

    cats = Counter(_failure_category(r) for r in failures)
    ws_sum = wb.create_sheet("Summary")
    ws_sum.append(["Metric", "Value"])
    ws_sum["A1"].font = Font(bold=True)
    ws_sum["B1"].font = Font(bold=True)
    summary_rows = [
        ("Deliveries Scanned", deliveries_scanned),
        ("WAV Files Scanned", len(rows)),
        ("Stereo Files Found", sum(1 for r in failures if _failure_category(r) == "stereo")),
        ("Sample Rate Failures", cats.get("sample_rate", 0)),
        ("Codec Failures", cats.get("codec", 0)),
        ("Channel Failures", cats.get("channels", 0)),
        ("Total Failures", len(failures)),
        ("", ""),
        ("Expected Spec", ""),
        ("Codec", INSPIRED_WAV_CODEC),
        ("Sample Rate", str(INSPIRED_WAV_SAMPLE_RATE)),
        ("Channels", str(INSPIRED_WAV_CHANNELS)),
        ("Bit Depth", "16"),
        ("Bitrate (kbps)", str(INSPIRED_WAV_BITRATE_KBPS)),
    ]
    for a, b in summary_rows:
        ws_sum.append([a, b])

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))


def run_audit(*, s3_uri: str | None, local_root: Path | None, output: Path) -> int:
    if not _openpyxl_ok():
        print("openpyxl is required: pip install openpyxl", file=sys.stderr)
        return 2
    if not shutil_which("ffprobe"):
        print("ffprobe is required on PATH", file=sys.stderr)
        return 2

    rows: list[dict[str, str]] = []
    deliveries_scanned = 0

    with tempfile.TemporaryDirectory(prefix="inspired_audio_audit_") as tmp:
        tmp_path = Path(tmp)
        if local_root is not None:
            delivery_dirs = sorted({p.parent for p in _list_local_wavs(local_root) if "media" in p.parts})
            deliveries_scanned = len({p for p in local_root.iterdir() if p.is_dir()}) if local_root.is_dir() else len(delivery_dirs)
            for wav in _list_local_wavs(local_root):
                rows.append(_audit_one(str(wav), wav))
        else:
            uri = s3_uri or DEFAULT_S3_URI
            keys = _list_s3_wavs(uri)
            delivery_names = {Path(u.replace("s3://", "").split("/", 1)[-1]).parts[0] for u in keys if "/" in u.replace("s3://", "")}
            deliveries_scanned = len(delivery_names)
            for i, s3_path in enumerate(keys, start=1):
                local = tmp_path / f"{i:06d}_{Path(s3_path).name}"
                print(f"Downloading ({i}/{len(keys)}): {s3_path}")
                _download_s3(s3_path, local)
                rows.append(_audit_one(s3_path, local))

    write_audit_xlsx(output, rows, deliveries_scanned=deliveries_scanned)
    fails = sum(1 for r in rows if r.get("PASS / FAIL") == "FAIL")
    print(f"Wrote {output} — {len(rows)} WAV(s), {fails} failure(s), {deliveries_scanned} delivery folder(s).")
    return 0 if fails == 0 else 1


def shutil_which(cmd: str) -> bool:
    from shutil import which

    return bool(which(cmd))


def main() -> int:
    parser = argparse.ArgumentParser(description="Audit Inspired delivery COM/CFX WAV files.")
    parser.add_argument(
        "--s3-uri",
        default=DEFAULT_S3_URI,
        help=f"S3 deliveries prefix (default: {DEFAULT_S3_URI})",
    )
    parser.add_argument(
        "--local-root",
        type=Path,
        help="Audit local delivery tree instead of S3 (recursive *_COM.wav / *_CFX.wav)",
    )
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        default=Path("Inspired_Audio_Audit.xlsx"),
        help="Output workbook path",
    )
    args = parser.parse_args()
    if args.local_root is None and not shutil_which("aws"):
        print("AWS CLI (aws) is required for S3 audit mode", file=sys.stderr)
        return 2
    return run_audit(s3_uri=args.s3_uri, local_root=args.local_root, output=args.output)


if __name__ == "__main__":
    raise SystemExit(main())
