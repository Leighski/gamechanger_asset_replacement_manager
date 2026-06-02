#!/usr/bin/env python3
"""
Inspired historical delivery stereo → mono WAV remediation.

Scans S3 (or a local delivery tree) for *_COM.wav and *_CFX.wav, converts
eligible stereo files to mono (PCM_S16LE, 48000 Hz, 16-bit), and writes
Remediation_Log.xlsx.

Default S3 prefix: s3://gcsuploads-mu/Inspired/DELIVERIES/

Modes (exactly one required):
  --report-only   Probe and log stereo candidates; no conversion or upload.
  --dry-run       Download, convert, validate locally; no upload.
  --execute       Backup original on S3, upload mono replacement.

Requires: ffmpeg, ffprobe, openpyxl, AWS CLI (``aws``) for S3 mode.

Does not modify MP4 files, metadata, or WAV filenames.
"""

from __future__ import annotations

import argparse
import shutil
import subprocess
import sys
import tempfile
from collections import Counter
from dataclasses import dataclass
from enum import Enum
from pathlib import Path

_SCRIPT_DIR = Path(__file__).resolve().parent
if str(_SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(_SCRIPT_DIR))

from services.inspired_audio_validator import (  # noqa: E402
    INSPIRED_WAV_BIT_DEPTH,
    INSPIRED_WAV_CODEC,
    INSPIRED_WAV_SAMPLE_RATE,
    validate_inspired_wav,
)

DEFAULT_S3_URI = "s3://gcsuploads-mu/Inspired/DELIVERIES/"
STEREO_CHANNELS = 2
BACKUP_SUFFIX = ".stereo_backup.wav"
FFMPEG_TIMEOUT_SEC = 600

LOG_HEADERS = (
    "File",
    "Original Channels",
    "New Channels",
    "Original Bitrate",
    "New Bitrate",
    "Status",
    "Notes",
)

# All statuses written by _process_one():
#   SKIPPED_ALREADY_MONO      — passes Inspired mono spec; no action taken
#   SKIPPED_NOT_ELIGIBLE      — not 48 kHz PCM_S16LE 16-bit stereo
#   REPORT_STEREO_CANDIDATE   — stereo eligible; report-only mode
#   DRY_RUN_OK                — converted + validated locally; dry-run mode
#   REMEDIATED                — backed up + uploaded; execute mode
#   FAILED                    — conversion, validation, or upload error

CHECKPOINT_INTERVAL = 100


class RunMode(str, Enum):
    REPORT_ONLY = "report-only"
    DRY_RUN = "dry-run"
    EXECUTE = "execute"


@dataclass
class RemediationRow:
    file: str
    original_channels: str = ""
    new_channels: str = ""
    original_bitrate: str = ""
    new_bitrate: str = ""
    status: str = ""
    notes: str = ""

    def as_dict(self) -> dict[str, str]:
        return {
            "File": self.file,
            "Original Channels": self.original_channels,
            "New Channels": self.new_channels,
            "Original Bitrate": self.original_bitrate,
            "New Bitrate": self.new_bitrate,
            "Status": self.status,
            "Notes": self.notes,
        }


def _log(message: str) -> None:
    print(message, flush=True)


def _tool_ok(name: str) -> bool:
    return bool(shutil.which(name))


def _status_counts(rows: list[RemediationRow]) -> dict[str, int]:
    """Aggregate row statuses for workbook Summary and console output."""
    counts = Counter(row.status for row in rows)
    stereo_candidates = (
        counts.get("REPORT_STEREO_CANDIDATE", 0)
        + counts.get("DRY_RUN_OK", 0)
        + counts.get("REMEDIATED", 0)
        + counts.get("FAILED", 0)
    )
    return {
        "files_scanned": len(rows),
        "stereo_candidates": stereo_candidates,
        "already_mono": counts.get("SKIPPED_ALREADY_MONO", 0),
        "remediated": counts.get("REMEDIATED", 0),
        "dry_run_converted": counts.get("DRY_RUN_OK", 0),
        "skipped_not_eligible": counts.get("SKIPPED_NOT_ELIGIBLE", 0),
        "failed": counts.get("FAILED", 0),
        "report_stereo_candidate": counts.get("REPORT_STEREO_CANDIDATE", 0),
    }


def _print_run_summary(rows: list[RemediationRow], *, mode: RunMode, output: Path) -> None:
    stats = _status_counts(rows)
    _log(f"Wrote {output}")
    _log(f"FILES SCANNED: {stats['files_scanned']}")
    _log(f"STEREO CANDIDATES: {stats['stereo_candidates']}")
    _log(f"ALREADY MONO: {stats['already_mono']}")
    _log(f"REMEDIATED: {stats['remediated']}")
    if mode == RunMode.DRY_RUN or stats["dry_run_converted"]:
        _log(f"DRY RUN CONVERTED: {stats['dry_run_converted']}")
    if stats["skipped_not_eligible"]:
        _log(f"SKIPPED (NOT ELIGIBLE): {stats['skipped_not_eligible']}")
    _log(f"FAILED: {stats['failed']}")


def _openpyxl_ok() -> bool:
    try:
        import openpyxl  # noqa: F401

        return True
    except ImportError:
        return False


def _is_target_wav(name: str) -> bool:
    lower = name.casefold()
    return lower.endswith("_com.wav") or lower.endswith("_cfx.wav")


def _list_s3_wavs(s3_uri: str) -> list[str]:
    if not s3_uri.endswith("/"):
        s3_uri += "/"
    proc = subprocess.run(
        ["aws", "s3", "ls", s3_uri, "--recursive"],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
    )
    if proc.returncode != 0:
        raise RuntimeError((proc.stderr or proc.stdout or "aws s3 ls failed").strip())
    bucket = s3_uri.replace("s3://", "").split("/", 1)[0]
    out: list[str] = []
    for line in proc.stdout.splitlines():
        parts = line.split()
        if len(parts) < 4:
            continue
        key = parts[-1]
        if _is_target_wav(Path(key).name):
            out.append(f"s3://{bucket}/{key}")
    return out


def _list_local_wavs(root: Path) -> list[Path]:
    return sorted(p for p in root.rglob("*.wav") if _is_target_wav(p.name))


def _download_s3(uri: str, dest: Path) -> None:
    dest.parent.mkdir(parents=True, exist_ok=True)
    proc = subprocess.run(
        ["aws", "s3", "cp", uri, str(dest)],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
    )
    if proc.returncode != 0:
        raise RuntimeError(f"aws s3 cp download failed for {uri}: {(proc.stderr or proc.stdout).strip()}")


def _upload_s3(local_path: Path, uri: str) -> None:
    proc = subprocess.run(
        ["aws", "s3", "cp", str(local_path), uri],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
    )
    if proc.returncode != 0:
        raise RuntimeError(f"aws s3 cp upload failed for {uri}: {(proc.stderr or proc.stdout).strip()}")


def _copy_s3(src_uri: str, dest_uri: str) -> None:
    proc = subprocess.run(
        ["aws", "s3", "cp", src_uri, dest_uri],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
    )
    if proc.returncode != 0:
        raise RuntimeError(
            f"aws s3 cp backup failed {src_uri} → {dest_uri}: {(proc.stderr or proc.stdout).strip()}"
        )


def _backup_s3_uri(original_uri: str) -> str:
    if not original_uri.lower().endswith(".wav"):
        return f"{original_uri}{BACKUP_SUFFIX}"
    return original_uri[:-4] + BACKUP_SUFFIX


def _int_field(fields: dict[str, str], key: str) -> int | None:
    raw = str(fields.get(key, "")).strip()
    if not raw:
        return None
    try:
        return int(float(raw))
    except (TypeError, ValueError):
        return None


def _is_stereo_remediation_source(fields: dict[str, str]) -> tuple[bool, str]:
    """True when file is 48 kHz PCM_S16LE 16-bit stereo."""
    sr = _int_field(fields, "sample_rate")
    ch = _int_field(fields, "channels")
    bd = _int_field(fields, "bit_depth")
    codec = str(fields.get("codec", "")).lower()

    problems: list[str] = []
    if sr != INSPIRED_WAV_SAMPLE_RATE:
        problems.append(f"sample_rate={sr!r} (expected {INSPIRED_WAV_SAMPLE_RATE})")
    if codec != INSPIRED_WAV_CODEC:
        problems.append(f"codec={codec!r} (expected {INSPIRED_WAV_CODEC})")
    if bd is not None and bd != INSPIRED_WAV_BIT_DEPTH:
        problems.append(f"bit_depth={bd!r} (expected {INSPIRED_WAV_BIT_DEPTH})")
    if ch != STEREO_CHANNELS:
        problems.append(f"channels={ch!r} (expected {STEREO_CHANNELS} for remediation)")

    if problems:
        return False, "; ".join(problems)
    return True, ""


def _ffmpeg_stereo_to_mono(src: Path, dest: Path) -> tuple[bool, str]:
    if not _tool_ok("ffmpeg"):
        return False, "ffmpeg not on PATH"
    dest.parent.mkdir(parents=True, exist_ok=True)
    cmd = [
        "ffmpeg",
        "-hide_banner",
        "-nostats",
        "-y",
        "-i",
        str(src),
        "-ac",
        "1",
        "-ar",
        str(INSPIRED_WAV_SAMPLE_RATE),
        "-sample_fmt",
        "s16",
        "-c:a",
        "pcm_s16le",
        str(dest),
    ]
    try:
        proc = subprocess.run(
            cmd,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=FFMPEG_TIMEOUT_SEC,
            check=False,
        )
    except (OSError, subprocess.SubprocessError) as exc:
        return False, str(exc)
    if proc.returncode != 0:
        tail = (proc.stderr or proc.stdout or "").strip()
        if len(tail) > 2000:
            tail = tail[-2000:]
        return False, f"ffmpeg exit {proc.returncode}: {tail}"
    if not dest.is_file():
        return False, "ffmpeg completed but output file missing"
    return True, ""


def write_remediation_log(path: Path, rows: list[RemediationRow], *, mode: RunMode) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Remediation Log"
    ws.append(list(LOG_HEADERS))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([row.as_dict().get(h, "") for h in LOG_HEADERS])

    ws_sum = wb.create_sheet("Summary")
    ws_sum.append(["Metric", "Value"])
    ws_sum["A1"].font = Font(bold=True)
    ws_sum["B1"].font = Font(bold=True)
    stats = _status_counts(rows)
    summary = [
        ("Mode", mode.value),
        ("Files Scanned", stats["files_scanned"]),
        ("Stereo Candidates Found", stats["stereo_candidates"]),
        ("Already Mono", stats["already_mono"]),
        ("Remediated", stats["remediated"]),
        ("Dry Run Converted", stats["dry_run_converted"]),
        ("Skipped (Not Eligible)", stats["skipped_not_eligible"]),
        ("Failed", stats["failed"]),
    ]
    for a, b in summary:
        ws_sum.append([a, b])

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))


def _process_one(
    file_label: str,
    local_src: Path,
    *,
    mode: RunMode,
    s3_uri: str | None,
) -> RemediationRow:
    row = RemediationRow(file=file_label)
    before = validate_inspired_wav(local_src)
    row.original_channels = before.get("channels", "")
    row.original_bitrate = before.get("bitrate", "")

    if str(before.get("validation_ok", "")).lower() == "yes":
        row.status = "SKIPPED_ALREADY_MONO"
        row.new_channels = row.original_channels
        row.new_bitrate = row.original_bitrate
        row.notes = "Already meets Inspired mono WAV specification."
        return row

    eligible, reason = _is_stereo_remediation_source(before)
    if not eligible:
        row.status = "SKIPPED_NOT_ELIGIBLE"
        row.new_channels = row.original_channels
        row.new_bitrate = row.original_bitrate
        row.notes = reason or "Not an eligible stereo remediation source."
        return row

    if mode == RunMode.REPORT_ONLY:
        row.status = "REPORT_STEREO_CANDIDATE"
        row.notes = "Stereo PCM_S16LE 48 kHz 16-bit; remediation required."
        return row

    out_path = local_src.with_name(local_src.stem + "_mono.wav")
    ok, err = _ffmpeg_stereo_to_mono(local_src, out_path)
    if not ok:
        row.status = "FAILED"
        row.notes = f"Conversion failed: {err}"
        return row

    after = validate_inspired_wav(out_path)
    row.new_channels = after.get("channels", "")
    row.new_bitrate = after.get("bitrate", "")

    if str(after.get("validation_ok", "")).lower() != "yes":
        row.status = "FAILED"
        row.notes = after.get("failure_reason") or after.get("validation_notes") or "Output validation failed"
        return row

    if mode == RunMode.DRY_RUN:
        row.status = "DRY_RUN_OK"
        row.notes = "Converted and validated locally; upload skipped (dry-run)."
        return row

    if not s3_uri:
        row.status = "FAILED"
        row.notes = "Execute mode requires S3 target URI."
        return row

    backup_uri = _backup_s3_uri(s3_uri)
    try:
        _copy_s3(s3_uri, backup_uri)
        _upload_s3(out_path, s3_uri)
    except RuntimeError as exc:
        row.status = "FAILED"
        row.notes = str(exc)
        return row

    row.status = "REMEDIATED"
    row.notes = f"Backed up to {backup_uri}; uploaded mono replacement."
    return row


def _maybe_checkpoint(output: Path, rows: list[RemediationRow], *, mode: RunMode, index: int, total: int) -> None:
    if index % CHECKPOINT_INTERVAL == 0 or index == total:
        write_remediation_log(output, rows, mode=mode)
        _log(f"Checkpoint: wrote {output} ({index}/{total} files processed).")


def run_remediation(
    *,
    mode: RunMode,
    s3_uri: str | None,
    local_root: Path | None,
    output: Path,
    limit: int | None,
) -> int:
    if not _openpyxl_ok():
        print("openpyxl is required: pip install openpyxl", file=sys.stderr)
        return 2
    if not _tool_ok("ffprobe"):
        print("ffprobe is required on PATH", file=sys.stderr)
        return 2
    if mode != RunMode.REPORT_ONLY and not _tool_ok("ffmpeg"):
        print("ffmpeg is required on PATH for --dry-run and --execute", file=sys.stderr)
        return 2
    if local_root is None and not _tool_ok("aws"):
        print("AWS CLI (aws) is required for S3 mode", file=sys.stderr)
        return 2

    rows: list[RemediationRow] = []

    with tempfile.TemporaryDirectory(prefix="inspired_mono_remediation_") as tmp:
        tmp_path = Path(tmp)
        if local_root is not None:
            _log(f"Enumerating local WAV files under {local_root}…")
            targets = _list_local_wavs(local_root)
            if limit is not None:
                targets = targets[:limit]
            total = len(targets)
            _log(f"Found {total} WAV file(s).")
            _log(f"Beginning remediation (mode={mode.value})…")
            for i, wav in enumerate(targets, start=1):
                label = str(wav)
                _log(f"Processing ({i}/{total}): {label}")
                if mode == RunMode.REPORT_ONLY:
                    rows.append(_process_one(label, wav, mode=mode, s3_uri=None))
                else:
                    work = tmp_path / f"{i:06d}_{wav.name}"
                    shutil.copy2(wav, work)
                    rows.append(_process_one(label, work, mode=mode, s3_uri=None))
                _maybe_checkpoint(output, rows, mode=mode, index=i, total=total)
        else:
            uri = s3_uri or DEFAULT_S3_URI
            _log("Enumerating S3 objects…")
            keys = _list_s3_wavs(uri)
            if limit is not None:
                keys = keys[:limit]
            total = len(keys)
            _log(f"Found {total} WAV file(s).")
            _log(f"Beginning remediation (mode={mode.value})…")
            for i, key in enumerate(keys, start=1):
                _log(f"Processing ({i}/{total}): {key}")
                local = tmp_path / f"{i:06d}_{Path(key).name}"
                _download_s3(key, local)
                rows.append(_process_one(key, local, mode=mode, s3_uri=key))
                _maybe_checkpoint(output, rows, mode=mode, index=i, total=total)

    write_remediation_log(output, rows, mode=mode)
    _print_run_summary(rows, mode=mode, output=output)
    stats = _status_counts(rows)
    return 1 if stats["failed"] else 0


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Remediate historical Inspired COM/CFX WAV files from stereo to mono.",
    )
    mode_group = parser.add_mutually_exclusive_group(required=True)
    mode_group.add_argument(
        "--report-only",
        action="store_true",
        help="Probe and log stereo candidates only (no conversion or upload).",
    )
    mode_group.add_argument(
        "--dry-run",
        action="store_true",
        help="Download, convert, and validate locally; do not upload.",
    )
    mode_group.add_argument(
        "--execute",
        action="store_true",
        help="Backup original on S3 and upload mono replacement.",
    )
    parser.add_argument(
        "--s3-uri",
        default=DEFAULT_S3_URI,
        help=f"S3 deliveries prefix (default: {DEFAULT_S3_URI})",
    )
    parser.add_argument(
        "--local-root",
        type=Path,
        help="Remediate local delivery tree instead of S3 (for testing).",
    )
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        default=Path("Remediation_Log.xlsx"),
        help="Remediation log workbook path (default: Remediation_Log.xlsx)",
    )
    parser.add_argument(
        "--limit",
        type=int,
        default=None,
        help="Process at most N WAV files (optional, for testing).",
    )
    args = parser.parse_args()

    if args.report_only:
        mode = RunMode.REPORT_ONLY
    elif args.dry_run:
        mode = RunMode.DRY_RUN
    else:
        mode = RunMode.EXECUTE

    if mode == RunMode.EXECUTE and args.local_root is not None:
        print(
            "--execute against --local-root is not supported; use S3 or --dry-run locally.",
            file=sys.stderr,
        )
        return 2

    return run_remediation(
        mode=mode,
        s3_uri=args.s3_uri,
        local_root=args.local_root,
        output=args.output,
        limit=args.limit,
    )


if __name__ == "__main__":
    raise SystemExit(main())
