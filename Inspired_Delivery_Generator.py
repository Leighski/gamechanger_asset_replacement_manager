#!/usr/bin/env python3
"""
Inspired Delivery Generator — standalone, macOS-friendly GUI.

UI: CustomTkinter (dark neon dashboard). Install: pip install customtkinter

Optional: pip install tkinterdnd2 for drag/drop.
Optional: pip install openpyxl for reports/*.xlsx and inspired_metadata.xlsx.
"""

from __future__ import annotations

import csv
import io
import json
import math
import os
import queue
import re
import shutil
import ssl
import stat
import subprocess
import sys
import threading
import time
import traceback
import urllib.error
import urllib.parse
import urllib.request
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Callable, Iterable, Sequence

from team_registry import (
    TeamRegistryService,
    normalize_team_value_to_abbreviation,
    validate_inspired_row_teams_vs_filename,
)
from team_registry.normalize import INSPIRED_TEAM_COLUMNS

import tkinter as tk
from tkinter import filedialog, messagebox, ttk

try:
    import customtkinter as ctk
except ImportError as _e:  # pragma: no cover
    raise SystemExit(
        "Inspired Delivery Generator requires CustomTkinter for the UI. "
        "Install with: pip install customtkinter\n"
        "See README.md for details."
    ) from _e

# ---------------------------------------------------------------------------
# Optional tkinterdnd2 (drag/drop)
# ---------------------------------------------------------------------------
try:
    from tkinterdnd2 import DND_FILES, TkinterDnD

    _TKDND_AVAILABLE = True
except ImportError:
    DND_FILES = None  # type: ignore[misc, assignment]
    TkinterDnD = None  # type: ignore[misc, assignment]
    _TKDND_AVAILABLE = False

# ---------------------------------------------------------------------------
# Optional openpyxl (reports/*.xlsx)
# ---------------------------------------------------------------------------
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    _OPENPYXL_AVAILABLE = True
except ImportError:
    openpyxl = None  # type: ignore[misc, assignment]
    Font = None  # type: ignore[misc, assignment]
    PatternFill = None  # type: ignore[misc, assignment]
    get_column_letter = None  # type: ignore[misc, assignment]
    _OPENPYXL_AVAILABLE = False


APP_NAME = "Inspired Delivery Generator"
SETTINGS_VERSION = 6

PER_PAGE = 100
MAX_SEARCH_PAGES = 40
HTTP_TIMEOUT_SEC = 120

# Iconik search/metadata structure dumps (stdout + delivery log). Off by default.
VERBOSE_METADATA_DIAGNOSTICS = False

_GOAL_TIME_SECONDS_SLUG_CF = "goal_time_seconds_engp"
# Inspired ENGP naming: ``…_G_01_ENGP.mp4`` = goal event; ``…_O_01_ENGP.mp4`` = non-goal (no goal time).
_ENGP_GOAL_EVENT_RE = re.compile(r"_G_\d", re.IGNORECASE)
_ENGP_NON_GOAL_EVENT_RE = re.compile(r"_O_\d", re.IGNORECASE)
_CFX_INVISIBLE_CHARS = ("\ufeff", "\u200b", "\u200c", "\u200d", "\u2060", "\u00a0")
_CFX_CLOUDMOUNTER_RESCAN_DELAY_SEC = 2.0


def clean_cfx_source_root_path(raw: str) -> str:
    """
    Normalize CFX root from GUI/settings: trim outer whitespace and line breaks,
    strip one or more layers of surrounding ASCII quotes.
    """
    s = str(raw or "")
    s = s.strip().strip("\n\r\t\v\f ")
    while len(s) >= 2 and s[0] == s[-1] and s[0] in ("'", '"'):
        s = s[1:-1].strip().strip("\n\r\t\v\f ")
    return s.strip()


def _default_settings_path() -> Path:
    mac_support = Path.home() / "Library" / "Application Support" / APP_NAME
    try:
        mac_support.mkdir(parents=True, exist_ok=True)
        return mac_support / "settings.json"
    except OSError:
        return Path.home() / ".inspired_delivery_generator_settings.json"


@dataclass
class AppSettings:
    output_folder: str = ""
    window_geometry: str = ""
    iconik_base_url: str = ""
    iconik_app_id: str = ""
    iconik_auth_token: str = ""
    iconik_collection_uuid: str = ""
    cfx_source_root_folder: str = ""
    metadata_csv_paths: list[str] = field(default_factory=list)
    lock_delivery_outputs: bool = False
    csv_only_delivery: bool = False
    version: int = SETTINGS_VERSION

    def to_dict(self) -> dict:
        return {
            "version": self.version,
            "output_folder": self.output_folder,
            "window_geometry": self.window_geometry,
            "iconik_base_url": self.iconik_base_url,
            "iconik_app_id": self.iconik_app_id,
            "iconik_auth_token": self.iconik_auth_token,
            "iconik_collection_uuid": self.iconik_collection_uuid,
            "cfx_source_root_folder": self.cfx_source_root_folder,
            "metadata_csv_paths": list(self.metadata_csv_paths),
            "lock_delivery_outputs": bool(self.lock_delivery_outputs),
            "csv_only_delivery": bool(self.csv_only_delivery),
        }

    @classmethod
    def from_dict(cls, data: dict) -> AppSettings:
        mcsv = data.get("metadata_csv_paths", [])
        if not isinstance(mcsv, list):
            mcsv = []
        return cls(
            version=int(data.get("version", SETTINGS_VERSION)),
            output_folder=str(data.get("output_folder", "")),
            window_geometry=str(data.get("window_geometry", "")),
            iconik_base_url=str(data.get("iconik_base_url", "")),
            iconik_app_id=str(data.get("iconik_app_id", "")),
            iconik_auth_token=str(data.get("iconik_auth_token", "")),
            iconik_collection_uuid=str(data.get("iconik_collection_uuid", "")),
            cfx_source_root_folder=clean_cfx_source_root_path(str(data.get("cfx_source_root_folder", ""))),
            metadata_csv_paths=[str(x) for x in mcsv if str(x).strip()],
            lock_delivery_outputs=bool(data.get("lock_delivery_outputs", False)),
            csv_only_delivery=bool(data.get("csv_only_delivery", False)),
        )


def load_settings(path: Path) -> AppSettings:
    if not path.is_file():
        return AppSettings()
    try:
        with path.open("r", encoding="utf-8") as f:
            raw = json.load(f)
        if not isinstance(raw, dict):
            return AppSettings()
        return AppSettings.from_dict(raw)
    except (OSError, json.JSONDecodeError):
        return AppSettings()


def save_settings(path: Path, settings: AppSettings) -> None:
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        with path.open("w", encoding="utf-8") as f:
            json.dump(settings.to_dict(), f, indent=2)
    except OSError:
        pass


def _is_mp4(path: Path) -> bool:
    return path.suffix.lower() == ".mp4"


def _join_api_base(base: str) -> str:
    return base.rstrip("/") + "/"


def _print_http_client_exception_diagnostic(exc: BaseException, *, note: str = "") -> None:
    """Diagnostics only: full exception context to stdout (not swallowed)."""
    print("\n========== HTTP client exception (diagnostic) ==========", flush=True)
    if note:
        print(f"NOTE:\n{note}", flush=True)
    print(f"EXCEPTION TYPE:\n{type(exc).__module__}.{type(exc).__qualname__}", flush=True)
    print(f"EXCEPTION MESSAGE:\n{exc!s}", flush=True)
    print("FULL TRACEBACK:", flush=True)
    traceback.print_exception(type(exc), exc, exc.__traceback__, chain=True)
    print("========== end HTTP client exception diagnostic ==========\n", flush=True)


def _http_request_json(
    method: str,
    url: str,
    *,
    headers: dict[str, str],
    body: bytes | None = None,
    timeout: float = HTTP_TIMEOUT_SEC,
) -> tuple[int, dict | list | None, str]:
    """Return (status_code, parsed_json_or_none, raw_text_snippet)."""
    ctx = ssl.create_default_context()
    req = urllib.request.Request(url, data=body, headers=headers, method=method)
    try:
        with urllib.request.urlopen(req, timeout=timeout, context=ctx) as resp:
            raw = resp.read().decode("utf-8", errors="replace")
            try:
                return resp.status, json.loads(raw), raw[:2000]
            except json.JSONDecodeError as e:
                _print_http_client_exception_diagnostic(
                    e,
                    note=(
                        "JSONDecodeError on HTTP response body (transport OK). "
                        "requests.JSONDecodeError is analogous when using the requests library."
                    ),
                )
                return resp.status, None, raw[:2000]
    except urllib.error.HTTPError as e:
        try:
            raw = e.read().decode("utf-8", errors="replace")
        except OSError:
            raw = ""
        try:
            parsed = json.loads(raw) if raw else None
        except json.JSONDecodeError as je:
            _print_http_client_exception_diagnostic(
                je,
                note="JSONDecodeError while parsing HTTP error response body.",
            )
            parsed = None
        return e.code, parsed, raw[:2000]
    except urllib.error.URLError as e:
        _print_http_client_exception_diagnostic(
            e,
            note=(
                "urllib.error.URLError: network / TLS / proxy / handshake failures often surface here. "
                "Analogous requests types: requests.RequestException, requests.ConnectionError, "
                "requests.SSLError, requests.ConnectTimeout, requests.ReadTimeout, requests.ProxyError "
                "(this app uses urllib, not requests)."
            ),
        )
        detail = f"{type(e).__qualname__}: {e}"
        if isinstance(e.reason, BaseException):
            detail += f" | reason={type(e.reason).__qualname__}: {e.reason}"
        return 0, None, detail
    except ssl.SSLError as e:
        _print_http_client_exception_diagnostic(
            e,
            note="ssl.SSLError: TLS / certificate failures (also often carried on URLError.reason).",
        )
        return 0, None, f"{type(e).__qualname__}: {e}"
    except TimeoutError as e:
        _print_http_client_exception_diagnostic(
            e,
            note="TimeoutError: deadline exceeded (requests: ConnectTimeout / ReadTimeout).",
        )
        return 0, None, f"{type(e).__qualname__}: {e}"
    except OSError as e:
        _print_http_client_exception_diagnostic(
            e,
            note="OSError: includes connection reset, DNS failures, etc., depending on platform.",
        )
        return 0, None, f"{type(e).__qualname__}: {e}"
    except Exception as e:
        _print_http_client_exception_diagnostic(
            e,
            note=(
                "Unexpected exception from urllib.request.urlopen. "
                "requests.RequestException is the broad analogue when using requests."
            ),
        )
        return 0, None, f"{type(e).__qualname__}: {e}"


def _diag_emit_line(msg: str, emit) -> None:
    print(msg, flush=True)
    if emit is not None:
        emit(msg)


def extract_iconik_version_id_from_search_object(hit: dict | None) -> str | None:
    """Best-effort version UUID from an Iconik search hit (deterministic precedence)."""
    if not hit or not isinstance(hit, dict):
        return None
    for key in ("version_id", "original_version_id", "latest_version_id", "current_version_id", "head_version_id"):
        v = hit.get(key)
        if isinstance(v, str) and v.strip():
            return v.strip()
        if isinstance(v, int) and not isinstance(v, bool):
            return str(v)
    for nest_key in ("version", "latest_version", "head_version"):
        nest = hit.get(nest_key)
        if isinstance(nest, dict):
            for vk in ("id", "version_id"):
                vv = nest.get(vk)
                if isinstance(vv, str) and vv.strip():
                    return vv.strip()
                if isinstance(vv, int) and not isinstance(vv, bool):
                    return str(vv)
    vers = hit.get("versions")
    if isinstance(vers, list) and vers:
        first = vers[0]
        if isinstance(first, dict):
            for vk in ("id", "version_id"):
                vv = first.get(vk)
                if isinstance(vv, str) and vv.strip():
                    return vv.strip()
                if isinstance(vv, int) and not isinstance(vv, bool):
                    return str(vv)
    return None


def _diag_json_snippet(val: object, *, max_len: int = 1600) -> str:
    try:
        if isinstance(val, (dict, list)):
            s = json.dumps(val, default=str, ensure_ascii=False)
        else:
            s = repr(val)
    except TypeError:
        s = repr(val)
    if len(s) > max_len:
        return s[:max_len] + "…"
    return s


_ID_LIKE_KEYS = frozenset(
    {
        "id",
        "asset_id",
        "parent_id",
        "collection_id",
        "file_id",
        "proxy_id",
        "version_id",
        "original_version_id",
        "latest_version_id",
        "current_version_id",
        "head_version_id",
        "storage_id",
        "job_id",
        "transcode_job_id",
        "original_file_id",
        "format_id",
    }
)


def _emit_search_hit_container_ids(hit: dict, name: str, emit) -> None:
    """Print id-like fields for list-of-dicts or a single dict container."""
    v = hit.get(name)
    if v is None:
        return
    if isinstance(v, dict):
        _diag_emit_line(f"  {name} (object):", emit)
        for ik in sorted(v.keys(), key=lambda x: str(x).casefold()):
            iv = v.get(ik)
            if ik in _ID_LIKE_KEYS or (isinstance(ik, str) and ik.endswith("_id") and not ik.startswith("_")):
                _diag_emit_line(f"    {ik} = {_diag_json_snippet(iv, max_len=800)}", emit)
        return
    if not isinstance(v, list):
        _diag_emit_line(f"  {name}: {_diag_json_snippet(v, max_len=1200)}", emit)
        return
    limit = 40
    for i, item in enumerate(v[:limit]):
        if not isinstance(item, dict):
            _diag_emit_line(f"  {name}[{i}] = {_diag_json_snippet(item, max_len=600)}", emit)
            continue
        id_parts = []
        for ik in sorted(item.keys(), key=lambda x: str(x).casefold()):
            if ik in _ID_LIKE_KEYS or (isinstance(ik, str) and ik.endswith("_id") and not ik.startswith("_")):
                id_parts.append(f"{ik}={item.get(ik)!r}")
        line = f"  {name}[{i}]: " + ("; ".join(id_parts) if id_parts else _diag_json_snippet(item, max_len=600))
        _diag_emit_line(line, emit)
    if len(v) > limit:
        _diag_emit_line(f"  … {name} truncated ({len(v)} items total)", emit)


def print_iconik_search_result_structure_diagnostics(hit: dict | None, emit) -> None:
    """Stdout + optional delivery_log: keys, version fields, related object ids (diagnostics only)."""
    if not VERBOSE_METADATA_DIAGNOSTICS:
        return
    if not hit or not isinstance(hit, dict):
        _diag_emit_line("SEARCH RESULT RAW KEYS: (absent)", emit)
        _diag_emit_line("SEARCH RESULT VERSION IDS:", emit)
        _diag_emit_line("extracted_primary_version_id: None", emit)
        _diag_emit_line("SEARCH RESULT RELATED OBJECT IDS:", emit)
        _diag_emit_line("  (absent)", emit)
        return

    _diag_emit_line("SEARCH RESULT RAW KEYS:", emit)
    for k in sorted(hit.keys(), key=lambda x: str(x).casefold()):
        _diag_emit_line(str(k), emit)

    _diag_emit_line("SEARCH RESULT VERSION IDS:", emit)
    vid = extract_iconik_version_id_from_search_object(hit)
    _diag_emit_line(f"extracted_primary_version_id: {vid!r}", emit)
    for cand in (
        "version_id",
        "original_version_id",
        "latest_version_id",
        "current_version_id",
        "head_version_id",
        "versions",
        "version",
        "latest_version",
        "head_version",
    ):
        if cand not in hit:
            continue
        _diag_emit_line(f"  {cand}: {_diag_json_snippet(hit[cand])}", emit)

    _diag_emit_line("SEARCH RESULT RELATED OBJECT IDS:", emit)
    for k in sorted(hit.keys(), key=lambda x: str(x).casefold()):
        v = hit.get(k)
        if k in _ID_LIKE_KEYS or (isinstance(k, str) and k.endswith("_id") and not k.startswith("_")):
            _diag_emit_line(f"  top-level.{k} = {_diag_json_snippet(v, max_len=1200)}", emit)
    for container in (
        "versions",
        "files",
        "proxies",
        "metadata_views",
        "keyframes",
        "segments",
        "storage_methods",
        "renditions",
        "formats",
        "transcode_jobs",
        "storage",
    ):
        _emit_search_hit_container_ids(hit, container, emit)

    md = hit.get("metadata")
    if md is not None:
        _diag_emit_line("  metadata (search hit):", emit)
        if isinstance(md, dict):
            keys = sorted(md.keys(), key=lambda x: str(x).casefold())
            _diag_emit_line(f"    key_count={len(keys)}", emit)
            cap = 60
            for mk in keys[:cap]:
                mv = md[mk]
                _diag_emit_line(f"    {mk}: {_diag_json_snippet(mv, max_len=1000)}", emit)
            if len(keys) > cap:
                _diag_emit_line(f"    … metadata keys truncated ({len(keys)} total)", emit)
        else:
            _diag_emit_line(f"    (non-dict): {_diag_json_snippet(md)}", emit)


class IconikClient:
    """Minimal Iconik REST client: search (collection_id as query param only) + asset metadata GET."""

    def __init__(self, *, base_url: str, app_id: str, auth_token: str, collection_id: str) -> None:
        self._base_url_raw = base_url
        self._base = _join_api_base(base_url)
        self._collection_id = collection_id.strip()
        self._headers = {
            "App-ID": app_id.strip(),
            "Auth-Token": auth_token.strip(),
            "Content-Type": "application/json",
            "Accept": "application/json",
        }

    @staticmethod
    def _print_iconik_request_diagnostics(
        *,
        kind: str,
        base_url_raw: str,
        base_url_normalized: str,
        final_url: str,
        method: str,
        params: dict[str, str] | None,
        body_bytes: bytes | None,
    ) -> None:
        """Temporary live diagnostics only (stdout)."""
        if not VERBOSE_METADATA_DIAGNOSTICS:
            return
        body_str = ""
        if body_bytes is not None:
            try:
                body_str = body_bytes.decode("utf-8", errors="replace")
            except Exception:
                body_str = repr(body_bytes)
        params_repr = json.dumps(params, sort_keys=True) if params is not None else "(none)"
        print(f"\n========== Iconik request diagnostics [{kind}] ==========", flush=True)
        print(f"FINAL REQUEST URL:\n{final_url}", flush=True)
        print(f"FINAL REQUEST METHOD:\n{method}", flush=True)
        print(f"FINAL REQUEST PARAMS:\n{params_repr}", flush=True)
        print(f"FINAL REQUEST BODY:\n{body_str if body_str else '(empty)'}", flush=True)
        print(f"BASE URL RAW:\n{base_url_raw!r}", flush=True)
        print(f"BASE URL NORMALIZED:\n{base_url_normalized!r}", flush=True)
        print(f"========== end diagnostics [{kind}] ==========\n", flush=True)

    def search_assets(self, query: str, *, page: int) -> tuple[int, list[dict], str | None]:
        """
        POST search/v1/search/ with body {doc_types, query} only.
        Query params: page, per_page, collection_id=<uuid>
        """
        path = urllib.parse.urljoin(self._base, "search/v1/search/")
        params = {
            "page": str(page),
            "per_page": str(PER_PAGE),
            "collection_id": self._collection_id,
        }
        url = path + "?" + urllib.parse.urlencode(params)
        body = json.dumps({"doc_types": ["assets"], "query": query}).encode("utf-8")
        self._print_iconik_request_diagnostics(
            kind="search",
            base_url_raw=self._base_url_raw,
            base_url_normalized=self._base,
            final_url=url,
            method="POST",
            params=params,
            body_bytes=body,
        )
        code, data, _snippet = _http_request_json("POST", url, headers=dict(self._headers), body=body)
        if code != 200 or not isinstance(data, dict):
            return code, [], None
        objs = data.get("objects")
        if not isinstance(objs, list):
            objs = []
        return code, [o for o in objs if isinstance(o, dict)], None

    def search_all_pages(self, query: str, cancel_event: threading.Event) -> tuple[list[dict], str | None]:
        """Fetch all pages for a query until empty or cap. Dedupe by asset id."""
        seen: set[str] = set()
        out: list[dict] = []
        err: str | None = None
        for page in range(1, MAX_SEARCH_PAGES + 1):
            if cancel_event.is_set():
                return out, "cancelled"
            code, batch, _ = self.search_assets(query, page=page)
            if code != 200:
                err = f"search HTTP {code} for query {query!r} page {page}"
                break
            if not batch:
                break
            for o in batch:
                aid = str(o.get("id") or o.get("asset_id") or "")
                if not aid or aid in seen:
                    continue
                seen.add(aid)
                out.append(o)
            if len(batch) < PER_PAGE:
                break
        return out, err

    def get_asset_metadata(
        self,
        asset_id: str,
        cancel_event: threading.Event,
        *,
        version_id: str | None = None,
    ) -> tuple[dict | None, bool, str | None]:
        """
        GET asset metadata. When ``version_id`` is set (from search hit), uses the same endpoint
        shape as Inspired_Highlight_version.py: ``metadata/v1/assets/{id}/versions/{version_id}/``,
        then falls back to asset root if that request is not successful.
        """
        if cancel_event.is_set():
            return None, False, "cancelled"

        def do_get(rel_path: str, diag_kind: str) -> tuple[int, dict | list | None, str]:
            path = urllib.parse.urljoin(self._base, rel_path)
            self._print_iconik_request_diagnostics(
                kind=diag_kind,
                base_url_raw=self._base_url_raw,
                base_url_normalized=self._base,
                final_url=path,
                method="GET",
                params=None,
                body_bytes=None,
            )
            return _http_request_json("GET", path, headers=dict(self._headers), body=None)

        vid = (version_id or "").strip()
        if vid:
            code, data, snippet = do_get(
                f"metadata/v1/assets/{asset_id}/versions/{vid}/",
                "metadata_asset_version",
            )
            if code == 200 and isinstance(data, dict):
                return data, True, None
            code2, data2, snippet2 = do_get(
                f"metadata/v1/assets/{asset_id}/",
                "metadata_asset_root_fallback",
            )
            if code2 == 200 and isinstance(data2, dict):
                return data2, True, None
            return (
                data2 if isinstance(data2, dict) else (data if isinstance(data, dict) else None),
                False,
                snippet2 or snippet or f"HTTP {code2}",
            )

        code, data, snippet = do_get(f"metadata/v1/assets/{asset_id}/", "metadata_asset_root")
        if code == 200 and isinstance(data, dict):
            return data, True, None
        return (data if isinstance(data, dict) else None), False, snippet or f"HTTP {code}"


def _asset_title(hit: dict) -> str:
    t = hit.get("title")
    if isinstance(t, str):
        return t.strip()
    return ""


def _asset_id(hit: dict) -> str:
    return str(hit.get("id") or hit.get("asset_id") or "")


def _asset_date_sort_key(hit: dict) -> float:
    for key in ("date_created", "created_date", "system_date_created", "indexed_date"):
        v = hit.get(key)
        if isinstance(v, str) and v.strip():
            try:
                s = v.replace("Z", "+00:00")
                return datetime.fromisoformat(s).timestamp()
            except ValueError:
                continue
    return 0.0


def _tier_for_hit(filename: str, stem: str, title: str) -> int:
    """
    Iconik hit ↔ local ENGP filename tiers (strict; no substring / fuzzy tiers).

    1 = exact filename match, 2 = filename stem match, 3 = title match, 4 = case-insensitive filename.
    99 = no match under these rules (hit is ignored for picking).
    """
    if not title:
        return 99
    fn = filename
    fn_cf = fn.casefold()
    t_cf = title.casefold()
    if title == fn:
        return 1
    try:
        title_stem = Path(title).stem
    except ValueError:
        title_stem = title
    if title_stem == stem:
        return 2
    if title == stem:
        return 3
    if t_cf == fn_cf:
        return 4
    return 99


def pick_best_asset(hits: list[dict], filename: str, stem: str) -> tuple[dict | None, bool, int]:
    """
    Deterministic winner among Iconik search hits using strict filename/title tiers only.

    Returns (winner_or_none, ambiguous, winning_tier). When multiple hits share the best
    tier, returns (None, True, tier) — callers must not silently pick metadata.
    """
    if not hits:
        return None, False, 0
    rows: list[tuple[int, float, str, str, dict]] = []
    for h in hits:
        title = _asset_title(h)
        aid = _asset_id(h)
        if not aid:
            continue
        tier = _tier_for_hit(filename, stem, title)
        if tier > 4:
            continue
        ts = _asset_date_sort_key(h)
        rows.append((tier, ts, aid, title, h))
    if not rows:
        return None, False, 0
    rows.sort(key=lambda r: (r[0], -r[1], r[2]))
    winner_tier = rows[0][0]
    same_tier = [r for r in rows if r[0] == winner_tier]
    if len(same_tier) > 1:
        return None, True, winner_tier
    return same_tier[0][4], False, winner_tier


def partial_fallback_query(stem: str) -> str:
    if len(stem) > 2:
        return stem[:-1]
    if stem:
        return stem[:1]
    return ""


def write_ambiguous_xlsx(path: Path, rows: list[dict[str, str | int]]) -> bool:
    if not _OPENPYXL_AVAILABLE or openpyxl is None or Font is None or get_column_letter is None:
        return False
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ambiguous"
    headers = [
        "ambiguity_kind",
        "source_filename",
        "resolved_asset_id",
        "resolved_title",
        "match_tier",
        "candidate_count",
        "candidates_id_title",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append(
            [
                xlsx_sanitize_cell(r.get("ambiguity_kind", "")),
                xlsx_sanitize_cell(r.get("source_filename", "")),
                xlsx_sanitize_cell(r.get("resolved_asset_id", "")),
                xlsx_sanitize_cell(r.get("resolved_title", "")),
                xlsx_sanitize_cell(r.get("match_tier", "")),
                xlsx_sanitize_cell(r.get("candidate_count", "")),
                xlsx_sanitize_cell(r.get("candidates_id_title", "")),
            ]
        )
    if not rows:
        ws.append(
            [
                "—",
                "—",
                "—",
                "—",
                "—",
                "0",
                "(no ambiguous metadata matches)",
            ]
        )
    nrow = max(len(rows), 1) + 1
    ncol = len(headers)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ncol)}{nrow}"
    widths = (22, 36, 36, 36, 12, 14, 72)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    wb.save(str(path))
    return True


def flatten_metadata_preview(
    meta_doc: dict | None = None,
    *,
    flat: dict[str, object] | None = None,
    max_pairs: int = 12,
    max_len: int = 80,
) -> str:
    if flat is not None:
        slugs = tuple(sorted(flat.keys(), key=lambda s: (str(s).casefold(), str(s))))
        if not flat:
            return ""
    elif meta_doc:
        flat, slugs = build_flat_metadata_map(meta_doc)
        if not flat:
            return _shorten(json.dumps(meta_doc, default=str), max_len * 2)
    else:
        return ""
    parts: list[str] = []
    for i, slug in enumerate(slugs):
        if i >= max_pairs:
            parts.append("…")
            break
        parts.append(f"{slug}={_shorten(json.dumps(flat.get(slug), default=str), max_len)}")
    return " | ".join(parts)


def _shorten(s: str, n: int) -> str:
    s = s.replace("\n", " ")
    return s if len(s) <= n else s[: n - 1] + "…"


# ---------------------------------------------------------------------------
# Phase 3 — Inspired delivery in-memory dataframe + deterministic normalization
# ---------------------------------------------------------------------------

# Exact public column set (stable shape; no extra keys).
INSPIRED_COLUMNS: tuple[str, ...] = (
    "Filename",
    "Id",
    "Country",
    "Season",
    "Match Date",
   # "Rights Holders",
    "Home Team",
    "Away Team",
    "Attacking Team",
    "Defending Team",
    "LeftToRight",
    "Action",
    "Goal Scored",
    "Penalty Shootout",
    "Reverse Goal",
    "Goal Type",
    "Goal Time (Frames)",
    "Duration (Frames)",
)

InspiredDataFrame = list[dict[str, str]]  # each dict: exactly INSPIRED_COLUMNS keys → str

BOOL_COLUMNS: frozenset[str] = frozenset({"Goal Scored"})
COERCED_ENGP_BOOL_COLUMNS: frozenset[str] = frozenset({"Penalty Shootout", "Reverse Goal"})
DATE_COLUMNS: frozenset[str] = frozenset({"Match Date"})
REQUIRED_COLUMNS_LOCAL_CSV: frozenset[str] = frozenset(
    {"Filename", "Country", "Season", "Match Date", "Home Team", "Away Team"}
)
REQUIRED_COLUMNS_ICONIK_API: frozenset[str] = frozenset(
    REQUIRED_COLUMNS_LOCAL_CSV | {"Id"}
)
# Back-compat alias for Iconik / live API validation paths.
REQUIRED_COLUMNS: frozenset[str] = REQUIRED_COLUMNS_ICONIK_API

_ID_SOURCE_FLAT_KEYS: tuple[str, ...] = ("Id", "id", "asset_id")


def required_columns_for_delivery(*, csv_only_delivery: bool) -> frozenset[str]:
    if csv_only_delivery:
        return REQUIRED_COLUMNS_LOCAL_CSV
    return REQUIRED_COLUMNS_ICONIK_API


def _id_value_from_flat(flat: dict[str, object], *, cf_index: dict[str, str]) -> str:
    """First non-blank Id / id / asset_id from flat (exact key or casefold index)."""
    for key in _ID_SOURCE_FLAT_KEYS:
        if key in flat and not _is_blankish(flat[key]):
            return _cleanup_csv_cell(flat[key])
        fk = cf_index.get(key.casefold())
        if fk is not None and not _is_blankish(flat.get(fk)):
            return _cleanup_csv_cell(flat[fk])
    return ""

# Explicit deterministic mapping: Inspired column → ordered Iconik / flat metadata keys to try (first match wins).
INSPIRED_COLUMN_SOURCES: dict[str, tuple[str, ...]] = {
    "Filename": (),
    "Id": (),
    # Exact-case Iconik slugs first (case-sensitive), then legacy lowercase / display aliases.
    "Country": ("Country_ENGP", "country", "Country", "COUNTRY"),
    "Season": ("Season_ENGP", "season", "Season", "SEASON"),
    "Match Date": ("Match_Date_ENGP", "match_date", "Match Date", "match date", "Date", "date"),
    "Rights Holders": ("Rights_Holders_ENGP", "rights_holders", "Rights Holders", "rights holders", "RightsHolder"),
    "Home Team": ("Home_Team_ENGP", "home_team", "Home Team", "home team", "HomeTeam"),
    "Away Team": ("Away_Team_ENGP", "away_team", "Away Team", "away team", "AwayTeam"),
    "Attacking Team": (
        "Attacking_Team_ENGP",
        "attacking_team",
        "Attacking Team",
        "attacking team",
        "AttackingTeam",
    ),
    "Defending Team": (
        "Defending_Team_ENGP",
        "defending_team",
        "Defending Team",
        "defending team",
        "DefendingTeam",
    ),
    # Derived only from Direction_of_Play_ENGP (see normalize_left_to_right_from_direction).
    "LeftToRight": ("Direction_of_Play_ENGP",),
    "Action": ("Action_ENGP", "action", "Action", "ACTION"),
    "Goal Scored": ("Goal_Scored_ENGP", "goal_scored", "Goal Scored", "goal scored", "GoalScored"),
    "Penalty Shootout": ("Penalty_Shootout_ENGP",),
    "Reverse Goal": ("Reverse_Goal_ENGP",),
    "Goal Type": ("Goal_Type_ENGP", "goal_type", "Goal Type", "goal type", "GoalType"),
    "Goal Time (Frames)": (
        "Goal_Time_Seconds_ENGP",
        "Goal_Time_seconds_ENGP",
        "goal_time_seconds_engp",
    ),
    "Duration (Frames)": (
        "Duration_Frames_ENGP",
        "Duration_ENGP",
        "duration_frames",
        "Duration (Frames)",
        "duration",
        "duration (frames)",
        "Duration",
    ),
}

_BLANK_TOKENS = frozenset(
    {
        "",
        "nan",
        "none",
        "null",
        "na",
        "n/a",
    }
)


def _is_blankish(value: object) -> bool:
    if value is None:
        return True
    if isinstance(value, float) and math.isnan(value):
        return True
    if isinstance(value, str):
        t = value.strip()
        if not t:
            return True
        return t.casefold() in _BLANK_TOKENS
    return False


def _unwrap_iconik_field_value(raw: object) -> object:
    """Reduce Iconik field objects (field_values or values arrays) to Python primitives or lists."""
    if not isinstance(raw, dict):
        return raw

    if isinstance(raw.get("field_values"), list):
        fvs = raw.get("field_values")
        vals: list[object] = []
        for fv in fvs:
            if isinstance(fv, dict) and "value" in fv:
                vals.append(fv.get("value"))
            else:
                vals.append(fv)
        if not vals:
            return None
        if len(vals) == 1:
            return vals[0]
        return vals

    if isinstance(raw.get("values"), list):
        vsl = raw.get("values")
        vals = []
        for x in vsl:
            if isinstance(x, dict) and "value" in x:
                vals.append(x.get("value"))
            else:
                vals.append(x)
        if not vals:
            return None
        if len(vals) == 1:
            return vals[0]
        return vals

    return raw


def _is_iconik_field_leaf_dict(node: object) -> bool:
    if not isinstance(node, dict):
        return False
    return isinstance(node.get("field_values"), list) or isinstance(node.get("values"), list)


def iconik_search_hit_metadata_version_id(hit: dict | None) -> str | None:
    """
    Version UUID for metadata GET — same precedence as Inspired_Highlight_version.py:
    ``latest_version_id`` then ``version_id`` on the search hit.
    """
    if not hit or not isinstance(hit, dict):
        return None
    v = hit.get("latest_version_id") or hit.get("version_id")
    if v is None:
        return None
    if isinstance(v, str):
        s = v.strip()
        return s or None
    if isinstance(v, int) and not isinstance(v, bool):
        return str(v)
    return None


def extract_iconik_tenant_metadata_flat(data: dict | None) -> dict[str, object]:
    """
    Flatten metadata from an Iconik JSON document using the same rules as
    ``Inspired_Highlight_version._extract_meta`` (known-good for this tenant).

    Priority (mutually exclusive branches):
    1. Top-level slug → dict with ``values`` list (Iconik metadata_values-style map).
    2. Else ``view_fields`` list: each item ``name`` + ``field_values`` (or ``value``).
    3. Else ``metadata`` dict: list / dict-with-field_values / scalar values.

    Keys keep Iconik casing; ``__*`` keys are skipped. Values are coerced like the reference (mostly str).
    """
    if not isinstance(data, dict):
        return {}
    fd: dict[str, object] = {}
    is_d = False
    for k, v in data.items():
        sk = str(k)
        if sk.startswith("__"):
            continue
        if isinstance(v, dict) and isinstance(v.get("values"), list):
            is_d = True
            vals = [str(x.get("value", "")) for x in v["values"] if str(x.get("value", ""))]
            fd[sk] = vals[0] if len(vals) == 1 else (vals if vals else "")
    if is_d:
        return fd
    if isinstance(data.get("view_fields"), list):
        res: dict[str, object] = {}
        for f in data["view_fields"]:
            if not isinstance(f, dict):
                continue
            n = f.get("name", "")
            if not n or (isinstance(n, str) and str(n).startswith("__")):
                continue
            fv = f.get("field_values", [])
            vals = (
                [str(x.get("value", "")) for x in fv if x.get("value")]
                if isinstance(fv, list)
                else []
            )
            res[str(n)] = vals[0] if len(vals) == 1 else (vals or f.get("value", ""))
        return res
    if isinstance(data.get("metadata"), dict):
        res: dict[str, object] = {}
        for k, v in data["metadata"].items():
            sk = str(k)
            if isinstance(v, list):
                res[sk] = v
            elif isinstance(v, dict):
                fv = v.get("field_values", [])
                vals = (
                    [str(x.get("value", "")) for x in fv if x.get("value")]
                    if isinstance(fv, list)
                    else []
                )
                res[sk] = vals[0] if len(vals) == 1 else (vals or "")
            else:
                res[sk] = v
        return res
    return {}


def build_flat_metadata_map(meta_doc: dict | None) -> tuple[dict[str, object], tuple[str, ...]]:
    """
    Flatten Iconik metadata by walking metadata_values, metadata, and view_fields (recursively).

    Keys in ``flat`` preserve the original slug / field name casing (first writer wins per exact key).
    Discovered slugs are the sorted list of those keys (case-insensitive sort for stable output).
    """
    flat: dict[str, object] = {}
    if not meta_doc or not isinstance(meta_doc, dict):
        return flat, ()

    MAX_DEPTH = 96

    def remember_scalar(slug: str, value: object) -> None:
        if not slug or slug in flat:
            return
        flat[slug] = value

    def remember(slug: str, leaf: dict) -> None:
        if not slug or slug in flat:
            return
        flat[slug] = _unwrap_iconik_field_value(leaf)

    def walk_mapping(node: object, depth: int) -> None:
        if depth > MAX_DEPTH or node is None:
            return
        if isinstance(node, dict):
            for k in sorted(node.keys(), key=lambda x: str(x).casefold()):
                v = node[k]
                sk = str(k)
                if isinstance(v, dict) and _is_iconik_field_leaf_dict(v):
                    remember(sk, v)
                elif isinstance(v, dict):
                    walk_mapping(v, depth + 1)
                elif isinstance(v, list):
                    for it in v:
                        walk_list_item(it, depth + 1)
                elif not isinstance(v, (dict, list)):
                    # Slug → primitive (some projections / merged views); includes False / 0.
                    remember_scalar(sk, v)
        elif isinstance(node, list):
            for it in node:
                walk_list_item(it, depth + 1)

    def walk_list_item(it: object, depth: int) -> None:
        if depth > MAX_DEPTH or it is None:
            return
        if isinstance(it, dict):
            name = (
                it.get("name")
                or it.get("slug")
                or it.get("field")
                or it.get("key")
                or it.get("field_name")
                or it.get("metadata_slug")
            )
            if isinstance(name, str) and _is_iconik_field_leaf_dict(it):
                remember(name, it)
                return
            walk_mapping(it, depth + 1)
        elif isinstance(it, list):
            walk_mapping(it, depth + 1)

    for root_key in ("metadata_values", "metadata", "view_fields"):
        root = meta_doc.get(root_key)
        if root is not None:
            walk_mapping(root, 0)

    ordered_slugs = tuple(sorted(flat.keys(), key=lambda s: (str(s).casefold(), str(s))))
    return flat, ordered_slugs


def merge_flat_metadata_prefer_primary(
    primary: dict[str, object],
    secondary: dict[str, object],
) -> dict[str, object]:
    """First map wins for non-blank values; secondary fills only missing or blankish keys (never blank-overwrites)."""
    out = dict(primary)
    for k, v in secondary.items():
        if _is_blankish(v):
            continue
        cur = out.get(k)
        if k not in out or _is_blankish(cur):
            out[k] = v
    return out


def _iter_search_hit_metadata_fragments(hit: dict) -> list[dict[str, object]]:
    """
    Collect minimal {metadata_values|metadata|view_fields} slices found anywhere on the search hit.
    Bounded traversal (Iconik search projections may nest these under asset/document/object, etc.).
    """
    fragments: list[dict[str, object]] = []
    roots = ("metadata_values", "metadata", "view_fields")
    seen: set[int] = set()
    stack: list[tuple[object, int]] = [(hit, 0)]
    while stack:
        node, depth = stack.pop()
        if depth > 14:
            continue
        if isinstance(node, dict):
            nid = id(node)
            if nid in seen:
                continue
            seen.add(nid)
            frag = {k: node[k] for k in roots if k in node and node[k] is not None}
            if frag:
                fragments.append(frag)
            for v in list(node.values())[:120]:
                if isinstance(v, dict):
                    stack.append((v, depth + 1))
                elif isinstance(v, list):
                    for it in v[:100]:
                        if isinstance(it, dict):
                            stack.append((it, depth + 1))
        elif isinstance(node, list):
            for it in node[:100]:
                if isinstance(it, dict):
                    stack.append((it, depth + 1))
    return fragments


def flatten_metadata_from_search_hit(hit: dict | None) -> tuple[dict[str, object], tuple[str, ...]]:
    """Flatten tenant metadata embedded in the Iconik search hit (projection / view_fields / metadata)."""
    if not hit or not isinstance(hit, dict):
        return {}, ()
    flat_acc: dict[str, object] = {}
    for frag in _iter_search_hit_metadata_fragments(hit):
        f, _ = build_flat_metadata_map(frag)
        flat_acc = merge_flat_metadata_prefer_primary(flat_acc, f)
    slugs = tuple(sorted(flat_acc.keys(), key=lambda s: (str(s).casefold(), str(s))))
    return flat_acc, slugs


def print_search_hit_flattened_metadata_diagnostics(
    flat_hit: dict[str, object],
    slugs_hit: tuple[str, ...],
    emit,
) -> None:
    if not VERBOSE_METADATA_DIAGNOSTICS:
        return
    _diag_emit_line("SEARCH HIT FLATTENED KEYS:", emit)
    if not flat_hit:
        _diag_emit_line("  (none)", emit)
    else:
        for k in sorted(flat_hit.keys(), key=lambda s: (str(s).casefold(), str(s))):
            _diag_emit_line(f"  {k}", emit)
    _diag_emit_line("SEARCH HIT DISCOVERED SLUGS:", emit)
    if not slugs_hit:
        _diag_emit_line("  (none)", emit)
    else:
        for s in slugs_hit:
            _diag_emit_line(f"  {s}", emit)


def print_metadata_flatten_compare_diagnostics(
    flat_hit: dict[str, object],
    slugs_hit: tuple[str, ...],
    flat_api: dict[str, object],
    slugs_api: tuple[str, ...],
    emit,
) -> None:
    if not VERBOSE_METADATA_DIAGNOSTICS:
        return
    _diag_emit_line("METADATA GET (endpoint) FLATTENED KEYS:", emit)
    if not flat_api:
        _diag_emit_line("  (none)", emit)
    else:
        for k in sorted(flat_api.keys(), key=lambda s: (str(s).casefold(), str(s))):
            _diag_emit_line(f"  {k}", emit)
    _diag_emit_line("METADATA GET DISCOVERED SLUGS:", emit)
    if not slugs_api:
        _diag_emit_line("  (none)", emit)
    else:
        for s in slugs_api:
            _diag_emit_line(f"  {s}", emit)
    only_hit = sorted(set(flat_hit) - set(flat_api), key=lambda s: (str(s).casefold(), str(s)))
    only_api = sorted(set(flat_api) - set(flat_hit), key=lambda s: (str(s).casefold(), str(s)))
    _diag_emit_line("FLATTEN DIFF — KEYS ONLY IN SEARCH HIT:", emit)
    _diag_emit_line("  " + (", ".join(only_hit) if only_hit else "(none)"), emit)
    _diag_emit_line("FLATTEN DIFF — KEYS ONLY IN METADATA GET:", emit)
    _diag_emit_line("  " + (", ".join(only_api) if only_api else "(none)"), emit)


def _enumerate_metadata_json_paths(obj: object, prefix: str = "") -> list[str]:
    """Deterministic depth-first paths (dot keys, bracket indices) for inspection only."""
    paths: list[str] = []
    if isinstance(obj, dict):
        for k in sorted(obj.keys(), key=lambda x: str(x).casefold()):
            seg = f"{prefix}.{k}" if prefix else str(k)
            paths.append(seg)
            paths.extend(_enumerate_metadata_json_paths(obj[k], seg))
    elif isinstance(obj, list):
        for i, item in enumerate(obj):
            seg = f"{prefix}[{i}]"
            paths.append(seg)
            paths.extend(_enumerate_metadata_json_paths(item, seg))
    return paths


def _find_first_field_values_under_doc(meta_doc: dict) -> str | None:
    for root_key in ("metadata_values", "metadata", "view_fields"):
        node = meta_doc.get(root_key)
        hit = _find_first_field_values_representation(node)
        if hit:
            return f"(under {root_key}) {hit}"
    return None


def _find_first_field_values_representation(node: object, *, max_depth: int = 32) -> str | None:
    """Return repr of the first ``field_values`` or ``values`` list found under node (diagnostics)."""
    if max_depth <= 0 or node is None:
        return None
    if isinstance(node, dict):
        fv = node.get("field_values")
        if isinstance(fv, list):
            return repr(fv)[:8000]
        vs = node.get("values")
        if isinstance(vs, list):
            return repr(vs)[:8000]
        for k in sorted(node.keys(), key=lambda x: str(x).casefold()):
            sub = _find_first_field_values_representation(node[k], max_depth=max_depth - 1)
            if sub:
                return f"(under key {k!r}) {sub}"
    elif isinstance(node, list):
        for i, it in enumerate(node[:20]):
            sub = _find_first_field_values_representation(it, max_depth=max_depth - 1)
            if sub:
                return f"(at index [{i}]) {sub}"
    return None


def gather_metadata_structure_diagnostic(meta_doc: dict) -> dict[str, object]:
    top = sorted(meta_doc.keys(), key=str.casefold)
    mv = meta_doc.get("metadata_values")
    mv_type = type(mv).__name__
    if mv is None:
        mv_len: object = "(absent)"
        first_mv_entry = None
    elif isinstance(mv, (dict, list)):
        mv_len = len(mv)
        if isinstance(mv, dict) and mv:
            fk = sorted(mv.keys(), key=lambda x: str(x).casefold())[0]
            first_mv_entry = repr({fk: mv[fk]})[:8000]
        elif isinstance(mv, list) and mv:
            first_mv_entry = repr(mv[0])[:8000]
        else:
            first_mv_entry = None
    else:
        mv_len = "(not dict/list)"
        first_mv_entry = repr(mv)[:4000]
    vf = meta_doc.get("view_fields")
    vf_type = type(vf).__name__
    vf_len: object
    if vf is None:
        vf_len = "(absent)"
    elif isinstance(vf, (dict, list)):
        vf_len = len(vf)
    else:
        vf_len = "(not dict/list)"
    first_fv = _find_first_field_values_under_doc(meta_doc)
    return {
        "top_level_metadata_keys": top,
        "metadata_values_type": mv_type,
        "metadata_values_length": mv_len,
        "first_metadata_values_entry_repr": first_mv_entry,
        "view_fields_type": vf_type,
        "view_fields_length": vf_len,
        "first_field_values_entry_repr": first_fv,
    }


def print_metadata_values_structure_diagnostic(meta_doc: dict, emit=None) -> None:
    """Temporary stdout (+ optional delivery_log) diagnostics for metadata_values shape."""
    if not VERBOSE_METADATA_DIAGNOSTICS:
        return
    d = gather_metadata_structure_diagnostic(meta_doc)

    def line(s: str = "") -> None:
        print(s, flush=True)
        if emit is not None:
            emit(s)

    line("TOP LEVEL METADATA KEYS:")
    for k in d["top_level_metadata_keys"]:
        line(k)
    line("METADATA_VALUES TYPE:")
    line(str(d.get("metadata_values_type", "")))
    line("METADATA_VALUES LENGTH:")
    line(str(d.get("metadata_values_length", "")))
    line("FIRST METADATA_VALUES ENTRY:")
    fme = d.get("first_metadata_values_entry_repr")
    line(fme if isinstance(fme, str) else "(none)")
    line("VIEW_FIELDS TYPE:")
    line(str(d.get("view_fields_type", "")))
    line("VIEW_FIELDS LENGTH:")
    line(str(d.get("view_fields_length", "")))
    line("FIRST FIELD_VALUES ENTRY:")
    ffv = d.get("first_field_values_entry_repr")
    line(ffv if isinstance(ffv, str) else "(none)")


def export_raw_metadata_debug_json(
    delivery_dir: Path,
    meta_doc: dict,
    *,
    resolved_asset_id: str,
    resolved_title: str,
    emit=None,
    flat_from_search_hit: dict[str, object] | None = None,
    slugs_from_search_hit: tuple[str, ...] | None = None,
    flat_from_metadata_endpoint: dict[str, object] | None = None,
    flat_merged: dict[str, object] | None = None,
) -> None:
    """Diagnostics only: first resolved asset with metadata; compares search projection vs GET body."""
    if VERBOSE_METADATA_DIAGNOSTICS:
        print_metadata_values_structure_diagnostic(meta_doc, emit=emit)
    flat_api = flat_from_metadata_endpoint
    if flat_api is None:
        flat_api, _ = build_flat_metadata_map(meta_doc)
    merged = flat_merged if flat_merged is not None else flat_api
    discovered_slugs = tuple(sorted(merged.keys(), key=lambda s: (str(s).casefold(), str(s))))
    paths = sorted(set(_enumerate_metadata_json_paths(meta_doc)), key=str.casefold)
    structure = gather_metadata_structure_diagnostic(meta_doc)
    if VERBOSE_METADATA_DIAGNOSTICS:
        print("DISCOVERED METADATA KEYS (merged canonical flat):", flush=True)
        if emit is not None:
            emit("DISCOVERED METADATA KEYS (merged canonical flat):")
        for slug in discovered_slugs:
            print(slug, flush=True)
            if emit is not None:
                emit(slug)
    payload: dict[str, object] = {
        "_debug_note": "First resolved asset with metadata GET 200 — search-hit vs endpoint flatten comparison.",
        "resolved_asset_id": resolved_asset_id,
        "resolved_title": resolved_title,
        "structure_diagnostic": structure,
        "raw_metadata_response": meta_doc,
        "flattened_metadata_endpoint": flat_api,
        "flattened_search_hit": flat_from_search_hit,
        "discovered_field_keys_search_hit": list(slugs_from_search_hit or ()),
        "flattened_metadata_merged": merged,
        "flattened_metadata": merged,
        "discovered_field_keys_merged": list(discovered_slugs),
        "discovered_field_keys": list(discovered_slugs),
        "discovered_metadata_paths": paths,
        "flattened_map_keys": sorted(merged.keys(), key=lambda s: (str(s).casefold(), str(s))),
    }
    out_path = delivery_dir / "raw_metadata_debug.json"
    out_path.write_text(json.dumps(payload, indent=2, ensure_ascii=False, default=str) + "\n", encoding="utf-8")


def export_raw_search_result_debug_json(
    delivery_dir: Path,
    raw_search_hit: dict,
    *,
    source_filename: str,
    resolved_asset_id: str,
    resolved_title: str,
    emit=None,
) -> None:
    """
    First resolved asset only: persist the search API hit dict unchanged (JSON envelope for context).
    """
    envelope: dict[str, object] = {
        "_debug_note": (
            "First resolved asset — raw search hit object from Iconik search/v1/search/ response "
            "(objects[] element). raw_search_hit is not flattened or filtered by this app."
        ),
        "source_filename": source_filename,
        "resolved_asset_id": resolved_asset_id,
        "resolved_title": resolved_title,
        "raw_search_hit": raw_search_hit,
    }
    out_path = delivery_dir / "raw_search_result_debug.json"
    out_path.write_text(json.dumps(envelope, indent=2, ensure_ascii=False, default=str) + "\n", encoding="utf-8")
    msg = f"Wrote raw_search_result_debug.json (first resolved asset, source file {source_filename!r})."
    print(msg, flush=True)
    if emit is not None:
        emit(msg)


def _casefold_first_key_index(flat: dict[str, object]) -> dict[str, str]:
    """Map lowercase → first original flat key (deterministic when multiple keys share casing)."""
    idx: dict[str, str] = {}
    for fk in sorted(flat.keys(), key=lambda x: (str(x).casefold(), str(x))):
        lk = str(fk).casefold()
        if lk not in idx:
            idx[lk] = fk
    return idx


def _get_mapped_raw(flat: dict[str, object], column: str, *, cf_index: dict[str, str]) -> object:
    if column in ("Filename", "Id"):
        return None
    if column == "Goal Time (Frames)":
        fk = cf_index.get(_GOAL_TIME_SECONDS_SLUG_CF)
        if fk is not None:
            return flat[fk]
    keys = INSPIRED_COLUMN_SOURCES.get(column, ())
    for k in keys:
        if k in flat:
            return flat[k]
    for k in keys:
        fk = cf_index.get(k.casefold())
        if fk is not None:
            return flat[fk]
    return None


def _flatten_list_to_csv(parts: list[str]) -> str:
    cleaned = [p.strip() for p in parts if p.strip() and p.strip().casefold() not in _BLANK_TOKENS]
    return ",".join(cleaned)


def _normalize_boolean(value: object) -> tuple[str, bool]:
    """
    Return (output, invalid).
    Output is only 'TRUE' or 'FALSE' for Inspired booleans (never blank).
    invalid True when the value was missing/blank (coerced to FALSE) or not recognized (coerced to FALSE).
    """
    if _is_blankish(value):
        return "FALSE", True
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE", False
    if isinstance(value, int) and not isinstance(value, bool):
        if value == 1:
            return "TRUE", False
        if value == 0:
            return "FALSE", False
        return "FALSE", True
    if isinstance(value, float) and not math.isnan(value):
        if value == 1.0:
            return "TRUE", False
        if value == 0.0:
            return "FALSE", False
        return "FALSE", True
    s = str(value).strip()
    cf = s.casefold()
    if cf in ("yes", "true", "1", "y", "t"):
        return "TRUE", False
    if cf in ("no", "false", "0", "n", "f"):
        return "FALSE", False
    return "FALSE", True


def _normalize_date_to_dd_mm_yyyy(value: object) -> tuple[str, bool]:
    """Return ('DD/MM/YYYY' or '', malformed)."""
    if _is_blankish(value):
        return "", False
    if isinstance(value, (datetime, date)):
        if isinstance(value, datetime):
            d = value.date()
        else:
            d = value
        return d.strftime("%d/%m/%Y"), False
    s = str(value).strip()
    if not s:
        return "", False
    if "T" in s:
        s = s.split("T", 1)[0].strip()
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", s):
        try:
            d = datetime.strptime(s, "%Y-%m-%d").date()
            return d.strftime("%d/%m/%Y"), False
        except ValueError:
            return "", True
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%Y/%m/%d"):
        try:
            d = datetime.strptime(s, fmt).date()
            return d.strftime("%d/%m/%Y"), False
        except ValueError:
            continue
    return "", True


def _normalize_string_separators(s: str) -> str:
    s = s.strip()
    s = re.sub(r"\s+", " ", s)
    s = re.sub(r",+", ",", s)
    s = re.sub(r";+", ";", s)
    s = re.sub(r"\|+", "|", s)
    s = re.sub(r",\s*,", ",", s)
    s = re.sub(r",\s*,", ",", s)
    s = s.strip(",;| ")
    return s


def _cleanup_csv_cell(value: object) -> str:
    """Normalize CSV-side text: newlines in quoted cells → comma list, whitespace, duplicate separators."""
    if value is None:
        return ""
    if isinstance(value, float) and math.isnan(value):
        return ""
    s = str(value).replace("\r\n", "\n").replace("\r", "\n")
    s = s.replace("\n", ",")
    return _normalize_string_separators(s)


def _split_field_tokens(s: str) -> list[str]:
    return [p.strip() for p in re.split(r"[,;|]+", s) if p.strip()]


def _unique_token_join(tokens: Iterable[str]) -> str:
    parts: list[str] = []
    seen: set[str] = set()
    for t in tokens:
        cf = t.casefold()
        if cf in seen:
            continue
        seen.add(cf)
        parts.append(t)
    return ",".join(parts)


def _normalize_miss_save_token_field(s: str) -> str:
    """If both Miss and Save appear as whole tokens, drop Miss (directional merge rule)."""
    toks = _split_field_tokens(s)
    cfs = {t.casefold() for t in toks}
    if "miss" in cfs and "save" in cfs:
        toks = [t for t in toks if t.casefold() != "miss"]
    return _unique_token_join(toks)


def _merge_csv_field_tokens(
    a: str,
    b: str,
    *,
    field_key: str,
    merge_warnings: list[str] | None = None,
) -> str:
    combined = _unique_token_join(_split_field_tokens(a) + _split_field_tokens(b))
    return _normalize_miss_save_token_field(combined)


def primary_flat_slug_for_inspired_column(col: str) -> str | None:
    """Canonical *_ENGP (or first) flat key used when overlaying CSV onto Iconik flat."""
    if col == "LeftToRight":
        return "Direction_of_Play_ENGP"
    if col in ("Filename", "Id"):
        return None
    tup = INSPIRED_COLUMN_SOURCES.get(col, ())
    return tup[0] if tup else None


def _csv_overlay_value_into_flat(
    out: dict[str, object],
    dest_key: str,
    raw_val: object,
    *,
    merge_warnings: list[str] | None = None,
) -> None:
    """Blank CSV never overwrites; differing non-blank values combine as unique comma-separated tokens."""
    vv = _cleanup_csv_cell(raw_val) if not isinstance(raw_val, (dict, list, tuple)) else _cleanup_csv_cell(str(raw_val))
    if not vv.strip():
        return
    cur = out.get(dest_key)
    if cur is None or _is_blankish(cur):
        out[dest_key] = vv
        return
    cs = str(cur).strip()
    if cs == vv or cs.casefold() == vv.casefold():
        return
    merged = _merge_csv_field_tokens(cs, vv, field_key=dest_key, merge_warnings=merge_warnings)
    out[dest_key] = merged


def _normalize_penalty_reverse_bool(raw: object) -> str:
    """1 / TRUE / YES → TRUE; anything else (including blank) → FALSE."""
    if _is_blankish(raw):
        return "FALSE"
    if isinstance(raw, bool):
        return "TRUE" if raw else "FALSE"
    if isinstance(raw, int) and not isinstance(raw, bool) and raw == 1:
        return "TRUE"
    if isinstance(raw, float) and not math.isnan(raw) and float(raw) == 1.0:
        return "TRUE"
    s = str(raw).strip()
    if s == "1":
        return "TRUE"
    cf = s.casefold()
    if cf in ("true", "yes"):
        return "TRUE"
    return "FALSE"


def _count_non_blank_csv_cells(rows: Iterable[dict[str, str]]) -> int:
    n = 0
    for r in rows:
        for k, v in r.items():
            if str(k).startswith("__"):
                continue
            if _cleanup_csv_cell(v).strip():
                n += 1
    return n


def _merged_row_has_inspired_column(merged: dict[str, str], column: str) -> bool:
    if column == "Id":
        for key in _ID_SOURCE_FLAT_KEYS:
            if merged.get(key, "").strip():
                return True
            for mk in merged:
                if mk.casefold() == key.casefold() and merged[mk].strip():
                    return True
        return False
    for slug in INSPIRED_COLUMN_SOURCES.get(column, ()):
        if merged.get(slug, "").strip():
            return True
    ck = column.casefold()
    for mk in merged:
        if mk.casefold() == ck and merged[mk].strip():
            return True
    return False


def merge_per_asset_csv_rows_into_one_record(
    rows: list[dict[str, str]],
    csv_path: Path,
    *,
    filename_column: str,
    derived_media_filename: str | None,
) -> tuple[dict[str, str], dict[str, str]]:
    """
    Collapse all data rows from one per-asset Iconik CSV into a single record.
    Blanks never overwrite; multiple values merge as unique comma-separated tokens (Miss+Save rule applied).
    """
    merge_warnings: list[str] = []
    keys_ever_nonblank: set[str] = set()
    for r in rows:
        for k, v in r.items():
            if str(k).startswith("__"):
                continue
            if _cleanup_csv_cell(v).strip():
                keys_ever_nonblank.add(k)

    merged: dict[str, str] = {}
    multi_row_keys: set[str] = set()
    for r in rows:
        for k, v in r.items():
            if str(k).startswith("__"):
                continue
            vv = _cleanup_csv_cell(v)
            if not vv.strip():
                continue
            if k not in merged or not merged[k].strip():
                merged[k] = vv
            elif merged[k].strip() == vv.strip():
                continue
            else:
                before = merged[k]
                merged[k] = _merge_csv_field_tokens(before, vv, field_key=k, merge_warnings=merge_warnings)
                multi_row_keys.add(k)

    for k in list(merged.keys()):
        before = merged[k]
        after = _normalize_miss_save_token_field(before)
        if after != before:
            merge_warnings.append(
                f"CSV {csv_path.name!s} field {k!r}: normalized {before!r} → {after!r} (Miss+Save rule)"
            )
            merged[k] = after

    if derived_media_filename is not None:
        merged[filename_column] = derived_media_filename

    derived_fn = _cleanup_csv_cell(merged.get(filename_column, ""))
    if not derived_fn and derived_media_filename:
        derived_fn = derived_media_filename.strip()

    row_count = len(rows)
    non_blank = _count_non_blank_csv_cells(rows)
    merged_field_count = sum(1 for k, v in merged.items() if not str(k).startswith("__") and v.strip())

    multi_fields = sorted(multi_row_keys)

    missing_after = sorted(k for k in keys_ever_nonblank if not merged.get(k, "").strip())

    missing_req: list[str] = []
    for req in sorted(REQUIRED_COLUMNS_LOCAL_CSV):
        if not _merged_row_has_inspired_column(merged, req):
            missing_req.append(req)

    diag: dict[str, str] = {
        "csv_path": str(csv_path),
        "derived_filename": derived_fn,
        "row_count": str(row_count),
        "non_blank_field_count": str(non_blank),
        "merged_field_count": str(merged_field_count),
        "missing_required_fields_after_merge": ", ".join(missing_req),
        "fields_merged_from_multiple_rows": ", ".join(multi_fields),
        "fields_present_in_csv_but_missing_after_merge": ", ".join(missing_after),
        "Goal_Time_Seconds_ENGP": merged.get("Goal_Time_Seconds_ENGP", "").strip(),
        "merge_warnings": "; ".join(merge_warnings),
    }
    return merged, diag


METADATA_CSV_NAMES: tuple[str, ...] = ("metadata.csv", "combined_metadata.csv", "inspired_metadata_source.csv")

# Iconik per-asset CSV exports often omit a filename column; the media name is only in the CSV
# basename (e.g. clip_ENGP.mp4.csv → clip_ENGP.mp4). Never treat arbitrary metadata slugs
# (e.g. Action_ENGP) as filename columns.
_CSV_EXPLICIT_FILENAME_KEYS_IN_ORDER: tuple[str, ...] = (
    "filename",
    "original_filename",
    "asset_name",
    "title",
)
_CSV_EXTRA_TITLE_KEYS: frozenset[str] = frozenset(
    {"display_title", "editorial_title", "asset_title", "match_title"}
)


def _csv_header_normalize(h: str) -> str:
    return h.strip().casefold().replace(" ", "_").replace("-", "_")


def media_filename_from_csv_basename(csv_path: Path) -> str:
    """Strip only a final .csv suffix (case-insensitive); remainder is the media filename."""
    name = csv_path.name
    if name.lower().endswith(".csv"):
        return name[:-4]
    return name


def pick_explicit_csv_filename_column(fieldnames: Iterable[str]) -> str:
    """First matching header (by normalized key), in product preference order; empty if none."""
    headers = [h.strip() for h in fieldnames if str(h).strip()]
    buckets: dict[str, str] = {}
    for h in headers:
        key = _csv_header_normalize(h)
        if key not in _CSV_EXPLICIT_FILENAME_KEYS_IN_ORDER:
            continue
        buckets.setdefault(key, h)
    for want in _CSV_EXPLICIT_FILENAME_KEYS_IN_ORDER:
        if want in buckets:
            return buckets[want]
    return ""


def pick_csv_title_column(headers: list[str], explicit_filename_header: str, assigned_filename_key: str) -> str:
    """
    Optional second column for tier-3 matching. ``explicit_filename_header`` is the real CSV
    header chosen as filename, or "" when filename is basename-derived (synthetic Filename).
    """
    ex_cf = explicit_filename_header.casefold() if explicit_filename_header else None
    assigned_cf = assigned_filename_key.casefold()
    for h in headers:
        if ex_cf is not None and h.casefold() == ex_cf:
            continue
        if h.casefold() == assigned_cf:
            continue
        kn = _csv_header_normalize(h)
        if kn in _CSV_EXTRA_TITLE_KEYS:
            return h
    for h in headers:
        if ex_cf is not None and h.casefold() == ex_cf:
            continue
        if h.casefold() == assigned_cf:
            continue
        kn = _csv_header_normalize(h)
        if kn == "title" and (not explicit_filename_header or _csv_header_normalize(explicit_filename_header) != "title"):
            return h
    return ""


def resolve_csv_identity_columns(
    fieldnames: Iterable[str], csv_path: Path
) -> tuple[str, str, str | None]:
    """
    Returns (filename_column_key_in_row_dicts, title_column_or_empty, derived_media_name_or_none).
    When no explicit filename header exists, filename_column_key is the synthetic ``Filename``
    and derived_media_name is the basename rule result (applied to every row for that file).
    """
    headers = [h.strip() for h in fieldnames if str(h).strip()]
    explicit = pick_explicit_csv_filename_column(headers)
    if explicit:
        fn_col = explicit
        derived: str | None = None
    else:
        fn_col = "Filename"
        derived = media_filename_from_csv_basename(csv_path)
    title_col = pick_csv_title_column(headers, explicit, fn_col)
    return fn_col, title_col, derived


def discover_optional_metadata_csv(mp4_paths: list[Path]) -> Path | None:
    """First existing sidecar CSV alongside input MP4s (parent dirs, list order)."""
    seen: set[str] = set()
    for p in mp4_paths:
        try:
            parent = p.parent
        except (OSError, ValueError):
            continue
        key = str(parent)
        if key in seen:
            continue
        seen.add(key)
        for name in METADATA_CSV_NAMES:
            cand = parent / name
            try:
                if cand.is_file():
                    return cand
            except OSError:
                continue
    return None


def load_merge_metadata_csv_bundle(
    explicit_paths: Sequence[Path | str],
    *,
    mp4_paths: list[Path],
) -> tuple[list[dict[str, str]], str, str, str, str | None, list[dict[str, str]]]:
    """
    Load operator-selected CSV path(s) in order. Each physical CSV is collapsed to **one** merged
    row (all data rows within that file combined). Each row includes ``__source_path__`` for reports.
    Returns ``csv_merge_debug_rows`` for ``reports/csv_merge_debug_report.xlsx``.
    """
    merged: list[dict[str, str]] = []
    fn_col = ""
    title_col = ""
    used_paths: list[str] = []
    err_parts: list[str] = []
    resolved_seen: set[str] = set()
    discovered: Path | None = None
    csv_merge_debug_rows: list[dict[str, str]] = []

    def add_file(path: Path) -> None:
        nonlocal fn_col, title_col, merged
        rows, err = load_delivery_metadata_csv(path)
        if err:
            err_parts.append(f"{path}: {err}")
            return
        if not rows:
            err_parts.append(f"{path}: empty")
            return
        fc_detect, tc_detect, derived_media = resolve_csv_identity_columns(list(rows[0].keys()), path)
        one_row, diag = merge_per_asset_csv_rows_into_one_record(
            rows,
            path,
            filename_column=fc_detect,
            derived_media_filename=derived_media,
        )
        csv_merge_debug_rows.append(diag)
        if not merged:
            fn_col, title_col = fc_detect, tc_detect
            one_row["__source_path__"] = str(path)
            merged.append(one_row)
        else:
            if derived_media is not None:
                one_row[fn_col] = derived_media
            else:
                one_row[fn_col] = _cleanup_csv_cell(one_row.get(fc_detect, ""))
            if title_col and tc_detect:
                one_row[title_col] = _cleanup_csv_cell(one_row.get(tc_detect, "")) if tc_detect else ""
            if tc_detect and title_col != tc_detect and tc_detect in one_row:
                del one_row[tc_detect]
            if fc_detect != fn_col and fc_detect in one_row:
                del one_row[fc_detect]
            one_row["__source_path__"] = str(path)
            merged.append(one_row)
        if str(path) not in used_paths:
            used_paths.append(str(path))

    for raw in explicit_paths:
        p = Path(os.path.expanduser(str(raw).strip()))
        if not str(p).strip():
            continue
        try:
            if not p.is_file():
                err_parts.append(f"{p}: not a file")
                continue
            key = str(p.resolve())
        except OSError as exc:
            err_parts.append(f"{p}: {exc}")
            continue
        if key in resolved_seen:
            continue
        resolved_seen.add(key)
        add_file(p)

    if not merged and mp4_paths:
        discovered = discover_optional_metadata_csv(mp4_paths)
        if discovered is not None:
            add_file(discovered)

    if used_paths:
        bundle = "; ".join(used_paths[:16]) + (" …" if len(used_paths) > 16 else "")
    elif discovered is not None:
        bundle = f"sidecar_auto:{discovered}"
    else:
        bundle = "(no CSV loaded)"
    first_err: str | None = None
    if not merged and err_parts:
        first_err = "; ".join(err_parts[:6]) + (" …" if len(err_parts) > 6 else "")
    elif not merged:
        first_err = "no CSV rows (explicit paths empty or invalid and no sidecar CSV found)"
    return merged, fn_col, title_col, bundle, first_err, csv_merge_debug_rows


def load_delivery_metadata_csv(path: Path) -> tuple[list[dict[str, str]], str | None]:
    """Read CSV as blank-safe string rows; returns (rows, error_or_none)."""
    try:
        raw = path.read_text(encoding="utf-8", errors="replace")
    except OSError as exc:
        return [], str(exc)
    rows_out: list[dict[str, str]] = []
    try:
        reader = csv.DictReader(io.StringIO(raw))
        if reader.fieldnames is None:
            return [], "CSV has no header row"
        clean_headers = []
        for h in reader.fieldnames:
            clean_headers.append(_cleanup_csv_cell(h) if h is not None else "")
        for raw_row in reader:
            row: dict[str, str] = {}
            for h, ch in zip(reader.fieldnames, clean_headers):
                key = (h or "").strip()
                if not key:
                    key = (ch or "").strip()
                if not key:
                    continue
                row[key] = _cleanup_csv_cell(raw_row.get(h, ""))
            rows_out.append(row)
    except csv.Error as exc:
        return [], str(exc)
    return rows_out, None


def _csv_row_tier(filename: str, stem: str, file_cell: str, title_cell: str) -> int:
    """Same tier semantics as Iconik matching: 1–4 strict, 99 = no match."""
    fn = filename
    fn_cf = fn.casefold()
    if file_cell == fn:
        return 1
    if file_cell:
        try:
            f_stem = Path(file_cell).stem
        except ValueError:
            f_stem = file_cell
        if f_stem == stem:
            return 2
    if title_cell == stem:
        return 3
    if title_cell == fn:
        return 3
    if file_cell and file_cell.casefold() == fn_cf:
        return 4
    if title_cell and title_cell.casefold() == fn_cf:
        return 4
    return 99


def csv_row_indices_for_mp4(
    rows: list[dict[str, str]],
    *,
    filename: str,
    stem: str,
    filename_col: str,
    title_col: str,
) -> tuple[list[int], int, bool]:
    """
    Deterministic CSV ↔ MP4 match. Returns (matching_row_indices, best_tier, ambiguous).
    ambiguous True when multiple CSV rows share the same best tier for this MP4.
    """
    scored: list[tuple[int, int]] = []
    for i, row in enumerate(rows):
        file_cell = _cleanup_csv_cell(row.get(filename_col, ""))
        title_cell = _cleanup_csv_cell(row.get(title_col, "")) if title_col else ""
        tier = _csv_row_tier(filename, stem, file_cell, title_cell)
        if tier <= 4:
            scored.append((tier, i))
    if not scored:
        return [], 99, False
    scored.sort(key=lambda t: (t[0], t[1]))
    best = scored[0][0]
    same = [idx for tier, idx in scored if tier == best]
    ambiguous = len(same) > 1
    return same, best, ambiguous


def merge_csv_row_into_iconik_flat(
    flat: dict[str, object],
    csv_row: dict[str, str],
    *,
    merge_warnings: list[str] | None = None,
) -> dict[str, object]:
    """
    Overlay CSV onto Iconik-style flat metadata. Preserves exact CSV header / slug keys; blanks never
    overwrite populated values; conflicting non-blank values merge as unique comma-separated tokens.
    """
    out: dict[str, object] = dict(flat)
    mw = merge_warnings
    cf_keys: dict[str, str] = {}
    for k in csv_row:
        if str(k).startswith("__"):
            continue
        cf_keys[str(k).casefold()] = k

    for k, v in csv_row.items():
        if str(k).startswith("__"):
            continue
        _csv_overlay_value_into_flat(out, str(k), v, merge_warnings=mw)

    for col in INSPIRED_COLUMNS:
        if col in ("Filename", "Id"):
            continue
        dest = primary_flat_slug_for_inspired_column(col)
        if not dest:
            continue
        if col == "LeftToRight":
            sk = "Direction_of_Play_ENGP".casefold()
            if sk in cf_keys:
                _csv_overlay_value_into_flat(out, dest, csv_row[cf_keys[sk]], merge_warnings=mw)
            continue
        for slug in INSPIRED_COLUMN_SOURCES.get(col, ()):
            sk = slug.casefold()
            if sk in cf_keys:
                _csv_overlay_value_into_flat(out, dest, csv_row[cf_keys[sk]], merge_warnings=mw)
                break
        ck = col.casefold()
        if ck in cf_keys:
            _csv_overlay_value_into_flat(out, dest, csv_row[cf_keys[ck]], merge_warnings=mw)
    return out


def flat_metadata_to_combined_csv_row(
    *,
    filename: str,
    resolved_asset_id: str,
    resolved_title: str,
    flat: dict[str, object] | None,
) -> dict[str, str]:
    """One UTF-8-safe string row for combined_iconik_metadata.csv (wide flattened metadata)."""
    row: dict[str, str] = {
        "Filename": xlsx_sanitize_cell(filename),
        "resolved_asset_id": xlsx_sanitize_cell(resolved_asset_id),
        "resolved_title": xlsx_sanitize_cell(resolved_title),
    }
    if not flat:
        return row
    for k, v in flat.items():
        key = str(k)
        if isinstance(v, (dict, list, tuple)):
            try:
                row[key] = json.dumps(v, ensure_ascii=False, default=str)
            except TypeError:
                row[key] = xlsx_sanitize_cell(v)
        else:
            row[key] = xlsx_sanitize_cell(v)
    return row


def write_combined_iconik_metadata_csv(path: Path, rows: list[dict[str, str]]) -> bool:
    """UTF-8 CSV of merged Iconik-style flat fields (one row per input file)."""
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
    except OSError:
        return False
    if not rows:
        path.write_text("Filename\n", encoding="utf-8")
        return True
    keys: list[str] = []
    seen: set[str] = set()

    def add_key(k: str) -> None:
        if k not in seen:
            seen.add(k)
            keys.append(k)

    for r in rows:
        for k in r:
            add_key(k)
    order_first = [k for k in ("Filename", "Id", "resolved_asset_id", "resolved_title") if k in seen]
    rest = sorted(k for k in keys if k not in order_first)
    fieldnames = order_first + rest
    try:
        with path.open("w", encoding="utf-8", newline="") as f:
            w = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
            w.writeheader()
            for r in rows:
                clean = {k: xlsx_sanitize_cell(r.get(k, "")) for k in fieldnames}
                w.writerow(clean)
    except OSError:
        return False
    return True


def write_unmatched_mp4_report_xlsx(path: Path, rows: list[dict[str, str]]) -> bool:
    if not _OPENPYXL_AVAILABLE or openpyxl is None or Font is None or get_column_letter is None:
        return False
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "unmatched_mp4"
    headers = ("mp4_filename", "csv_path", "reason")
    ws.append(list(headers))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append([xlsx_sanitize_cell(r.get(h, "")) for h in headers])
    if not rows:
        ws.append(["—", "—", "(no MP4 files lacked a CSV row)"])
    nrow = max(len(rows), 1) + 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{nrow}"
    for i, h in enumerate(headers, start=1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = min(56, max(12, len(h) + 2))
    wb.save(str(path))
    return True


def write_unmatched_csv_rows_report_xlsx(path: Path, rows: list[dict[str, str]]) -> bool:
    if not _OPENPYXL_AVAILABLE or openpyxl is None or Font is None or get_column_letter is None:
        return False
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "unmatched_csv"
    headers = ("csv_row_index", "csv_path", "filename_cell", "title_cell", "reason")
    ws.append(list(headers))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append([xlsx_sanitize_cell(r.get(h, "")) for h in headers])
    if not rows:
        ws.append(["—", "—", "—", "—", "(no CSV rows lacked a matching MP4)"])
    nrow = max(len(rows), 1) + 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{nrow}"
    for i, h in enumerate(headers, start=1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = min(56, max(12, len(h) + 2))
    wb.save(str(path))
    return True


def write_csv_merge_debug_report_xlsx(path: Path, rows: list[dict[str, str]]) -> bool:
    if not _OPENPYXL_AVAILABLE or openpyxl is None or Font is None or get_column_letter is None:
        return False
    headers = (
        "csv_path",
        "derived_filename",
        "row_count",
        "non_blank_field_count",
        "merged_field_count",
        "missing_required_fields_after_merge",
        "fields_merged_from_multiple_rows",
        "fields_present_in_csv_but_missing_after_merge",
        "Goal_Time_Seconds_ENGP",
        "merge_warnings",
    )
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "csv_merge"
    ws.append(list(headers))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append([xlsx_sanitize_cell(r.get(h, "")) for h in headers])
    if not rows:
        ws.append(["(no CSV files merged in this run)", "", "", "", "", "", "", "", "", ""])
    nrow = max(len(rows), 1) + 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{nrow}"
    for i, h in enumerate(headers, start=1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = min(56, max(14, len(h) + 2))
    wb.save(str(path))
    return True


def build_delivery_validation_rows(
    mp4_paths: list[Path],
    inspired_rows: InspiredDataFrame,
    media_rows: list[dict[str, object]],
    warnings_by_filename: dict[str, list[str]],
    *,
    completion: str,
    ffmpeg_available: bool,
    csv_only_delivery: bool = False,
) -> list[dict[str, str]]:
    """One signoff row per input clip (same order as ``mp4_paths``)."""
    rows: list[dict[str, str]] = []
    n = len(mp4_paths)
    comp_u = str(completion).strip().upper()
    for i in range(n):
        mp4 = mp4_paths[i]
        fn = mp4.name
        row = inspired_rows[i] if i < len(inspired_rows) else {}
        mr = media_rows[i] if i < len(media_rows) else {}
        has_id = bool(str(row.get("Id", "")).strip())
        rud = any(
            str(row.get(c, "")).strip()
            for c in ("Country", "Season", "Match Date", "Home Team", "Away Team")
        )
        if csv_only_delivery:
            meta_found = bool(str(row.get("Filename", "")).strip()) and all(
                str(row.get(c, "")).strip()
                for c in ("Country", "Season", "Match Date", "Home Team", "Away Team")
            )
        else:
            meta_found = has_id or rud
        meta_tf = "TRUE" if meta_found else "FALSE"
        v_ok = str(mr.get("video_encode_ok", "")).strip().casefold() == "yes"
        c_ok = str(mr.get("com_wav_ok", "")).strip().casefold() == "yes"
        x_ok = str(mr.get("cfx_wav_ok", "")).strip().casefold() == "yes"
        dur_ok = bool(str(row.get("Duration (Frames)", "")).strip())
        warns = warnings_for_run_heuristic(warnings_by_filename.get(fn, []), filename=fn)
        warn_s = "; ".join(warns)

        if comp_u == "CANCELLED":
            clip_status = "CANCELLED"
        elif comp_u == "FAILED":
            clip_status = "FAILED"
        else:
            fail = (not meta_found) or (ffmpeg_available and not v_ok)
            if fail:
                clip_status = "FAIL"
            elif warns or (not c_ok) or (not x_ok) or (not dur_ok):
                clip_status = "WARN"
            else:
                clip_status = "OK"

        rows.append(
            {
                "clip_filename": fn,
                "metadata_row_found": meta_tf,
                "engp_video_generated": "TRUE" if v_ok else "FALSE",
                "com_wav_generated": "TRUE" if c_ok else "FALSE",
                "cfx_wav_generated": "TRUE" if x_ok else "FALSE",
                "duration_frames_populated": "TRUE" if dur_ok else "FALSE",
                "penalty_shootout": str(row.get("Penalty Shootout", "")).strip(),
                "reverse_goal": str(row.get("Reverse Goal", "")).strip(),
                "validation_warnings": warn_s,
                "final_delivery_status": clip_status,
            }
        )
    return rows


def write_metadata_qc_report_xlsx(path: Path, issues: list[dict[str, str]]) -> bool:
    if not _OPENPYXL_AVAILABLE or openpyxl is None or Font is None or get_column_letter is None:
        return False
    headers = ("filename", "rule_id", "severity", "field", "value", "message")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "metadata_qc"
    ws.append(list(headers))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for issue in issues:
        ws.append([xlsx_sanitize_cell(issue.get(h, "")) for h in headers])
    if not issues:
        ws.append(["—", "—", "—", "—", "—", "(no metadata QC issues)"])
    nrow = max(len(issues), 1) + 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{nrow}"
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(56, max(12, len(h) + 2))
    wb.save(str(path))
    return True


def write_delivery_validation_report_xlsx(path: Path, rows: list[dict[str, str]]) -> bool:
    if not _OPENPYXL_AVAILABLE or openpyxl is None or Font is None or get_column_letter is None:
        return False
    headers = (
        "clip_filename",
        "metadata_row_found",
        "engp_video_generated",
        "com_wav_generated",
        "cfx_wav_generated",
        "duration_frames_populated",
        "penalty_shootout",
        "reverse_goal",
        "validation_warnings",
        "final_delivery_status",
    )
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "validation"
    ws.append(list(headers))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append([xlsx_sanitize_cell(r.get(h, "")) for h in headers])
    if not rows:
        ws.append(["—", "FALSE", "FALSE", "FALSE", "FALSE", "FALSE", "", "", "(no clips)", "—"])
    nrow = max(len(rows), 1) + 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{nrow}"
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(56, max(12, len(h) + 2))
    wb.save(str(path))
    return True


def write_delivery_manifest_txt(
    path: Path,
    *,
    delivery_root: Path,
    completion: str,
    counters: ResolutionCounters,
    media_stats: dict[str, int],
    n_mp4: int,
    run_started_clock: datetime,
    run_duration_sec: float,
) -> bool:
    """Human-readable manifest: counts, timing, then every file with size and mtime."""
    lines: list[str] = [
        "Inspired Delivery Generator — delivery manifest",
        f"completion_state: {completion}",
        f"run_started_local: {run_started_clock.isoformat(timespec='seconds')}",
        f"run_duration_seconds: {run_duration_sec:.3f}",
        f"input_mp4_count: {n_mp4}",
        "",
        "Resolution counters:",
        f"  requested: {counters.requested}",
        f"  resolved: {counters.resolved}",
        f"  unresolved: {counters.unresolved}",
        f"  ambiguous: {counters.ambiguous}",
        f"  metadata_failures: {counters.metadata_failures}",
        f"  unmatched_mp4_vs_csv: {counters.unmatched_mp4}",
        f"  unmatched_csv_rows_vs_mp4: {counters.unmatched_csv}",
        "",
        "Media generation:",
        f"  video_ok: {media_stats.get('video_ok', 0)}",
        f"  video_fail: {media_stats.get('video_fail', 0)}",
        f"  com_ok: {media_stats.get('com_ok', 0)}",
        f"  com_fail: {media_stats.get('com_fail', 0)}",
        f"  cfx_ok: {media_stats.get('cfx_ok', 0)}",
        f"  cfx_fail: {media_stats.get('cfx_fail', 0)}",
        f"  cfx_missing_source: {media_stats.get('cfx_missing_source', 0)}",
        "",
        "Files (relative_path, size_bytes, mtime_local):",
    ]
    try:
        all_files = sorted(delivery_root.rglob("*"), key=lambda p: str(p).casefold())
    except OSError:
        all_files = []
    for fp in all_files:
        if not fp.is_file():
            continue
        if fp == path:
            continue
        try:
            rel = fp.relative_to(delivery_root)
        except ValueError:
            rel = fp
        try:
            st = fp.stat()
            sz = st.st_size
            mt = datetime.fromtimestamp(st.st_mtime).isoformat(timespec="seconds")
        except OSError:
            sz, mt = -1, ""
        lines.append(f"{rel}\t{sz}\t{mt}")
    lines.append("")
    lines.append(f"(manifest self: {path.name})")
    try:
        path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    except OSError:
        return False
    return True


def lock_delivery_outputs_tree(root: Path, emit: Callable[[str], None] | None = None) -> None:
    """Best-effort read-only: Windows sets files read-only; POSIX chmods files and directories."""
    em = emit or (lambda _s: None)
    marker = root / ".delivery_locked"
    iso = datetime.now().isoformat(timespec="seconds")
    try:
        marker.write_text(
            f"locked_at_local={iso}\n"
            "This delivery tree was sealed by Inspired Delivery Generator (read-only where supported).\n",
            encoding="utf-8",
        )
    except OSError as exc:
        em(f"[WARN] Could not write .delivery_locked marker: {exc}")
        return

    if os.name == "nt":
        for dirpath, _dirnames, filenames in os.walk(root):
            for name in filenames:
                fp = Path(dirpath) / name
                try:
                    os.chmod(fp, stat.S_IREAD)
                except OSError as exc:
                    em(f"[WARN] lock: could not set read-only on {fp}: {exc}")
        em(f"Delivery outputs locked under {root} (Windows: files set read-only).")
        return

    file_mode = 0o444
    dir_mode = 0o555
    for dirpath, dirnames, filenames in os.walk(root, topdown=False):
        for name in filenames:
            fp = Path(dirpath) / name
            try:
                os.chmod(fp, file_mode)
            except OSError as exc:
                em(f"[WARN] lock: could not chmod file {fp}: {exc}")
        for name in dirnames:
            dp = Path(dirpath) / name
            try:
                os.chmod(dp, dir_mode)
            except OSError as exc:
                em(f"[WARN] lock: could not chmod directory {dp}: {exc}")
    try:
        os.chmod(root, dir_mode)
    except OSError as exc:
        em(f"[WARN] lock: could not chmod delivery root {root}: {exc}")
    em(f"Delivery outputs locked under {root} (POSIX: read-only tree).")


def _scalar_to_string(value: object) -> str:
    if _is_blankish(value):
        return ""
    if isinstance(value, float):
        if math.isnan(value):
            return ""
        if value.is_integer():
            return str(int(value))
        return str(value).strip()
    if isinstance(value, int) and not isinstance(value, bool):
        return str(value)
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return _normalize_string_separators(str(value))


def _direction_of_play_raw_to_text(raw: object) -> str:
    """Flatten Iconik / CSV direction values to a single searchable string."""
    if raw is None or _is_blankish(raw):
        return ""
    if isinstance(raw, (list, tuple)):
        parts: list[str] = []
        for item in raw:
            t = _direction_of_play_raw_to_text(item)
            if t:
                parts.append(t)
        return " ".join(parts)
    return str(raw).strip()


def normalize_left_to_right_from_direction(raw: object) -> tuple[str, list[str]]:
    """
    Map Inspired ``LeftToRight`` from ``Direction_of_Play_ENGP`` only (not from legacy LeftToRight_* slugs).

    Left To Right / Right To Left substrings → TRUE / FALSE; blank → blank + missing warning;
    anything else (including ambiguous both) → blank + invalid warning. Never coerce blank to FALSE.
    """
    s = _direction_of_play_raw_to_text(raw)
    if not s:
        return "", ["Missing Direction_of_Play_ENGP"]
    scf = s.casefold()
    has_ltr = "left to right" in scf
    has_rtl = "right to left" in scf
    if has_ltr and not has_rtl:
        return "TRUE", []
    if has_rtl and not has_ltr:
        return "FALSE", []
    return "", ["Invalid Direction_of_Play_ENGP value"]


def is_non_goal_engp_clip(filename: str) -> bool:
    """True when the ENGP basename marks a non-goal event (``_O_<n>`` token, case-insensitive)."""
    return bool(_ENGP_NON_GOAL_EVENT_RE.search(Path(filename).name))


def engp_clip_requires_goal_time(filename: str) -> bool:
    """True when the ENGP basename marks a goal event (``_G_<n>`` token, not overridden by ``_O_<n>``)."""
    name = Path(filename).name
    if _ENGP_NON_GOAL_EVENT_RE.search(name):
        return False
    return bool(_ENGP_GOAL_EVENT_RE.search(name))


def _is_goal_time_validation_warning(message: str) -> bool:
    cf = message.casefold()
    return "goal_time_seconds_engp" in cf or "goal time (frames)" in cf


def warnings_for_run_heuristic(warnings: list[str], *, filename: str) -> list[str]:
    """Warnings that should affect preview, validation reports, and COMPLETED WITH WARNINGS."""
    if is_non_goal_engp_clip(filename):
        return [w for w in warnings if not _is_goal_time_validation_warning(w)]
    return list(warnings)


def normalize_goal_time_frames_from_seconds(
    raw: object,
    *,
    engp_filename: str = "",
) -> tuple[str, list[str]]:
    """
    Map Inspired ``Goal Time (Frames)`` from ``Goal_Time_Seconds_ENGP`` only: ``round(seconds * 25)``.
    Inspired delivery assumes 25 fps. Non-goal clips (``_O_<n>`` in filename) keep goal time blank with no warning.
    """
    if is_non_goal_engp_clip(engp_filename):
        return "", []

    if raw is None or _is_blankish(raw):
        if engp_clip_requires_goal_time(engp_filename):
            return "", ["missing Goal_Time_Seconds_ENGP for goal clip"]
        return "", []
    if isinstance(raw, bool):
        return "", ["invalid Goal_Time_Seconds_ENGP for goal clip"]
    if isinstance(raw, (list, tuple, dict)):
        return "", ["invalid Goal_Time_Seconds_ENGP for goal clip"]
    sec: float | None = None
    if isinstance(raw, (int, float)):
        if isinstance(raw, float) and math.isnan(raw):
            return "", ["invalid Goal_Time_Seconds_ENGP for goal clip"]
        sec = float(raw)
    else:
        s = str(raw).strip()
        if not s:
            if engp_clip_requires_goal_time(engp_filename):
                return "", ["missing Goal_Time_Seconds_ENGP for goal clip"]
            return "", []
        try:
            sec = float(s)
        except ValueError:
            return "", ["invalid Goal_Time_Seconds_ENGP for goal clip"]
    if sec is None or math.isnan(sec) or math.isinf(sec):
        return "", ["invalid Goal_Time_Seconds_ENGP for goal clip"]
    frames = int(round(sec * 25.0))
    return str(frames), []


# ---------------------------------------------------------------------------
# Team registry (shared) — per-delivery normalization state
# ---------------------------------------------------------------------------

_delivery_team_registry: TeamRegistryService | None = None
_unknown_team_records: list[dict[str, str]] = []


def begin_delivery_team_normalization() -> TeamRegistryService:
    """Load canonical team registry once per delivery run (before row build / export)."""
    global _delivery_team_registry, _unknown_team_records
    _delivery_team_registry = TeamRegistryService()
    _unknown_team_records = []
    return _delivery_team_registry


def get_delivery_team_registry() -> TeamRegistryService:
    if _delivery_team_registry is None:
        return begin_delivery_team_normalization()
    return _delivery_team_registry


def record_unknown_team(
    *,
    filename: str,
    column: str,
    original_value: str,
    detail: str = "",
) -> None:
    _unknown_team_records.append(
        {
            "filename": filename,
            "column": column,
            "original_value": original_value,
            "detail": detail,
        }
    )


def take_unknown_team_records() -> list[dict[str, str]]:
    return list(_unknown_team_records)


def write_unknown_teams_csv(path: Path, records: list[dict[str, str]]) -> bool:
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        with path.open("w", encoding="utf-8", newline="") as fh:
            writer = csv.DictWriter(
                fh,
                fieldnames=["filename", "column", "original_value", "detail"],
                lineterminator="\n",
            )
            writer.writeheader()
            for row in records:
                writer.writerow(
                    {
                        "filename": row.get("filename", ""),
                        "column": row.get("column", ""),
                        "original_value": row.get("original_value", ""),
                        "detail": row.get("detail", ""),
                    }
                )
        return True
    except OSError:
        return False


def normalize_inspired_cell(
    column: str,
    raw: object,
    *,
    engp_filename: str = "",
) -> tuple[str, list[str]]:
    """Normalize a single Inspired column value; returns (cell, warnings)."""
    warns: list[str] = []
    if column == "Filename":
        if raw is None:
            return "", warns
        return str(raw).strip(), warns

    if column == "LeftToRight":
        return normalize_left_to_right_from_direction(raw)

    if column in COERCED_ENGP_BOOL_COLUMNS:
        r = raw
        if isinstance(raw, (list, tuple)):
            parts = [p for p in (_cleanup_csv_cell(x) for x in raw) if p.strip()]
            r = ",".join(parts) if parts else None
        return _normalize_penalty_reverse_bool(r), []

    if column == "Goal Time (Frames)":
        return normalize_goal_time_frames_from_seconds(raw, engp_filename=engp_filename)

    if column == "Country":
        def walk_lists_country(v: object) -> str:
            if isinstance(v, (list, tuple)):
                parts: list[str] = []
                for item in v:
                    if isinstance(item, (list, tuple)):
                        inner = walk_lists_country(item)
                        if inner:
                            parts.append(inner)
                    else:
                        if _is_blankish(item):
                            continue
                        parts.append(_scalar_to_string(item))
                return _flatten_list_to_csv(parts)
            if _is_blankish(v):
                return ""
            return _scalar_to_string(v)

        if isinstance(raw, (list, tuple)):
            out = walk_lists_country(raw)
        elif _is_blankish(raw):
            return "", warns
        else:
            out = _scalar_to_string(raw)
        if out and re.search(r"english", out, re.IGNORECASE):
            return "ENG", warns
        return out, warns

    if column in INSPIRED_TEAM_COLUMNS:

        def walk_lists_team(v: object) -> str:
            if isinstance(v, (list, tuple)):
                parts: list[str] = []
                for item in v:
                    if isinstance(item, (list, tuple)):
                        inner = walk_lists_team(item)
                        if inner:
                            parts.append(inner)
                    else:
                        if _is_blankish(item):
                            continue
                        parts.append(_scalar_to_string(item))
                return _flatten_list_to_csv(parts)
            if _is_blankish(v):
                return ""
            return _scalar_to_string(v)

        if isinstance(raw, (list, tuple)):
            out = walk_lists_team(raw)
        elif _is_blankish(raw):
            return "", warns
        else:
            out = _scalar_to_string(raw)
        if not out:
            return "", warns
        norm = normalize_team_value_to_abbreviation(
            out,
            registry=get_delivery_team_registry(),
            allow_alias_enrichment=True,
        )
        if not norm.resolved:
            warns.append(norm.warning or f"unknown team: {out!r}")
            record_unknown_team(
                filename=engp_filename,
                column=column,
                original_value=out,
                detail=norm.warning or "",
            )
            return out, warns
        return norm.output_value, warns

    def walk_lists(v: object) -> str:
        if isinstance(v, (list, tuple)):
            parts: list[str] = []
            for item in v:
                if isinstance(item, (list, tuple)):
                    inner = walk_lists(item)
                    if inner:
                        parts.append(inner)
                else:
                    if _is_blankish(item):
                        continue
                    parts.append(_scalar_to_string(item))
            return _flatten_list_to_csv(parts)
        if _is_blankish(v):
            return ""
        return _scalar_to_string(v)

    if isinstance(raw, (list, tuple)):
        out = walk_lists(raw)
        if column in COERCED_ENGP_BOOL_COLUMNS:
            return _normalize_penalty_reverse_bool(out if out else None), []
        if column in BOOL_COLUMNS:
            warns.append(f"{column}: list value in boolean column; flattened before boolean check is invalid")
            return "", warns
        if column in DATE_COLUMNS:
            warns.append(f"{column}: list value in date column; cannot normalize to single date")
            return "", warns
        return out, warns

    if column in BOOL_COLUMNS:
        out, inv = _normalize_boolean(raw)
        if inv:
            if _is_blankish(raw):
                warns.append(f"{column}: blank boolean coerced to FALSE")
            else:
                warns.append(f"{column}: invalid boolean ({raw!r}) coerced to FALSE")
        return out, warns

    if column in DATE_COLUMNS:
        out, bad = _normalize_date_to_dd_mm_yyyy(raw)
        if bad:
            warns.append(f"{column}: malformed date ({raw!r})")
        return out, warns

    if isinstance(raw, (int, float)) and not isinstance(raw, bool):
        if column == "Duration (Frames)":
            if isinstance(raw, float) and math.isnan(raw):
                return "", warns
            if isinstance(raw, float) and raw.is_integer():
                return str(int(raw)), warns
            return str(raw), warns

    return walk_lists(raw), warns


def build_inspired_row(
    *,
    filename: str,
    resolved_title: str,
    asset_id: str,
    meta_doc: dict | None,
    metadata_ok: bool,
    flat: dict[str, object] | None = None,
    csv_only_delivery: bool = False,
) -> tuple[dict[str, str], list[str]]:
    """
    Build one Inspired dataframe row (exactly INSPIRED_COLUMNS) plus validation warnings.
    Does not expose arbitrary metadata keys — only mapped columns.

    If ``flat`` is provided (merged search-hit + optional endpoint flatten), it is the canonical
    metadata source for mapping; ``meta_doc`` is not re-flattened.
    """
    warns: list[str] = []
    if flat is not None:
        flat_use = dict(flat)
        if not metadata_ok and not csv_only_delivery:
            warns.append("metadata GET failed or skipped (search projection / merged flat used when present)")
        if not flat_use:
            warns.append("empty metadata: no extractable fields from search projection or metadata API")
    else:
        flat_use, _discovered_slugs = build_flat_metadata_map(meta_doc if metadata_ok and meta_doc else None)
        if metadata_ok and meta_doc and not flat_use:
            warns.append("empty metadata: no extractable fields")
        if not metadata_ok and not csv_only_delivery:
            warns.append("metadata request failed or skipped")

    cf_index = _casefold_first_key_index(flat_use)

    row: dict[str, str] = {c: "" for c in INSPIRED_COLUMNS}
    row["Filename"] = (filename or "").strip()
    aid = asset_id.strip()
    if not aid:
        aid = _id_value_from_flat(flat_use, cf_index=cf_index)
    row["Id"] = aid

    for col in INSPIRED_COLUMNS:
        if col in ("Filename", "Id"):
            continue
        raw = _get_mapped_raw(flat_use, col, cf_index=cf_index)
        cell, w = normalize_inspired_cell(col, raw, engp_filename=filename)
        row[col] = cell
        warns.extend(w)

    warns = warnings_for_run_heuristic(warns, filename=filename)

    for req in sorted(required_columns_for_delivery(csv_only_delivery=csv_only_delivery)):
        if row.get(req, "") == "":
            warns.append(f"missing required field: {req}")

    return row, warns


def inspired_norm_summary(warnings: list[str], *, filename: str = "") -> str:
    w = warnings_for_run_heuristic(warnings, filename=filename) if filename else warnings
    if not w:
        return "OK"
    return "WARN: " + "; ".join(w[:3]) + ("…" if len(w) > 3 else "")


def count_inspired_rows_with_heuristic_warnings(
    warnings_by_filename: dict[str, list[str]],
) -> int:
    """Row warning count for UI / COMPLETED WITH WARNINGS (excludes goal-time gaps on ``_O_`` clips)."""
    return sum(1 for fn, w in warnings_by_filename.items() if warnings_for_run_heuristic(w, filename=fn))


def build_normalized_preview_document(
    *,
    rows: InspiredDataFrame,
    warnings_by_filename: dict[str, list[str]],
    completion: str,
    duration_frames_populated: int = 0,
) -> dict[str, object]:
    total = len(rows)
    rows_with_warns = count_inspired_rows_with_heuristic_warnings(warnings_by_filename)
    return {
        "completion": completion,
        "columns": list(INSPIRED_COLUMNS),
        "rows": rows,
        "warnings_by_filename": {k: v for k, v in sorted(warnings_by_filename.items())},
        "stats": {
            "total_normalized_rows": total,
            "rows_with_warnings": rows_with_warns,
            "duration_frames_populated": duration_frames_populated,
        },
    }


# ---------------------------------------------------------------------------
# Phase 4 — FFprobe duration (frames) + Inspired XLSX
# ---------------------------------------------------------------------------

FFPROBE_CACHE: dict[str, tuple[str, str | None]] = {}
FFPROBE_TIMEOUT_SEC = 120


def parse_frame_rate(value: object) -> float | None:
    if value is None:
        return None
    t = str(value).strip()
    if not t or t.upper() in ("N/A", "0/0", "NAN"):
        return None
    if "/" in t:
        a, _, b = t.partition("/")
        try:
            aa, bb = float(a), float(b)
            if bb == 0:
                return None
            out = aa / bb
            return out if out > 0 else None
        except ValueError:
            return None
    try:
        out = float(t)
        return out if out > 0 else None
    except ValueError:
        return None


def ffprobe_duration_frames(
    path: Path,
    cache: dict[str, tuple[str, str | None]],
    cancel_event: threading.Event,
) -> tuple[str, str | None]:
    """
    Return (duration_frames_str, error_message_or_None).
    error_message is operator-facing: ffprobe missing, file missing, invalid fps, probe failed, …
    """
    try:
        key = str(path.resolve())
    except OSError:
        key = str(path)
    if key in cache:
        return cache[key]

    if cancel_event.is_set():
        return ("", "cancelled")

    if not shutil.which("ffprobe"):
        out = ("", "ffprobe missing (install ffmpeg / ffprobe and ensure it is on PATH)")
        cache[key] = out
        return out

    if not path.is_file():
        out = ("", "file missing")
        cache[key] = out
        return out

    cmd = [
        "ffprobe",
        "-v",
        "error",
        "-hide_banner",
        "-print_format",
        "json",
        "-show_format",
        "-show_streams",
        str(path),
    ]
    try:
        proc = subprocess.run(
            cmd,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=FFPROBE_TIMEOUT_SEC,
            check=False,
        )
    except (OSError, subprocess.SubprocessError) as e:
        out = ("", f"probe failed: {e}")
        cache[key] = out
        return out

    if proc.returncode != 0:
        err = (proc.stderr or proc.stdout or "").strip()[:500]
        out = ("", f"probe failed (exit {proc.returncode}): {err or 'no stderr'}")
        cache[key] = out
        return out

    try:
        doc = json.loads(proc.stdout or "{}")
    except json.JSONDecodeError as e:
        out = ("", f"probe failed: invalid JSON ({e})")
        cache[key] = out
        return out

    if not isinstance(doc, dict):
        out = ("", "probe failed: unexpected ffprobe output")
        cache[key] = out
        return out

    fmt = doc.get("format")
    if not isinstance(fmt, dict):
        out = ("", "probe failed: no format section")
        cache[key] = out
        return out

    dur_raw = fmt.get("duration")
    try:
        duration_s = float(dur_raw) if dur_raw is not None else None
    except (TypeError, ValueError):
        duration_s = None
    if duration_s is None or duration_s < 0 or math.isnan(duration_s):
        out = ("", "probe failed: invalid duration_seconds")
        cache[key] = out
        return out

    video: dict | None = None
    streams = doc.get("streams")
    if isinstance(streams, list):
        for s in streams:
            if isinstance(s, dict) and s.get("codec_type") == "video":
                video = s
                break

    fps: float | None = None
    if video:
        fps = parse_frame_rate(video.get("avg_frame_rate"))
        if fps is None:
            fps = parse_frame_rate(video.get("r_frame_rate"))

    if fps is None or fps <= 0 or math.isnan(fps):
        out = ("", "invalid fps (avg_frame_rate and r_frame_rate unusable)")
        cache[key] = out
        return out

    if cancel_event.is_set():
        return ("", "cancelled")

    frames = int(round(duration_s * fps))
    if frames < 0:
        out = ("", "probe failed: negative frame count")
        cache[key] = out
        return out

    out = (str(frames), None)
    cache[key] = out
    return out


# ---------------------------------------------------------------------------
# Phase 5 — Inspired media generation + packaging (FFmpeg / ffprobe validation)
# ---------------------------------------------------------------------------

MEDIA_FFMPEG_TIMEOUT_SEC = 6 * 3600
MEDIA_FFPROBE_VALIDATE_TIMEOUT_SEC = 180


def cfx_companion_path_for_engp(engp: Path) -> Path:
    """Legacy: ``NAME_ENGP.mp4`` → sibling ``NAME_CFX.mp4`` (case-insensitive on ``_ENGP.mp4`` suffix)."""
    return engp.parent / re.sub(r"(?i)_ENGP\.mp4$", "_CFX.mp4", engp.name)


def cfx_primary_engp_cfx_path(engp: Path) -> Path:
    """Canonical: ``NAME_ENGP.mp4`` → sibling ``NAME_ENGP_CFX.mp4`` (stem + ``_CFX.mp4``)."""
    return engp.parent / f"{engp.stem}_CFX.mp4"


def cfx_expected_primary_report_path(engp: Path) -> Path:
    """Path shown as canonical expected CFX MP4 beside this ENGP (primary naming rule)."""
    return cfx_primary_engp_cfx_path(engp)


def _cfx_resolved_file(p: Path) -> Path:
    try:
        return p.resolve()
    except OSError:
        return p


def _sibling_casefold_mp4_match(parent: Path, basename_folds: tuple[str, ...]) -> tuple[Path | None, list[Path]]:
    """First basename_fold in order that matches any sibling ``*.mp4`` (case-insensitive basename)."""
    try:
        entries = list(parent.iterdir())
    except OSError:
        return None, []
    mp4s = [p for p in entries if p.is_file() and p.suffix.lower() == ".mp4"]
    for fold in basename_folds:
        hits = sorted([p for p in mp4s if p.name.casefold() == fold], key=lambda p: str(p))
        if hits:
            return hits[0], hits[1:]
    return None, []


def _sorted_pick(paths: list[Path]) -> tuple[Path | None, list[Path]]:
    if not paths:
        return None, []
    s = sorted(paths, key=lambda p: str(p))
    return s[0], s[1:]


@dataclass
class CfxTreeIndex:
    """Recursive CFX tree: exact basename, casefold basename, and flat MP4 list."""

    by_basename: dict[str, list[Path]] = field(default_factory=dict)
    by_basename_fold: dict[str, list[Path]] = field(default_factory=dict)
    all_mp4_paths: list[Path] = field(default_factory=list)


@dataclass
class CfxIndexBuildStats:
    expanded_root: str = ""
    dirs_scanned: int = 0
    mp4_indexed: int = 0
    rescan_attempted: bool = False


def normalize_cfx_basename(name: str) -> str:
    """Strip whitespace, line breaks, and common invisible Unicode from MP4 basenames."""
    s = str(name or "")
    s = s.replace("\r", "").replace("\n", "").replace("\t", "")
    for ch in _CFX_INVISIBLE_CHARS:
        s = s.replace(ch, "")
    return s.strip()


def _cfx_engp_event_stem(engp: Path) -> str:
    """Match stem without trailing ``_ENGP`` (e.g. ``ARS_BHA_20130126_G_01``)."""
    return re.sub(r"(?i)_engp$", "", engp.stem)


def _cfx_walk_onerror(exc: OSError) -> None:
    """Allow ``os.walk`` to continue on FUSE / CloudMounter transient errors."""
    return


def _scan_cfx_tree_into_index(base_str: str, out: CfxTreeIndex) -> int:
    """Walk ``base_str``; return number of directories visited."""
    dirs_seen = 0
    for dirpath, dirnames, filenames in os.walk(base_str, topdown=True, onerror=_cfx_walk_onerror):
        dirs_seen += 1
        dirnames.sort()
        for fn in filenames:
            if not fn.lower().endswith(".mp4"):
                continue
            full = Path(dirpath) / fn
            try:
                full_str = os.path.join(dirpath, os.path.basename(str(full)))
            except (TypeError, ValueError):
                full_str = str(full)
            full_path = Path(full_str)
            bn = normalize_cfx_basename(os.path.basename(full_str))
            if not bn.lower().endswith(".mp4"):
                continue
            out.all_mp4_paths.append(full_path)
            out.by_basename.setdefault(bn, []).append(full_path)
            out.by_basename_fold.setdefault(bn.casefold(), []).append(full_path)
    return dirs_seen


def _finalize_cfx_index(out: CfxTreeIndex) -> None:
    for k, paths in out.by_basename.items():
        out.by_basename[k] = sorted(paths, key=lambda p: str(p))
    for k, paths in out.by_basename_fold.items():
        out.by_basename_fold[k] = sorted(paths, key=lambda p: str(p))
    out.all_mp4_paths.sort(key=lambda p: str(p))


def build_cfx_basename_index(root: Path | str) -> tuple[CfxTreeIndex, CfxIndexBuildStats]:
    """
    Walk ``root`` (no ``Path.resolve`` on root). Index every ``*.mp4`` by normalized basename
    and by ``basename.casefold()``. Retries once after a short delay when CloudMounter returns
    an empty listing on first pass.
    """
    stats = CfxIndexBuildStats()
    out = CfxTreeIndex()
    base_str = os.path.expanduser(str(root).strip())
    stats.expanded_root = base_str
    if not cfx_source_root_usable_for_scan(base_str):
        return out, stats
    try:
        stats.dirs_scanned = _scan_cfx_tree_into_index(base_str, out)
        stats.mp4_indexed = len(out.all_mp4_paths)
        if stats.mp4_indexed == 0:
            info = inspect_cfx_source_root_path(base_str)
            if info.get("exists") and info.get("readable"):
                time.sleep(_CFX_CLOUDMOUNTER_RESCAN_DELAY_SEC)
                stats.rescan_attempted = True
                retry = CfxTreeIndex()
                stats.dirs_scanned = _scan_cfx_tree_into_index(base_str, retry)
                if len(retry.all_mp4_paths) > 0:
                    out = retry
                    stats.mp4_indexed = len(out.all_mp4_paths)
        _finalize_cfx_index(out)
    except OSError:
        return CfxTreeIndex(), stats
    return out, stats


def nearest_cfx_candidates_by_contains(engp: Path, tree: CfxTreeIndex) -> list[Path]:
    """Basenames under the CFX tree that contain the ENGP event stem and end with CFX naming."""
    stem = _cfx_engp_event_stem(engp)
    if not stem:
        return []
    stem_fold = stem.casefold()
    hits: list[Path] = []
    seen: set[str] = set()
    for p in tree.all_mp4_paths:
        bn = normalize_cfx_basename(os.path.basename(str(p)))
        bn_cf = bn.casefold()
        if stem_fold not in bn_cf:
            continue
        if not (bn_cf.endswith("_engp_cfx.mp4") or bn_cf.endswith("_cfx.mp4")):
            continue
        key = str(p)
        if key in seen:
            continue
        seen.add(key)
        hits.append(p)
    return sorted(hits, key=lambda x: str(x))


def write_cfx_index_debug_txt(path: Path, tree: CfxTreeIndex, header_lines: list[str]) -> bool:
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        lines = list(header_lines)
        lines.append("")
        lines.append("indexed_mp4_files (basename TAB full_path):")
        for p in tree.all_mp4_paths:
            bn = normalize_cfx_basename(os.path.basename(str(p)))
            lines.append(f"{bn}\t{p}")
        if not tree.all_mp4_paths:
            lines.append("(no MP4 files indexed)")
        path.write_text("\n".join(lines) + "\n", encoding="utf-8")
        return True
    except OSError:
        return False


def emit_cfx_index_prebuild_diagnostics(
    emit: Callable[[str], None],
    *,
    raw: str,
    cleaned: str,
    stats: CfxIndexBuildStats,
) -> None:
    info = inspect_cfx_source_root_path(cleaned)
    emit("CFX index: pre-build diagnostics")
    emit(f"  raw configured CFX root: {raw!r}")
    emit(f"  normalized/expanded CFX root: {stats.expanded_root!r}")
    emit(f"  exists: {info.get('exists')!r}")
    emit(f"  is_dir (os.path.isdir): {info.get('is_dir_os')!r}")
    emit(f"  is_dir (Path.is_dir): {info.get('is_dir_path')!r}")
    emit(f"  readable: {info.get('readable')!r}")
    emit(f"  parent_exists: {info.get('parent_exists')!r}")
    emit(f"  root usable for scan: {cfx_source_root_usable_for_scan(cleaned)!r}")


def emit_cfx_index_postbuild_diagnostics(
    emit: Callable[[str], None],
    *,
    stats: CfxIndexBuildStats,
    tree: CfxTreeIndex,
) -> None:
    emit("CFX index: post-build summary")
    emit(f"  total recursive directories scanned: {stats.dirs_scanned}")
    emit(f"  total MP4 files indexed: {stats.mp4_indexed}")
    if stats.rescan_attempted:
        emit(f"  CloudMounter re-scan attempted after {_CFX_CLOUDMOUNTER_RESCAN_DELAY_SEC}s delay")


def emit_cfx_lookup_diagnostics(
    emit: Callable[[str], None],
    *,
    engp: Path,
    tree: CfxTreeIndex,
    primary: Path,
    legacy: Path,
    result: Path | None,
    discovery: str,
    duplicates: list[Path],
) -> None:
    prim_bn = normalize_cfx_basename(primary.name)
    leg_bn = normalize_cfx_basename(legacy.name)
    exact_prim = tree.by_basename.get(prim_bn, [])
    exact_leg = tree.by_basename.get(leg_bn, [])
    fold_prim = tree.by_basename_fold.get(prim_bn.casefold(), [])
    fold_leg = tree.by_basename_fold.get(leg_bn.casefold(), [])
    near = nearest_cfx_candidates_by_contains(engp, tree)
    emit(f"CFX lookup: requested ENGP filename: {engp.name!r}")
    emit(f"  expected ENGP_CFX basename: {prim_bn!r}")
    emit(f"  expected legacy CFX basename: {leg_bn!r}")
    emit(f"  contains stem: {_cfx_engp_event_stem(engp)!r}")
    emit(
        f"  exact basename lookup (primary): {len(exact_prim)} hit(s)"
        + (f" → {exact_prim[0]}" if exact_prim else "")
    )
    emit(
        f"  exact basename lookup (legacy): {len(exact_leg)} hit(s)"
        + (f" → {exact_leg[0]}" if exact_leg else "")
    )
    emit(
        f"  casefold basename lookup (primary): {len(fold_prim)} hit(s)"
        + (f" → {fold_prim[0]}" if fold_prim else "")
    )
    emit(
        f"  casefold basename lookup (legacy): {len(fold_leg)} hit(s)"
        + (f" → {fold_leg[0]}" if fold_leg else "")
    )
    emit(f"  nearest_candidates_by_contains: {len(near)}")
    for i, p in enumerate(near[:12]):
        emit(f"    [{i + 1}] {normalize_cfx_basename(os.path.basename(str(p)))!r} → {p}")
    if len(near) > 12:
        emit(f"    … (+{len(near) - 12} more)")
    if result is not None:
        emit(f"  resolved: {discovery!r} → {result}")
        if duplicates:
            emit(f"  duplicate candidates ({len(duplicates)}): " + "; ".join(str(d) for d in duplicates[:8]))
    else:
        emit(f"  resolved: (none) discovery={discovery!r}")


def resolve_cfx_source_mp4(
    engp: Path,
    tree: CfxTreeIndex,
    *,
    emit: Callable[[str], None] | None = None,
) -> tuple[Path | None, Path, str, list[Path]]:
    """
    Resolve CFX-side source MP4 for this ENGP.

    Primary: ``NAME_ENGP_CFX.mp4`` (sibling then recursive under CFX root, exact / casefold).
    Fallback: ``NAME_CFX.mp4`` with the same lookup order. If still missing, exactly one
    ``nearest_candidates_by_contains`` match may be used (stem contains + CFX suffix).
    """
    primary = cfx_primary_engp_cfx_path(engp)
    legacy = cfx_companion_path_for_engp(engp)
    expect = cfx_expected_primary_report_path(engp)
    prim_bn = normalize_cfx_basename(primary.name)
    leg_bn = normalize_cfx_basename(legacy.name)

    def ret(p: Path, disc: str, dups: list[Path]) -> tuple[Path | None, Path, str, list[Path]]:
        resolved = _cfx_resolved_file(p)
        if emit is not None:
            emit_cfx_lookup_diagnostics(
                emit,
                engp=engp,
                tree=tree,
                primary=primary,
                legacy=legacy,
                result=resolved,
                discovery=disc,
                duplicates=dups,
            )
        return resolved, expect, disc, dups

    def finish_none(disc: str, dups: list[Path] | None = None) -> tuple[Path | None, Path, str, list[Path]]:
        if emit is not None:
            emit_cfx_lookup_diagnostics(
                emit,
                engp=engp,
                tree=tree,
                primary=primary,
                legacy=legacy,
                result=None,
                discovery=disc,
                duplicates=list(dups or []),
            )
        return None, expect, disc, list(dups or [])

    if primary.is_file():
        return ret(primary, "sibling_engp_cfx", [])
    if legacy.is_file():
        return ret(legacy, "sibling_legacy_cfx", [])

    folds = (prim_bn.casefold(), leg_bn.casefold())
    hit, dups = _sibling_casefold_mp4_match(engp.parent, folds)
    if hit is not None:
        return ret(hit, "sibling_casefold", dups)

    chosen, rest = _sorted_pick(tree.by_basename.get(prim_bn, []))
    if chosen is not None:
        return ret(chosen, "cfx_root_engp_cfx", rest)

    chosen, rest = _sorted_pick(tree.by_basename.get(leg_bn, []))
    if chosen is not None:
        return ret(chosen, "cfx_root_legacy_cfx", rest)

    chosen, rest = _sorted_pick(tree.by_basename_fold.get(prim_bn.casefold(), []))
    if chosen is not None:
        return ret(chosen, "cfx_root_casefold_engp_cfx", rest)

    chosen, rest = _sorted_pick(tree.by_basename_fold.get(leg_bn.casefold(), []))
    if chosen is not None:
        return ret(chosen, "cfx_root_casefold_legacy_cfx", rest)

    near = nearest_cfx_candidates_by_contains(engp, tree)
    if len(near) == 1:
        if emit is not None:
            emit(f"[WARN] {engp.name}: CFX matched by nearest stem fallback → {near[0]}")
        return ret(near[0], "nearest_stem_fallback", [])
    if len(near) > 1:
        return finish_none("ambiguous_nearest_contains", near)

    return finish_none("missing", [])


def inspect_cfx_source_root_path(cleaned: str) -> dict[str, object]:
    """
    Filesystem probes for diagnostics (no Path.resolve — CloudMounter / FUSE paths
    may not resolve reliably).
    """
    p = os.path.expanduser(cleaned.strip())
    path_obj = Path(p)
    parent = path_obj.parent
    exists_p = False
    ok_os_dir = False
    ok_path_dir = False
    readable = False
    parent_exists = False
    try:
        exists_p = os.path.exists(p)
    except OSError:
        exists_p = False
    try:
        ok_os_dir = os.path.isdir(p)
    except OSError:
        ok_os_dir = False
    try:
        ok_path_dir = path_obj.is_dir()
    except OSError:
        ok_path_dir = False
    try:
        readable = os.access(p, os.R_OK)
    except OSError:
        readable = False
    try:
        parent_exists = parent.exists()
    except OSError:
        parent_exists = False
    return {
        "expanded_path": p,
        "exists": exists_p,
        "is_dir_os": ok_os_dir,
        "is_dir_path": ok_path_dir,
        "readable": readable,
        "parent_exists": parent_exists,
    }


def cfx_source_root_usable_for_scan(cleaned: str) -> bool:
    """
    True when the CFX root can be walked. CloudMounter / FUSE mounts may fail one of the two
    directory probes; accept if either reports a directory, or the path exists and is readable.
    """
    info = inspect_cfx_source_root_path(cleaned)
    if not info.get("exists"):
        return False
    if bool(info.get("is_dir_os")) or bool(info.get("is_dir_path")):
        return True
    return bool(info.get("readable"))


def emit_cfx_root_rejection_diagnostics(
    emit: Callable[[str], None],
    *,
    raw: str,
    cleaned: str,
) -> None:
    info = inspect_cfx_source_root_path(cleaned)
    emit("[WARN] CFX source root rejected (tree search skipped).")
    emit(f"  RAW CFX ROOT: {raw!r}")
    emit(f"  CLEANED CFX ROOT: {cleaned!r}")
    emit(f"  expanded_path: {info['expanded_path']!r}")
    emit(f"  exists: {info['exists']!r}")
    emit(f"  is_dir (os.path.isdir): {info['is_dir_os']!r}")
    emit(f"  is_dir (Path.is_dir): {info['is_dir_path']!r}")
    emit(f"  readable: {info['readable']!r}")
    emit(f"  parent_exists: {info['parent_exists']!r}")


def media_packaged_output_paths(engp: Path, media_dir: Path) -> tuple[Path, Path, Path]:
    """Return (output ENGP mp4, COM wav, CFX-side wav) under ``media_dir``; basenames match Inspired naming."""
    stem = engp.stem
    return (
        media_dir / engp.name,
        media_dir / f"{stem}_COM.wav",
        media_dir / f"{stem}_CFX.wav",
    )


def _ffmpeg_available() -> bool:
    return bool(shutil.which("ffmpeg"))


def _ffprobe_available() -> bool:
    return bool(shutil.which("ffprobe"))


def count_engp_cfx_mp4s_under_root(root: Path | str | None) -> int:
    """Count ``*.mp4`` files under ``root`` whose basename ends with ``_ENGP_CFX.mp4`` (case-insensitive)."""
    if root is None:
        return 0
    cleaned = clean_cfx_source_root_path(str(root))
    if not cleaned or not cfx_source_root_usable_for_scan(cleaned):
        return 0
    base = os.path.expanduser(cleaned.strip())
    suf = "_engp_cfx.mp4"
    n = 0
    try:
        for _d, _subs, files in os.walk(base, topdown=True, onerror=_cfx_walk_onerror):
            for fn in files:
                bn = normalize_cfx_basename(os.path.basename(fn))
                if bn.casefold().endswith(suf):
                    n += 1
    except OSError:
        return 0
    return n


def _ffmpeg_run(
    cmd: list[str],
    *,
    timeout_sec: int,
    cancel_event: threading.Event,
) -> tuple[int, str]:
    if cancel_event.is_set():
        return -1, "cancelled"
    try:
        proc = subprocess.run(
            cmd,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=timeout_sec,
            check=False,
        )
    except subprocess.TimeoutExpired:
        return -124, "ffmpeg timeout"
    except (OSError, subprocess.SubprocessError) as e:
        return -1, str(e)
    tail = (proc.stderr or proc.stdout or "").strip()
    if len(tail) > 4000:
        tail = tail[-4000:]
    return proc.returncode, tail


def ffmpeg_transcode_inspired_engp_video(
    engp_source: Path,
    out_mp4: Path,
    cancel_event: threading.Event,
) -> tuple[bool, str]:
    """
    Deterministic H.264 1920×1080 @ 25fps, BT.709, ~8.5 Mbps VBR, video-only MP4 (Inspired spec).
    """
    if not _ffmpeg_available():
        return False, "ffmpeg not on PATH"
    out_mp4.parent.mkdir(parents=True, exist_ok=True)
    vf = "scale=1920:1080,setsar=1"
    cmd = [
        "ffmpeg",
        "-hide_banner",
        "-nostats",
        "-y",
        "-i",
        str(engp_source),
        "-an",
        "-map_metadata",
        "-1",
        "-c:v",
        "libx264",
        "-profile:v",
        "high",
        "-pix_fmt",
        "yuv420p",
        "-r",
        "25",
        "-vf",
        vf,
        "-color_primaries",
        "bt709",
        "-color_trc",
        "bt709",
        "-colorspace",
        "bt709",
        "-b:v",
        "8500k",
        "-maxrate",
        "9000k",
        "-bufsize",
        "18000k",
        "-preset",
        "medium",
        "-movflags",
        "+faststart",
        str(out_mp4),
    ]
    rc, err = _ffmpeg_run(cmd, timeout_sec=MEDIA_FFMPEG_TIMEOUT_SEC, cancel_event=cancel_event)
    if cancel_event.is_set():
        return False, "cancelled"
    if rc != 0:
        return False, f"ffmpeg exit {rc}: {err}"
    return True, ""


def ffmpeg_extract_inspired_wav(
    source_media: Path,
    out_wav: Path,
    cancel_event: threading.Event,
) -> tuple[bool, str]:
    """48 kHz mono PCM s16le WAV (Inspired spec)."""
    if not _ffmpeg_available():
        return False, "ffmpeg not on PATH"
    out_wav.parent.mkdir(parents=True, exist_ok=True)
    cmd = [
        "ffmpeg",
        "-hide_banner",
        "-nostats",
        "-y",
        "-i",
        str(source_media),
        "-vn",
        "-map_metadata",
        "-1",
        "-c:a",
        "pcm_s16le",
        "-ar",
        "48000",
        "-ac",
        "1",
        "-sample_fmt",
        "s16",
        str(out_wav),
    ]
    rc, err = _ffmpeg_run(cmd, timeout_sec=MEDIA_FFMPEG_TIMEOUT_SEC, cancel_event=cancel_event)
    if cancel_event.is_set():
        return False, "cancelled"
    if rc != 0:
        return False, f"ffmpeg exit {rc}: {err}"
    return True, ""


def _ffprobe_json_file(path: Path) -> dict | None:
    if not shutil.which("ffprobe") or not path.is_file():
        return None
    cmd = [
        "ffprobe",
        "-v",
        "error",
        "-hide_banner",
        "-print_format",
        "json",
        "-show_format",
        "-show_streams",
        str(path),
    ]
    try:
        proc = subprocess.run(
            cmd,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=MEDIA_FFPROBE_VALIDATE_TIMEOUT_SEC,
            check=False,
        )
    except (OSError, subprocess.SubprocessError):
        return None
    if proc.returncode != 0:
        return None
    try:
        doc = json.loads(proc.stdout or "{}")
    except json.JSONDecodeError:
        return None
    return doc if isinstance(doc, dict) else None


def validate_packaged_mp4_inspired(path: Path) -> dict[str, str]:
    """Post-encode checks for packaged Inspired MP4 (delegates to shared video validator)."""
    from services.inspired_video_validator import validate_inspired_mp4

    return validate_inspired_mp4(path)


def validate_packaged_wav_inspired(path: Path) -> dict[str, str]:
    """Post-encode Inspired WAV checks (delegates to shared audio validator)."""
    from services.inspired_audio_validator import validate_inspired_wav

    return validate_inspired_wav(path)


def write_media_generation_report_xlsx(path: Path, rows: list[dict[str, object]]) -> bool:
    if not _OPENPYXL_AVAILABLE or openpyxl is None or Font is None or get_column_letter is None:
        return False
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "media_generation"
    if rows:
        headers = list(rows[0].keys())
    else:
        headers = [
            "engp_filename",
            "engp_source_path",
            "cfx_expected_path",
            "cfx_resolved_source_path",
            "cfx_discovery",
            "cfx_duplicate_paths",
            "cfx_found",
            "output_video",
            "output_com_wav",
            "output_cfx_wav",
            "video_encode_ok",
            "video_encode_error",
            "com_wav_ok",
            "com_wav_error",
            "cfx_wav_ok",
            "cfx_wav_error",
        ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append([xlsx_sanitize_cell(r.get(h, "")) for h in headers])
    if not rows:
        ws.append([xlsx_sanitize_cell("(no media generation rows)") if i == 0 else "" for i in range(len(headers))])
    nrow = max(len(rows), 1) + 1
    ncol = len(headers)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ncol)}{nrow}"
    for i in range(1, ncol + 1):
        letter = get_column_letter(i)
        hname = headers[i - 1]
        ws.column_dimensions[letter].width = min(56, max(10, len(str(hname)) + 2))
    wb.save(str(path))
    return True


def write_unresolved_media_report_xlsx(path: Path, rows: list[dict[str, object]]) -> bool:
    if not _OPENPYXL_AVAILABLE or openpyxl is None or Font is None or get_column_letter is None:
        return False
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "unresolved_media"
    headers = (
        "engp_filename",
        "cfx_expected_path",
        "cfx_resolved_source_path",
        "cfx_discovery",
        "cfx_duplicate_paths",
        "issue",
        "video_error",
        "com_wav_error",
        "cfx_wav_error",
    )
    ws.append(list(headers))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append([xlsx_sanitize_cell(r.get(h, "")) for h in headers])
    if not rows:
        ws.append(["—", "—", "—", "—", "—", "(no unresolved media issues)", "—", "—", "—"])
    nrow = max(len(rows), 1) + 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{nrow}"
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(56, max(10, len(h) + 2))
    wb.save(str(path))
    return True


def xlsx_sanitize_cell(value: object) -> str:
    """Blank-safe, NaN-safe string for XLSX cells."""
    if value is None:
        return ""
    if isinstance(value, float) and math.isnan(value):
        return ""
    if isinstance(value, str):
        t = value.strip()
        if t.casefold() in ("nan", "none", "null"):
            return ""
        return t
    return str(value)


def write_inspired_metadata_xlsx(
    path: Path,
    rows: InspiredDataFrame,
    *,
    warnings_by_filename: dict[str, list[str]] | None = None,
) -> bool:
    """Inspired columns only, exact order, bold header, freeze, autofilter, widths, red row fill on warnings."""
    if not _OPENPYXL_AVAILABLE or openpyxl is None or Font is None or get_column_letter is None or PatternFill is None:
        return False
    warn_fill = PatternFill(fill_type="solid", fgColor="FFFFC7CE")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inspired"
    headers = list(INSPIRED_COLUMNS)
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    warns_in = warnings_by_filename or {}
    for row in rows:
        ws.append([xlsx_sanitize_cell(row.get(c, "")) for c in INSPIRED_COLUMNS])
    ws.freeze_panes = "A2"
    nrow = len(rows) + 1
    ncol = len(INSPIRED_COLUMNS)
    ws.auto_filter.ref = f"A1:{get_column_letter(ncol)}{nrow}"
    col_widths: dict[str, float] = {
        "Filename": 44,
        "Id": 38,
        "Country": 14,
        "Season": 10,
        "Match Date": 14,
        "Rights Holders": 22,
        "Home Team": 18,
        "Away Team": 18,
        "Attacking Team": 18,
        "Defending Team": 18,
        "LeftToRight": 14,
        "Action": 24,
        "Goal Scored": 14,
        "Penalty Shootout": 16,
        "Reverse Goal": 14,
        "Goal Type": 16,
        "Goal Time (Frames)": 18,
        "Duration (Frames)": 18,
    }
    for i, name in enumerate(headers, start=1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = col_widths.get(name, 14)
    for ri, row in enumerate(rows, start=2):
        fn = row.get("Filename", "").strip()
        if warns_in.get(fn):
            for ci in range(1, ncol + 1):
                ws.cell(row=ri, column=ci).fill = warn_fill
    wb.save(str(path))
    return True


@dataclass
class ResolutionCounters:
    requested: int = 0
    resolved: int = 0
    unresolved: int = 0
    ambiguous: int = 0
    metadata_failures: int = 0
    unmatched_mp4: int = 0
    unmatched_csv: int = 0

    def to_dict(self) -> dict[str, int]:
        return {
            "requested": self.requested,
            "resolved": self.resolved,
            "unresolved": self.unresolved,
            "ambiguous": self.ambiguous,
            "metadata_failures": self.metadata_failures,
            "unmatched_mp4": self.unmatched_mp4,
            "unmatched_csv": self.unmatched_csv,
        }


class DeliveryWorker(threading.Thread):
    """Resolve Iconik assets + metadata; write delivery outputs; thread-safe queue updates."""

    def __init__(
        self,
        *,
        mp4_paths: list[Path],
        output_parent: Path,
        message_queue: queue.Queue,
        cancel_event: threading.Event,
        iconik_base_url: str,
        iconik_app_id: str,
        iconik_auth_token: str,
        iconik_collection_uuid: str,
        cfx_source_root_folder: str,
        cfx_source_root_folder_raw: str = "",
        metadata_csv_paths: Sequence[Path | str] | None = None,
        lock_delivery_outputs: bool = False,
        csv_only_delivery: bool = False,
    ) -> None:
        super().__init__(daemon=True)
        self._mp4_paths = mp4_paths
        self._output_parent = output_parent
        self._q = message_queue
        self._cancel = cancel_event
        self._iconik_base_url = iconik_base_url.strip()
        self._iconik_app_id = iconik_app_id.strip()
        self._iconik_auth_token = iconik_auth_token.strip()
        self._iconik_collection_uuid = iconik_collection_uuid.strip()
        self._cfx_source_root_folder = clean_cfx_source_root_path(cfx_source_root_folder)
        self._cfx_source_root_raw = cfx_source_root_folder_raw
        self._metadata_csv_paths: list[Path] = [Path(os.path.expanduser(str(p).strip())) for p in (metadata_csv_paths or ()) if str(p).strip()]
        self._lock_delivery_outputs = bool(lock_delivery_outputs)
        self._csv_only_delivery = bool(csv_only_delivery)

    def _send(self, payload: dict) -> None:
        self._q.put(payload)

    def _counters(self, c: ResolutionCounters) -> None:
        self._send({"type": "counters", **c.to_dict()})

    def _emit_status(self, delivery_log: list[str] | None, text: str) -> None:
        if delivery_log is not None:
            ts = datetime.now().isoformat(timespec="seconds")
            delivery_log.append(f"{ts}\t{text}")
        self._send({"type": "status_line", "text": text})

    def _run_frame_alignment_qc(
        self,
        delivery_dir: Path,
        delivery_log: list[str],
        media_rows: list[dict[str, object]],
    ) -> bool:
        """Post-render QC: MP4 frame count vs COM/CFX WAV frame-equivalents (±1 frame @ 25 fps)."""
        from services.frame_alignment_validator import (
            validate_delivery_frame_alignment,
            write_frame_alignment_qc_report,
        )

        media_dir = delivery_dir / "media"
        self._emit_status(
            delivery_log,
            "Frame alignment QC: validating rendered MP4 vs COM/CFX WAV (ffprobe)…",
        )
        result = validate_delivery_frame_alignment(self._mp4_paths, media_dir, media_rows)
        report_path = delivery_dir / "reports" / "frame_alignment_qc_report.txt"
        if write_frame_alignment_qc_report(report_path, result):
            self._emit_status(delivery_log, "Wrote reports/frame_alignment_qc_report.txt.")
        else:
            self._emit_status(delivery_log, "[WARN] Could not write reports/frame_alignment_qc_report.txt.")
        self._emit_status(delivery_log, result.summary)
        for block in result.log_blocks:
            self._emit_status(delivery_log, block)
        return result.passed

    def _run_metadata_goal_qc(
        self,
        delivery_log: list[str],
        inspired_rows: InspiredDataFrame,
        warnings_by_filename: dict[str, list[str]],
    ):
        from services.inspired_metadata_qc import validate_inspired_metadata_batch

        self._emit_status(delivery_log, "Metadata QC: validating goal/filename/action rules…")
        result = validate_inspired_metadata_batch(self._mp4_paths, inspired_rows)
        for issue in result.issues:
            warnings_by_filename.setdefault(issue.filename, []).append(issue.message)
            tag = "ERROR" if issue.severity == "error" else "WARN"
            self._emit_status(delivery_log, f"[{tag}] {issue.filename}: {issue.message}")
        self._emit_status(
            delivery_log,
            f"Metadata QC complete: {result.error_count} error(s), {result.warning_count} warning(s).",
        )
        return result

    def _run_video_spec_qc(
        self,
        delivery_log: list[str],
        media_rows: list[dict[str, object]],
    ) -> bool:
        from services.inspired_video_validator import (
            validate_delivery_mp4_fps_consistency,
            validate_inspired_mp4,
        )

        self._emit_status(
            delivery_log,
            "Video spec QC: validating MP4 (H264, no audio, 25/50 fps, consistent delivery frame rate)…",
        )
        passed = True
        for row in media_rows:
            if str(row.get("video_encode_ok", "")).strip().casefold() != "yes":
                continue
            fn = str(row.get("engp_filename", ""))
            mp4_path = Path(str(row.get("output_video", "")))
            if not mp4_path.is_file():
                passed = False
                msg = f"Delivery failed: packaged MP4 missing for {fn}"
                self._emit_status(delivery_log, f"[ERROR] {msg}")
                continue
            val = validate_inspired_mp4(mp4_path)
            row["mp4_container"] = val.get("container", "")
            row["mp4_codec"] = val.get("codec", "")
            row["mp4_fps"] = val.get("fps", "")
            row["mp4_audio_stream_count"] = val.get("audio_stream_count", "")
            row["mp4_validation_ok"] = val.get("validation_ok", "")
            row["mp4_validation_summary"] = val.get("validation_notes", "")
            if str(val.get("validation_ok", "")).lower() != "yes":
                passed = False
                reason = val.get("failure_reason") or val.get("validation_notes") or "video spec failed"
                self._emit_status(delivery_log, f"[ERROR] {fn}: Delivery failed: MP4 — {reason}")

        fps_ok, fps_msg, _ = validate_delivery_mp4_fps_consistency(media_rows)
        if not fps_ok:
            passed = False
            self._emit_status(delivery_log, f"[ERROR] Delivery failed: {fps_msg}")
        self._emit_status(
            delivery_log,
            "Video spec QC: PASSED" if passed else "Video spec QC: FAILED — delivery blocked",
        )
        return passed

    def _run_audio_spec_qc(
        self,
        delivery_log: list[str],
        media_rows: list[dict[str, object]],
    ) -> bool:
        from services.inspired_audio_validator import wav_spec_failure_message
        from services.inspired_video_validator import validate_audio_duration_vs_mp4

        self._emit_status(
            delivery_log,
            "Audio spec QC: validating COM/CFX WAV (PCM_S16LE, 48000 Hz, mono, 16-bit, duration ≤ MP4)…",
        )
        passed = True
        for row in media_rows:
            fn = str(row.get("engp_filename", ""))
            mp4_path = Path(str(row.get("output_video", "")))
            for label, ok_key, path_key, prefix in (
                ("COM.wav", "com_wav_ok", "output_com_wav", "com_wav"),
                ("CFX.wav", "cfx_wav_ok", "output_cfx_wav", "cfx_wav"),
            ):
                if str(row.get(ok_key, "")).strip().casefold() != "yes":
                    continue
                wav_path = Path(str(row.get(path_key, "")))
                if not wav_path.is_file():
                    passed = False
                    msg = f"Delivery failed: {label} missing on disk for {fn}"
                    self._emit_status(delivery_log, f"[ERROR] {msg}")
                    row[f"{prefix}_validation_ok"] = "no"
                    continue
                val = validate_packaged_wav_inspired(wav_path)
                for spec_key in ("sample_rate", "channels", "codec", "bit_depth", "bitrate", "duration"):
                    row[f"{prefix}_{spec_key}"] = val.get(spec_key, "")
                row[f"{prefix}_validation_ok"] = val.get("validation_ok", "")
                row[f"{prefix}_validation_summary"] = val.get("validation_notes", "")
                if str(val.get("validation_ok", "")).lower() != "yes":
                    passed = False
                    msg = wav_spec_failure_message(label, val)
                    self._emit_status(delivery_log, f"[ERROR] {fn}: {msg}")
                    continue
                if mp4_path.is_file():
                    dur_val = validate_audio_duration_vs_mp4(wav_path, mp4_path, label=label)
                    row[f"{prefix}_mp4_duration"] = dur_val.get("mp4_duration", "")
                    row[f"{prefix}_duration_ok"] = dur_val.get("duration_ok", "")
                    if str(dur_val.get("duration_ok", "")).lower() != "yes":
                        passed = False
                        msg = f"Delivery failed: {label} — {dur_val.get('failure_reason', 'Audio duration exceeds corresponding video duration.')}"
                        self._emit_status(delivery_log, f"[ERROR] {fn}: {msg}")
        self._emit_status(
            delivery_log,
            "Audio spec QC: PASSED" if passed else "Audio spec QC: FAILED — delivery blocked",
        )
        return passed

    def _run_media_generation_phase(
        self,
        delivery_dir: Path,
        delivery_log: list[str],
    ) -> tuple[list[dict[str, object]], list[dict[str, object]], dict[str, int]]:
        """
        Phase 5: package ENGP video + COM/CFX WAVs under ``delivery_dir/media/``.
        Does not abort the delivery on per-file failure; returns rows for XLSX reports.
        """
        media_dir = delivery_dir / "media"
        media_dir.mkdir(parents=True, exist_ok=True)
        n = len(self._mp4_paths)
        counts: dict[str, int] = {
            "video_ok": 0,
            "video_fail": 0,
            "com_ok": 0,
            "com_fail": 0,
            "cfx_ok": 0,
            "cfx_fail": 0,
            "cfx_missing_source": 0,
        }
        media_rows: list[dict[str, object]] = []
        unresolved_rows: list[dict[str, object]] = []

        def send_counts() -> None:
            self._send({"type": "media_counts", **counts})

        self._send({"type": "media_phase_begin", "total": n})
        send_counts()

        if n == 0:
            return media_rows, unresolved_rows, counts

        if not _ffmpeg_available():
            msg = "[WARN] ffmpeg not on PATH; skipping media transcoding and WAV extraction."
            self._emit_status(delivery_log, msg)
            for mp4 in self._mp4_paths:
                cfx_exp = cfx_expected_primary_report_path(mp4)
                media_rows.append(
                    {
                        "engp_filename": mp4.name,
                        "engp_source_path": str(mp4.resolve()),
                        "cfx_expected_path": str(cfx_exp),
                        "cfx_resolved_source_path": "",
                        "cfx_discovery": "n/a",
                        "cfx_duplicate_paths": "",
                        "cfx_found": "n/a",
                        "video_encode_ok": "no",
                        "video_encode_error": "ffmpeg missing",
                        "com_wav_ok": "no",
                        "com_wav_error": "ffmpeg missing",
                        "cfx_wav_ok": "no",
                        "cfx_wav_error": "ffmpeg missing",
                        "mp4_validation_summary": "",
                        "com_wav_validation_summary": "",
                        "cfx_wav_validation_summary": "",
                    }
                )
                unresolved_rows.append(
                    {
                        "engp_filename": mp4.name,
                        "cfx_expected_path": str(cfx_exp),
                        "cfx_resolved_source_path": "",
                        "cfx_discovery": "n/a",
                        "cfx_duplicate_paths": "",
                        "issue": "ffmpeg missing on PATH",
                        "video_error": "ffmpeg missing",
                        "com_wav_error": "ffmpeg missing",
                        "cfx_wav_error": "ffmpeg missing",
                    }
                )
                counts["video_fail"] += 1
                counts["com_fail"] += 1
                counts["cfx_fail"] += 1
            send_counts()
            return media_rows, unresolved_rows, counts

        self._emit_status(delivery_log, f"Media generation: writing packaged assets under {media_dir}…")

        cfx_root_cfg = clean_cfx_source_root_path(self._cfx_source_root_folder)
        cfx_tree = CfxTreeIndex()
        cfx_index_stats = CfxIndexBuildStats()
        cfx_root_expanded = ""
        cfx_log = lambda t: self._emit_status(delivery_log, t)
        if cfx_root_cfg:
            cfx_root_expanded = os.path.expanduser(cfx_root_cfg)
            emit_cfx_index_prebuild_diagnostics(
                cfx_log,
                raw=self._cfx_source_root_raw,
                cleaned=cfx_root_cfg,
                stats=CfxIndexBuildStats(expanded_root=cfx_root_expanded),
            )
            if cfx_source_root_usable_for_scan(cfx_root_cfg):
                cfx_tree, cfx_index_stats = build_cfx_basename_index(cfx_root_cfg)
                emit_cfx_index_postbuild_diagnostics(cfx_log, stats=cfx_index_stats, tree=cfx_tree)
                n_cfx_mp4 = cfx_index_stats.mp4_indexed
                self._emit_status(
                    delivery_log,
                    f"Media: CFX source root {cfx_root_expanded!r} — indexed {n_cfx_mp4} MP4 path(s) by basename "
                    f"({cfx_index_stats.dirs_scanned} director{'y' if cfx_index_stats.dirs_scanned == 1 else 'ies'} scanned).",
                )
                reports_dir = delivery_dir / "reports"
                try:
                    reports_dir.mkdir(parents=True, exist_ok=True)
                except OSError:
                    reports_dir = delivery_dir
                dbg_header = [
                    f"CFX index debug — {datetime.now().isoformat(timespec='seconds')}",
                    f"raw_root: {self._cfx_source_root_raw!r}",
                    f"cleaned_root: {cfx_root_cfg!r}",
                    f"expanded_root: {cfx_index_stats.expanded_root!r}",
                    f"directories_scanned: {cfx_index_stats.dirs_scanned}",
                    f"mp4_indexed: {cfx_index_stats.mp4_indexed}",
                    f"cloudmounter_rescan_attempted: {cfx_index_stats.rescan_attempted}",
                ]
                dbg_path = reports_dir / "cfx_index_debug.txt"
                if write_cfx_index_debug_txt(dbg_path, cfx_tree, dbg_header):
                    self._emit_status(delivery_log, f"Wrote reports/cfx_index_debug.txt ({n_cfx_mp4} MP4 entries).")
                else:
                    self._emit_status(delivery_log, "[WARN] Could not write reports/cfx_index_debug.txt.")
            else:
                emit_cfx_root_rejection_diagnostics(
                    cfx_log,
                    raw=self._cfx_source_root_raw,
                    cleaned=cfx_root_cfg,
                )

        for i, mp4 in enumerate(self._mp4_paths, start=1):
            if self._cancel.is_set():
                self._emit_status(delivery_log, "Media generation cancelled by operator.")
                break

            fn = mp4.name
            prim_cfx = cfx_primary_engp_cfx_path(mp4)
            leg_cfx = cfx_companion_path_for_engp(mp4)
            cfx_src, sibling_exp, cfx_discovery, cfx_dupe_paths = resolve_cfx_source_mp4(
                mp4, cfx_tree, emit=cfx_log
            )
            cfx_ok = cfx_src is not None
            dup_join = ""
            if cfx_dupe_paths:
                dup_join = "; ".join(str(p) for p in cfx_dupe_paths[:24])
                if len(cfx_dupe_paths) > 24:
                    dup_join += f"; … (+{len(cfx_dupe_paths) - 24} more)"

            if cfx_ok:
                if cfx_discovery == "nearest_stem_fallback":
                    self._emit_status(
                        delivery_log,
                        f"[WARN] {fn}: CFX matched by nearest stem fallback: {cfx_src}",
                    )
                self._emit_status(
                    delivery_log,
                    f"Media: {fn}: CFX source found ({cfx_discovery}): {cfx_src}",
                )
                if cfx_dupe_paths:
                    self._emit_status(
                        delivery_log,
                        f"[WARN] {fn}: duplicate CFX source candidates ({cfx_discovery}); using {cfx_src}. "
                        f"Other path(s) ({len(cfx_dupe_paths)}): {dup_join}",
                    )
            else:
                miss = (
                    f"[WARN] {fn}: CFX source missing (tried {prim_cfx.name!r} then {leg_cfx.name!r} beside ENGP"
                )
                if cfx_root_cfg:
                    miss += "; exhausted CFX tree (exact basename / casefold for ENGP_CFX and legacy CFX names)"
                    if cfx_root_expanded:
                        miss += f"; root {cfx_root_expanded!r}"
                miss += ")"
                self._emit_status(delivery_log, miss)

            out_mp4, out_com, out_cfx = media_packaged_output_paths(mp4, media_dir)

            v_err = ""
            ok_v, err_v = ffmpeg_transcode_inspired_engp_video(mp4, out_mp4, self._cancel)
            if ok_v:
                counts["video_ok"] += 1
                self._emit_status(delivery_log, f"Media: encoded video {out_mp4.name}")
            else:
                counts["video_fail"] += 1
                v_err = err_v
                self._emit_status(delivery_log, f"[WARN] {fn}: video encode failed: {err_v}")
            self._send({"type": "media_video_progress", "current": i, "total": n, "filename": fn})
            send_counts()

            c_err = ""
            ok_c, err_c = ffmpeg_extract_inspired_wav(mp4, out_com, self._cancel)
            if ok_c:
                counts["com_ok"] += 1
                self._emit_status(delivery_log, f"Media: wrote COM WAV {out_com.name}")
            else:
                counts["com_fail"] += 1
                c_err = err_c
                self._emit_status(delivery_log, f"[WARN] {fn}: COM WAV failed: {err_c}")
            self._send({"type": "media_com_progress", "current": i, "total": n, "filename": fn})
            send_counts()

            x_err = ""
            ok_x = False
            if not cfx_ok:
                counts["cfx_missing_source"] += 1
                x_err = f"CFX source not found (tried {prim_cfx.name!r} then {leg_cfx.name!r}"
                if cfx_root_cfg:
                    x_err += "; CFX tree had no match (exact basename / casefold for ENGP_CFX and legacy CFX names)"
                x_err += ")"
            else:
                ok_x, err_x = ffmpeg_extract_inspired_wav(cfx_src, out_cfx, self._cancel)
                if ok_x:
                    counts["cfx_ok"] += 1
                    self._emit_status(delivery_log, f"Media: wrote CFX WAV {out_cfx.name}")
                else:
                    counts["cfx_fail"] += 1
                    x_err = err_x
                    self._emit_status(delivery_log, f"[WARN] {fn}: CFX WAV failed: {err_x}")
            self._send({"type": "media_cfx_progress", "current": i, "total": n, "filename": fn})
            send_counts()

            mp4_val: dict[str, str] = {}
            com_val: dict[str, str] = {}
            cfx_val: dict[str, str] = {}
            if ok_v and out_mp4.is_file():
                mp4_val = validate_packaged_mp4_inspired(out_mp4)
            if ok_c and out_com.is_file():
                com_val = validate_packaged_wav_inspired(out_com)
            if ok_x and out_cfx.is_file():
                cfx_val = validate_packaged_wav_inspired(out_cfx)

            def _sumv(d: dict[str, str]) -> str:
                parts = [f"{k}={v}" for k, v in d.items() if v and k not in ("validation_ok", "validation_notes")]
                if d.get("validation_notes"):
                    parts.append(f"notes={d['validation_notes']}")
                return "; ".join(parts)

            media_rows.append(
                {
                    "engp_filename": fn,
                    "engp_source_path": str(mp4.resolve()),
                    "cfx_expected_path": str(sibling_exp),
                    "cfx_resolved_source_path": str(cfx_src) if cfx_src else "",
                    "cfx_discovery": cfx_discovery,
                    "cfx_duplicate_paths": dup_join,
                    "cfx_found": "yes" if cfx_ok else "no",
                    "output_video": str(out_mp4),
                    "output_com_wav": str(out_com),
                    "output_cfx_wav": str(out_cfx),
                    "video_encode_ok": "yes" if ok_v else "no",
                    "video_encode_error": v_err,
                    "com_wav_ok": "yes" if ok_c else "no",
                    "com_wav_error": c_err,
                    "cfx_wav_ok": "yes" if ok_x else "no",
                    "cfx_wav_error": x_err,
                    "mp4_validation_ok": mp4_val.get("validation_ok", ""),
                    "mp4_validation_summary": _sumv(mp4_val),
                    "com_wav_validation_ok": com_val.get("validation_ok", ""),
                    "com_wav_validation_summary": _sumv(com_val),
                    "cfx_wav_validation_ok": cfx_val.get("validation_ok", ""),
                    "cfx_wav_validation_summary": _sumv(cfx_val),
                }
            )

            issues: list[str] = []
            if not cfx_ok:
                issues.append("cfx_source_missing")
            if not ok_v:
                issues.append("video_encode_failed")
            if not ok_c:
                issues.append("com_wav_failed")
            if cfx_ok and not ok_x:
                issues.append("cfx_wav_failed")
            if issues:
                unresolved_rows.append(
                    {
                        "engp_filename": fn,
                        "cfx_expected_path": str(sibling_exp),
                        "cfx_resolved_source_path": str(cfx_src) if cfx_src else "",
                        "cfx_discovery": cfx_discovery,
                        "cfx_duplicate_paths": dup_join,
                        "issue": ", ".join(issues),
                        "video_error": v_err,
                        "com_wav_error": c_err,
                        "cfx_wav_error": x_err,
                    }
                )

        self._emit_status(
            delivery_log,
            "Media generation summary: "
            f"video ok/fail {counts['video_ok']}/{counts['video_fail']}, "
            f"COM ok/fail {counts['com_ok']}/{counts['com_fail']}, "
            f"CFX ok/fail {counts['cfx_ok']}/{counts['cfx_fail']}, "
            f"CFX missing source {counts['cfx_missing_source']}.",
        )
        return media_rows, unresolved_rows, counts

    def run(self) -> None:
        stamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
        delivery_name = f"Inspired_Delivery_{stamp}"
        delivery_dir: Path | None = None
        counters = ResolutionCounters(requested=len(self._mp4_paths))
        ambiguous_xlsx_rows: list[dict[str, str | int]] = []
        resolved_json: list[dict[str, object]] = []

        try:
            run_perf_start = time.perf_counter()
            run_clock_start = datetime.now()
            self._send({"type": "state", "state": "RUNNING"})
            self._send({"type": "preview_clear"})
            self._counters(counters)

            if not self._mp4_paths:
                self._send({"type": "status_line", "text": "No input files."})
                self._send({"type": "finished", "completion": "FAILED", "delivery_dir": None})
                return

            if not self._output_parent.is_dir():
                self._send({"type": "status_line", "text": "Output folder is not valid."})
                self._send({"type": "finished", "completion": "FAILED", "delivery_dir": None})
                return

            if not self._csv_only_delivery:
                if not all(
                    [
                        self._iconik_base_url,
                        self._iconik_app_id,
                        self._iconik_auth_token,
                        self._iconik_collection_uuid,
                    ]
                ):
                    self._send({"type": "status_line", "text": "Iconik Base URL, App ID, Auth Token, and Collection UUID are required."})
                    self._send({"type": "finished", "completion": "FAILED", "delivery_dir": None})
                    return
                client = IconikClient(
                    base_url=self._iconik_base_url,
                    app_id=self._iconik_app_id,
                    auth_token=self._iconik_auth_token,
                    collection_id=self._iconik_collection_uuid,
                )
            else:
                client = None  # Local CSV-only: no Iconik HTTP or search.

            delivery_dir = self._output_parent / delivery_name
            delivery_dir.mkdir(parents=False, exist_ok=False)
            delivery_log: list[str] = []
            self._emit_status(delivery_log, f"Delivery folder: {delivery_dir}")

            team_reg = begin_delivery_team_normalization()
            reg_summary = team_reg.validate_registry()
            self._emit_status(
                delivery_log,
                f"Team registry: {reg_summary.get('team_count', 0)} teams from {reg_summary.get('registry_path', '')}",
            )

            inspired_rows: InspiredDataFrame = []
            warnings_by_filename: dict[str, list[str]] = {}
            metadata_qc_result = None
            metadata_debug_exported = False
            search_result_debug_exported = False
            combined_iconik_rows: list[dict[str, str]] = []
            unmatched_mp4_rows: list[dict[str, str]] = []
            unmatched_csv_rows: list[dict[str, str]] = []
            csv_rows, csv_fn_col, csv_title_col, csv_bundle_display, csv_load_err, csv_merge_debug_rows = (
                load_merge_metadata_csv_bundle(
                    self._metadata_csv_paths,
                    mp4_paths=self._mp4_paths,
                )
            )
            csv_row_matched: set[int] = set()
            csv_active = bool(csv_rows and csv_fn_col and not csv_load_err)
            if csv_load_err and not csv_rows:
                self._emit_status(delivery_log, f"[WARN] CSV metadata: {csv_load_err}")
            elif csv_rows:
                sample_fn = _cleanup_csv_cell(csv_rows[0].get(csv_fn_col, "")) if csv_fn_col else ""
                tail = f"; filename sample={sample_fn!r}" if sample_fn else ""
                self._emit_status(
                    delivery_log,
                    f"CSV metadata ({len(csv_rows)} merged record(s) from {len(csv_merge_debug_rows)} file(s)) — {csv_bundle_display}; "
                    f"filename column={csv_fn_col!r}, title column={csv_title_col!r}{tail}",
                )
            else:
                self._emit_status(
                    delivery_log,
                    "CSV metadata: none loaded (add files under “Choose Metadata CSV(s)” or use sidecar discovery).",
                )

            if self._csv_only_delivery:
                self._emit_status(
                    delivery_log,
                    "Local CSV-only mode: Iconik search and metadata API are disabled for this run.",
                )

            n = len(self._mp4_paths)
            for i, mp4 in enumerate(self._mp4_paths):
                if self._cancel.is_set():
                    self._emit_status(delivery_log, "Cancelled by operator.")
                    self._write_outputs(
                        delivery_dir,
                        counters,
                        resolved_json,
                        ambiguous_xlsx_rows,
                        inspired_rows,
                        warnings_by_filename,
                        delivery_log,
                        duration_frames_populated=0,
                        completion="CANCELLED",
                        media_rows=[],
                        unresolved_media_rows=[],
                        media_stats={},
                        combined_iconik_rows=combined_iconik_rows,
                        unmatched_mp4_rows=unmatched_mp4_rows,
                        unmatched_csv_rows=unmatched_csv_rows,
                        csv_merge_debug_rows=csv_merge_debug_rows,
                        run_perf_start=run_perf_start,
                        run_clock_start=run_clock_start,
                        lock_delivery_outputs=self._lock_delivery_outputs,
                    )
                    self._send({"type": "finished", "completion": "CANCELLED", "delivery_dir": str(delivery_dir)})
                    return

                pct = int((i / max(n, 1)) * 85)
                self._send({"type": "progress", "value": pct})
                fn = mp4.name
                stem = mp4.stem
                self._emit_status(delivery_log, f"Resolving {fn} ({i + 1}/{n})…")

                csv_indices: list[int] = []
                csv_best_tier = 99
                csv_amb_local = False
                if csv_active:
                    csv_indices, csv_best_tier, csv_amb_local = csv_row_indices_for_mp4(
                        csv_rows,
                        filename=fn,
                        stem=stem,
                        filename_col=csv_fn_col,
                        title_col=csv_title_col,
                    )

                def register_unmatched_mp4_if_needed() -> None:
                    nonlocal counters
                    if not csv_active:
                        return
                    if csv_amb_local:
                        return
                    if csv_indices:
                        return
                    counters.unmatched_mp4 += 1
                    self._counters(counters)
                    unmatched_mp4_rows.append(
                        {
                            "mp4_filename": fn,
                            "csv_path": csv_bundle_display,
                            "reason": "no_csv_row_match_under_strict_tiers",
                        }
                    )

                def _emit_inspired_row(
                    resolved_title: str,
                    asset_id: str,
                    meta: dict | None,
                    meta_ok: bool,
                    *,
                    flat: dict[str, object] | None = None,
                ) -> None:
                    row, w = build_inspired_row(
                        filename=fn,
                        resolved_title=resolved_title,
                        asset_id=asset_id,
                        meta_doc=meta,
                        metadata_ok=meta_ok,
                        flat=flat,
                        csv_only_delivery=self._csv_only_delivery,
                    )
                    inspired_rows.append(row)
                    warnings_by_filename[fn] = w
                    for line in w:
                        self._emit_status(delivery_log, f"[WARN] {fn}: {line}")
                    self._send(
                        {
                            "type": "inspired_preview_row",
                            "filename": row.get("Filename", fn),
                            "resolved_title": resolved_title,
                            "country": row.get("Country", ""),
                            "match_date": row.get("Match Date", ""),
                            "home_team": row.get("Home Team", ""),
                            "away_team": row.get("Away Team", ""),
                            "goal_scored": row.get("Goal Scored", ""),
                            "goal_time_frames": row.get("Goal Time (Frames)", ""),
                            "normalization": inspired_norm_summary(w, filename=fn),
                        }
                    )
                    self._send(
                        {
                            "type": "inspired_stats",
                            "total": len(inspired_rows),
                            "with_warnings": count_inspired_rows_with_heuristic_warnings(
                                warnings_by_filename
                            ),
                        }
                    )

                if self._csv_only_delivery:
                    flat_csv: dict[str, object] = {}
                    if csv_active:
                        if csv_amb_local and csv_indices:
                            counters.ambiguous += 1
                            self._counters(counters)
                            amb_parts = [
                                f"csv#{ri + 1}:{csv_rows[ri].get(csv_fn_col, '')!s} ({_cleanup_csv_cell(csv_rows[ri].get('__source_path__', ''))})"
                                for ri in sorted(csv_indices)
                            ]
                            ambiguous_xlsx_rows.append(
                                {
                                    "ambiguity_kind": "csv_duplicate_metadata",
                                    "source_filename": fn,
                                    "resolved_asset_id": "",
                                    "resolved_title": "",
                                    "match_tier": csv_best_tier,
                                    "candidate_count": len(csv_indices),
                                    "candidates_id_title": " | ".join(amb_parts),
                                }
                            )
                            self._emit_status(
                                delivery_log,
                                f"[WARN] {fn}: ambiguous CSV metadata match (local CSV-only mode); merge skipped.",
                            )
                        elif csv_indices:
                            pick = min(csv_indices)
                            csv_mw: list[str] = []
                            flat_csv = merge_csv_row_into_iconik_flat({}, csv_rows[pick], merge_warnings=csv_mw)
                            for line in csv_mw:
                                warnings_by_filename.setdefault(fn, []).append(line)
                            csv_row_matched.add(pick)
                            counters.resolved += 1
                            self._counters(counters)
                            src_hint = _cleanup_csv_cell(csv_rows[pick].get("__source_path__", ""))
                            self._emit_status(
                                delivery_log,
                                f"{fn}: merged local CSV metadata (bundle slot {pick + 1}; {src_hint or csv_bundle_display}).",
                            )

                    if not flat_csv:
                        if not csv_active:
                            counters.unresolved += 1
                            self._counters(counters)
                            self._emit_status(
                                delivery_log,
                                f"[WARN] {fn}: local CSV-only mode has no active CSV metadata (add CSV paths or sidecar files).",
                            )
                        elif not (csv_amb_local and csv_indices) and not csv_indices:
                            counters.unresolved += 1
                            self._counters(counters)
                            self._emit_status(
                                delivery_log,
                                f"[WARN] {fn}: unresolved in local CSV-only mode (no CSV row match).",
                            )

                    resolved_json.append(
                        {
                            "filename": fn,
                            "resolved_asset_id": None,
                            "resolved_title": None,
                            "metadata_success": bool(flat_csv),
                        }
                    )
                    preview_csv = flatten_metadata_preview(flat=flat_csv) if flat_csv else ""
                    self._send(
                        {
                            "type": "preview_row",
                            "filename": fn,
                            "asset_id": "",
                            "title": "",
                            "metadata_ok": "yes" if flat_csv else "no",
                            "metadata_preview": preview_csv,
                        }
                    )
                    _emit_inspired_row("", "", None, False, flat=flat_csv if flat_csv else None)
                    combined_iconik_rows.append(
                        flat_metadata_to_combined_csv_row(
                            filename=fn,
                            resolved_asset_id="",
                            resolved_title="",
                            flat=flat_csv if flat_csv else None,
                        )
                    )
                    register_unmatched_mp4_if_needed()
                    continue

                assert client is not None

                hits: list[dict] = []
                search_err: str | None = None
                for q in (fn, stem, partial_fallback_query(stem)):
                    if self._cancel.is_set():
                        break
                    if q == "" and not hits:
                        continue
                    hits, err = client.search_all_pages(q, self._cancel)
                    if err:
                        search_err = err
                        self._emit_status(delivery_log, f"[WARN] {fn}: {err}")
                    if hits:
                        break

                if not hits:
                    counters.unresolved += 1
                    self._counters(counters)
                    self._emit_status(delivery_log, f"[WARN] {fn}: unresolved (no search hits).")
                    resolved_json.append(
                        {
                            "filename": fn,
                            "resolved_asset_id": None,
                            "resolved_title": None,
                            "metadata_success": False,
                        }
                    )
                    self._send(
                        {
                            "type": "preview_row",
                            "filename": fn,
                            "asset_id": "",
                            "title": "",
                            "metadata_ok": "no",
                            "metadata_preview": search_err or "",
                        }
                    )
                    _emit_inspired_row("", "", None, False)
                    combined_iconik_rows.append(
                        flat_metadata_to_combined_csv_row(
                            filename=fn,
                            resolved_asset_id="",
                            resolved_title="",
                            flat=None,
                        )
                    )
                    register_unmatched_mp4_if_needed()
                    continue

                winner, ambiguous_iconik, win_tier = pick_best_asset(hits, fn, stem)
                if winner is None and ambiguous_iconik:
                    counters.ambiguous += 1
                    self._counters(counters)
                    cand_lines: list[str] = []
                    for h in hits:
                        if _tier_for_hit(fn, stem, _asset_title(h)) == win_tier and _asset_id(h):
                            cand_lines.append(f"{_asset_id(h)}:{_asset_title(h)}")
                    ambiguous_xlsx_rows.append(
                        {
                            "ambiguity_kind": "iconik_search",
                            "source_filename": fn,
                            "resolved_asset_id": "",
                            "resolved_title": "",
                            "match_tier": win_tier,
                            "candidate_count": len(cand_lines),
                            "candidates_id_title": " | ".join(sorted(cand_lines)),
                        }
                    )
                    self._emit_status(
                        delivery_log,
                        f"[WARN] {fn}: ambiguous Iconik search match (tier {win_tier}); "
                        "no asset chosen — see reports/ambiguous_metadata_matches.xlsx.",
                    )
                    resolved_json.append(
                        {
                            "filename": fn,
                            "resolved_asset_id": None,
                            "resolved_title": None,
                            "metadata_success": False,
                            "ambiguous_iconik_tier": win_tier,
                        }
                    )
                    self._send(
                        {
                            "type": "preview_row",
                            "filename": fn,
                            "asset_id": "",
                            "title": "",
                            "metadata_ok": "no",
                            "metadata_preview": "ambiguous Iconik match (no winner)",
                        }
                    )
                    _emit_inspired_row("", "", None, False, flat=None)
                    combined_iconik_rows.append(
                        flat_metadata_to_combined_csv_row(
                            filename=fn,
                            resolved_asset_id="",
                            resolved_title="",
                            flat=None,
                        )
                    )
                    register_unmatched_mp4_if_needed()
                    continue

                if winner is None:
                    counters.unresolved += 1
                    self._counters(counters)
                    self._emit_status(
                        delivery_log,
                        f"[WARN] {fn}: unresolved (no Iconik hit matched filename/title under strict tiers).",
                    )
                    resolved_json.append(
                        {
                            "filename": fn,
                            "resolved_asset_id": None,
                            "resolved_title": None,
                            "metadata_success": False,
                        }
                    )
                    self._send(
                        {
                            "type": "preview_row",
                            "filename": fn,
                            "asset_id": "",
                            "title": "",
                            "metadata_ok": "no",
                            "metadata_preview": "",
                        }
                    )
                    _emit_inspired_row("", "", None, False)
                    combined_iconik_rows.append(
                        flat_metadata_to_combined_csv_row(
                            filename=fn,
                            resolved_asset_id="",
                            resolved_title="",
                            flat=None,
                        )
                    )
                    register_unmatched_mp4_if_needed()
                    continue

                counters.resolved += 1
                self._counters(counters)

                aid = _asset_id(winner)
                title_guess = _asset_title(winner)

                if not search_result_debug_exported:
                    search_result_debug_exported = True
                    export_raw_search_result_debug_json(
                        delivery_dir,
                        winner,
                        source_filename=fn,
                        resolved_asset_id=aid,
                        resolved_title=title_guess,
                        emit=lambda m: self._emit_status(delivery_log, m),
                    )
                    print_iconik_search_result_structure_diagnostics(
                        winner,
                        emit=lambda m: self._emit_status(delivery_log, m),
                    )

                emit_d = lambda m: self._emit_status(delivery_log, m)
                flat_hit, slugs_hit = flatten_metadata_from_search_hit(winner)
                print_search_hit_flattened_metadata_diagnostics(flat_hit, slugs_hit, emit_d)

                meta_vid = iconik_search_hit_metadata_version_id(winner)
                meta_doc, meta_ok, meta_err = client.get_asset_metadata(
                    aid,
                    self._cancel,
                    version_id=meta_vid,
                )
                flat_api_highlight = (
                    extract_iconik_tenant_metadata_flat(meta_doc) if meta_ok and meta_doc else {}
                )
                flat_api_bm, _ = build_flat_metadata_map(meta_doc if meta_ok and meta_doc else None)
                flat_api = merge_flat_metadata_prefer_primary(flat_api_highlight, flat_api_bm)
                slugs_api = tuple(sorted(flat_api.keys(), key=lambda s: (str(s).casefold(), str(s))))
                print_metadata_flatten_compare_diagnostics(flat_hit, slugs_hit, flat_api, slugs_api, emit_d)

                flat_merged = merge_flat_metadata_prefer_primary(flat_hit, flat_api)

                if not meta_ok and not flat_merged:
                    counters.metadata_failures += 1
                    self._counters(counters)
                    self._emit_status(delivery_log, f"[WARN] {fn}: metadata request failed ({meta_err}).")
                elif not meta_ok:
                    self._emit_status(
                        delivery_log,
                        f"[WARN] {fn}: metadata GET failed ({meta_err}); using search-hit projection only.",
                    )

                title_final = title_guess
                if meta_doc:
                    t2 = meta_doc.get("title")
                    if isinstance(t2, str) and t2.strip():
                        title_final = t2.strip()

                if meta_ok and isinstance(meta_doc, dict) and not metadata_debug_exported:
                    metadata_debug_exported = True
                    export_raw_metadata_debug_json(
                        delivery_dir,
                        meta_doc,
                        resolved_asset_id=aid,
                        resolved_title=title_final,
                        emit=emit_d,
                        flat_from_search_hit=flat_hit,
                        slugs_from_search_hit=slugs_hit,
                        flat_from_metadata_endpoint=flat_api,
                        flat_merged=flat_merged,
                    )
                    self._emit_status(
                        delivery_log,
                        f"Wrote raw_metadata_debug.json for first metadata GET 200 (asset {aid}).",
                    )

                flat_for_row: dict[str, object] = dict(flat_merged) if flat_merged else {}
                if csv_active:
                    if csv_amb_local:
                        counters.ambiguous += 1
                        self._counters(counters)
                        parts: list[str] = []
                        for ri in sorted(csv_indices):
                            src_p = _cleanup_csv_cell(csv_rows[ri].get("__source_path__", ""))
                            parts.append(f"csv#{ri + 1}:{csv_rows[ri].get(csv_fn_col, '')!s} ({src_p})")
                        ambiguous_xlsx_rows.append(
                            {
                                "ambiguity_kind": "csv_duplicate_metadata",
                                "source_filename": fn,
                                "resolved_asset_id": aid,
                                "resolved_title": title_final,
                                "match_tier": csv_best_tier,
                                "candidate_count": len(csv_indices),
                                "candidates_id_title": " | ".join(parts),
                            }
                        )
                        self._emit_status(
                            delivery_log,
                            f"[WARN] {fn}: ambiguous CSV metadata match; Iconik metadata kept, CSV merge skipped.",
                        )
                    elif csv_indices:
                        pick = min(csv_indices)
                        csv_mw2: list[str] = []
                        flat_for_row = merge_csv_row_into_iconik_flat(flat_for_row, csv_rows[pick], merge_warnings=csv_mw2)
                        for line in csv_mw2:
                            warnings_by_filename.setdefault(fn, []).append(line)
                        csv_row_matched.add(pick)
                        self._emit_status(
                            delivery_log,
                            f"{fn}: merged optional CSV metadata (bundle slot {pick + 1}; "
                            f"{_cleanup_csv_cell(csv_rows[pick].get('__source_path__', '')) or csv_bundle_display}).",
                        )

                resolved_json.append(
                    {
                        "filename": fn,
                        "resolved_asset_id": aid,
                        "resolved_title": title_final,
                        "metadata_success": bool(flat_merged) or meta_ok,
                    }
                )
                has_meta = bool(flat_for_row)
                preview = (
                    flatten_metadata_preview(flat=flat_for_row)
                    if has_meta
                    else ((meta_err or "") if not meta_ok else flatten_metadata_preview(meta_doc=meta_doc))
                )
                self._send(
                    {
                        "type": "preview_row",
                        "filename": fn,
                        "asset_id": aid,
                        "title": title_final,
                        "metadata_ok": "yes" if has_meta or meta_ok else "no",
                        "metadata_preview": preview,
                    }
                )
                merged_for_row = flat_for_row if flat_for_row else None
                _emit_inspired_row(title_final, aid, meta_doc, meta_ok, flat=merged_for_row)
                combined_iconik_rows.append(
                    flat_metadata_to_combined_csv_row(
                        filename=fn,
                        resolved_asset_id=aid,
                        resolved_title=title_final,
                        flat=flat_for_row if flat_for_row else None,
                    )
                )
                register_unmatched_mp4_if_needed()

            self._emit_status(
                delivery_log,
                "Team normalization QC: comparing metadata Home/Away pair to filename team pair (unordered)…",
            )
            for mp4, row in zip(self._mp4_paths, inspired_rows):
                fn = mp4.name
                for msg in validate_inspired_row_teams_vs_filename(
                    fn,
                    row,
                    registry=get_delivery_team_registry(),
                ):
                    warnings_by_filename.setdefault(fn, []).append(msg)
                    self._emit_status(delivery_log, f"[WARN] {fn}: {msg}")

            metadata_qc_result = self._run_metadata_goal_qc(
                delivery_log,
                inspired_rows,
                warnings_by_filename,
            )

            if not self._cancel.is_set() and metadata_qc_result.error_count > 0:
                self._emit_status(
                    delivery_log,
                    f"Delivery aborted: metadata QC failed ({metadata_qc_result.error_count} error(s)).",
                )
                self._write_outputs(
                    delivery_dir,
                    counters,
                    resolved_json,
                    ambiguous_xlsx_rows,
                    inspired_rows,
                    warnings_by_filename,
                    delivery_log,
                    duration_frames_populated=0,
                    completion="FAILED",
                    media_rows=[],
                    unresolved_media_rows=[],
                    media_stats={},
                    combined_iconik_rows=combined_iconik_rows,
                    unmatched_mp4_rows=unmatched_mp4_rows,
                    unmatched_csv_rows=unmatched_csv_rows,
                    csv_merge_debug_rows=csv_merge_debug_rows,
                    run_perf_start=run_perf_start,
                    run_clock_start=run_clock_start,
                    lock_delivery_outputs=self._lock_delivery_outputs,
                    metadata_qc_result=metadata_qc_result,
                    video_spec_ok=False,
                    audio_spec_ok=False,
                    frame_alignment_ok=False,
                )
                self._send({"type": "finished", "completion": "FAILED", "delivery_dir": str(delivery_dir)})
                return

            if csv_active:
                for idx, row in enumerate(csv_rows):
                    if idx in csv_row_matched:
                        continue
                    counters.unmatched_csv += 1
                    fc = _cleanup_csv_cell(row.get(csv_fn_col, ""))
                    tc = _cleanup_csv_cell(row.get(csv_title_col, "")) if csv_title_col else ""
                    src = _cleanup_csv_cell(row.get("__source_path__", "")) or csv_bundle_display
                    unmatched_csv_rows.append(
                        {
                            "csv_row_index": str(idx + 1),
                            "csv_path": src,
                            "filename_cell": fc,
                            "title_cell": tc,
                            "reason": "no_matching_mp4_under_strict_tiers",
                        }
                    )
                if unmatched_csv_rows:
                    self._counters(counters)

            # FFprobe: Duration (Frames) — Iconik-resolved rows use Id; local CSV-only probes every input MP4.
            if self._csv_only_delivery:
                resolved_total = len(self._mp4_paths)
                self._emit_status(
                    delivery_log,
                    "FFprobe: enriching Duration (Frames) for all input files (local CSV-only mode)…",
                )
            else:
                resolved_total = sum(1 for r in inspired_rows if r.get("Id", "").strip())
                self._emit_status(delivery_log, "FFprobe: enriching Duration (Frames) for resolved assets…")
            populated_frames = 0
            done_probe = 0
            self._send({"type": "ffprobe_progress", "current": 0, "total": max(resolved_total, 1), "populated": 0})
            for mp4, row in zip(self._mp4_paths, inspired_rows):
                if self._cancel.is_set():
                    break
                if not self._csv_only_delivery and not row.get("Id", "").strip():
                    continue
                done_probe += 1
                fn = mp4.name
                frames, err = ffprobe_duration_frames(mp4, FFPROBE_CACHE, self._cancel)
                if err:
                    row["Duration (Frames)"] = ""
                    if err != "cancelled":
                        msg = f"[WARN] {fn}: {err}"
                        self._emit_status(delivery_log, msg)
                        warnings_by_filename.setdefault(fn, []).append(err)
                else:
                    row["Duration (Frames)"] = frames
                    populated_frames += 1
                self._send(
                    {
                        "type": "ffprobe_progress",
                        "current": done_probe,
                        "total": max(resolved_total, 1),
                        "populated": populated_frames,
                    }
                )
            self._send({"type": "duration_frames_count", "count": populated_frames})
            self._emit_status(
                delivery_log,
                f"FFprobe pass complete: Duration (Frames) populated for {populated_frames} / {max(resolved_total, 0)} resolved file(s).",
            )

            media_rows: list[dict[str, object]] = []
            unresolved_media_rows: list[dict[str, object]] = []
            media_stats: dict[str, int] = {}
            if not self._cancel.is_set():
                media_rows, unresolved_media_rows, media_stats = self._run_media_generation_phase(
                    delivery_dir,
                    delivery_log,
                )

            video_spec_ok = True
            if not self._cancel.is_set() and media_rows:
                video_spec_ok = self._run_video_spec_qc(delivery_log, media_rows)

            audio_spec_ok = True
            if not self._cancel.is_set() and media_rows:
                audio_spec_ok = self._run_audio_spec_qc(delivery_log, media_rows)

            frame_alignment_ok = True
            if not self._cancel.is_set() and self._mp4_paths:
                frame_alignment_ok = self._run_frame_alignment_qc(
                    delivery_dir,
                    delivery_log,
                    media_rows,
                )

            if self._cancel.is_set():
                self._emit_status(delivery_log, "Cancelled by operator.")
                self._write_outputs(
                    delivery_dir,
                    counters,
                    resolved_json,
                    ambiguous_xlsx_rows,
                    inspired_rows,
                    warnings_by_filename,
                    delivery_log,
                    duration_frames_populated=populated_frames,
                    completion="CANCELLED",
                    media_rows=media_rows,
                    unresolved_media_rows=unresolved_media_rows,
                    media_stats=media_stats,
                    combined_iconik_rows=combined_iconik_rows,
                    unmatched_mp4_rows=unmatched_mp4_rows,
                    unmatched_csv_rows=unmatched_csv_rows,
                    csv_merge_debug_rows=csv_merge_debug_rows,
                    run_perf_start=run_perf_start,
                    run_clock_start=run_clock_start,
                    lock_delivery_outputs=self._lock_delivery_outputs,
                )
                self._send({"type": "finished", "completion": "CANCELLED", "delivery_dir": str(delivery_dir)})
                return

            if not video_spec_ok:
                self._emit_status(
                    delivery_log,
                    "Delivery aborted: video specification QC failed — delivery blocked.",
                )
                self._write_outputs(
                    delivery_dir,
                    counters,
                    resolved_json,
                    ambiguous_xlsx_rows,
                    inspired_rows,
                    warnings_by_filename,
                    delivery_log,
                    duration_frames_populated=populated_frames,
                    completion="FAILED",
                    media_rows=media_rows,
                    unresolved_media_rows=unresolved_media_rows,
                    media_stats=media_stats,
                    combined_iconik_rows=combined_iconik_rows,
                    unmatched_mp4_rows=unmatched_mp4_rows,
                    unmatched_csv_rows=unmatched_csv_rows,
                    csv_merge_debug_rows=csv_merge_debug_rows,
                    run_perf_start=run_perf_start,
                    run_clock_start=run_clock_start,
                    lock_delivery_outputs=self._lock_delivery_outputs,
                    metadata_qc_result=metadata_qc_result,
                    video_spec_ok=False,
                    audio_spec_ok=audio_spec_ok,
                    frame_alignment_ok=frame_alignment_ok,
                )
                self._send({"type": "finished", "completion": "FAILED", "delivery_dir": str(delivery_dir)})
                return

            if not audio_spec_ok:
                self._emit_status(
                    delivery_log,
                    "Delivery aborted: audio WAV specification QC failed — delivery blocked.",
                )
                self._write_outputs(
                    delivery_dir,
                    counters,
                    resolved_json,
                    ambiguous_xlsx_rows,
                    inspired_rows,
                    warnings_by_filename,
                    delivery_log,
                    duration_frames_populated=populated_frames,
                    completion="FAILED",
                    media_rows=media_rows,
                    unresolved_media_rows=unresolved_media_rows,
                    media_stats=media_stats,
                    combined_iconik_rows=combined_iconik_rows,
                    unmatched_mp4_rows=unmatched_mp4_rows,
                    unmatched_csv_rows=unmatched_csv_rows,
                    csv_merge_debug_rows=csv_merge_debug_rows,
                    run_perf_start=run_perf_start,
                    run_clock_start=run_clock_start,
                    lock_delivery_outputs=self._lock_delivery_outputs,
                    metadata_qc_result=metadata_qc_result,
                    video_spec_ok=video_spec_ok,
                    audio_spec_ok=False,
                    frame_alignment_ok=frame_alignment_ok,
                )
                self._send({"type": "finished", "completion": "FAILED", "delivery_dir": str(delivery_dir)})
                return

            if not frame_alignment_ok:
                self._emit_status(
                    delivery_log,
                    "Delivery aborted: frame alignment QC failed — upload/delivery completion blocked.",
                )
                self._write_outputs(
                    delivery_dir,
                    counters,
                    resolved_json,
                    ambiguous_xlsx_rows,
                    inspired_rows,
                    warnings_by_filename,
                    delivery_log,
                    duration_frames_populated=populated_frames,
                    completion="FAILED",
                    media_rows=media_rows,
                    unresolved_media_rows=unresolved_media_rows,
                    media_stats=media_stats,
                    combined_iconik_rows=combined_iconik_rows,
                    unmatched_mp4_rows=unmatched_mp4_rows,
                    unmatched_csv_rows=unmatched_csv_rows,
                    csv_merge_debug_rows=csv_merge_debug_rows,
                    run_perf_start=run_perf_start,
                    run_clock_start=run_clock_start,
                    lock_delivery_outputs=self._lock_delivery_outputs,
                    metadata_qc_result=metadata_qc_result,
                    video_spec_ok=video_spec_ok,
                    audio_spec_ok=audio_spec_ok,
                    frame_alignment_ok=False,
                )
                self._send({"type": "finished", "completion": "FAILED", "delivery_dir": str(delivery_dir)})
                return

            self._write_outputs(
                delivery_dir,
                counters,
                resolved_json,
                ambiguous_xlsx_rows,
                inspired_rows,
                warnings_by_filename,
                delivery_log,
                duration_frames_populated=populated_frames,
                completion="COMPLETED",
                media_rows=media_rows,
                unresolved_media_rows=unresolved_media_rows,
                media_stats=media_stats,
                combined_iconik_rows=combined_iconik_rows,
                unmatched_mp4_rows=unmatched_mp4_rows,
                unmatched_csv_rows=unmatched_csv_rows,
                csv_merge_debug_rows=csv_merge_debug_rows,
                run_perf_start=run_perf_start,
                run_clock_start=run_clock_start,
                lock_delivery_outputs=self._lock_delivery_outputs,
            metadata_qc_result=metadata_qc_result,
            video_spec_ok=video_spec_ok,
            audio_spec_ok=audio_spec_ok,
            frame_alignment_ok=frame_alignment_ok,
        )

            self._send({"type": "progress", "value": 100})
            self._emit_status(delivery_log, f"Delivery outputs written under {delivery_dir}")
            self._send({"type": "finished", "completion": "COMPLETED", "delivery_dir": str(delivery_dir)})
        except FileExistsError:
            self._send({"type": "status_line", "text": "Delivery folder already exists; aborting."})
            self._send({"type": "finished", "completion": "FAILED", "delivery_dir": None})
        except Exception as exc:  # noqa: BLE001
            self._send({"type": "status_line", "text": f"Error: {exc}"})
            self._send({"type": "finished", "completion": "FAILED", "delivery_dir": str(delivery_dir) if delivery_dir else None})

    def _write_outputs(
        self,
        delivery_dir: Path,
        counters: ResolutionCounters,
        resolved_json: list[dict[str, object]],
        ambiguous_xlsx_rows: list[dict[str, str | int]],
        inspired_rows: InspiredDataFrame,
        warnings_by_filename: dict[str, list[str]],
        delivery_log: list[str],
        *,
        duration_frames_populated: int = 0,
        completion: str,
        media_rows: list[dict[str, object]] | None = None,
        unresolved_media_rows: list[dict[str, object]] | None = None,
        media_stats: dict[str, int] | None = None,
        combined_iconik_rows: list[dict[str, str]] | None = None,
        unmatched_mp4_rows: list[dict[str, str]] | None = None,
        unmatched_csv_rows: list[dict[str, str]] | None = None,
        csv_merge_debug_rows: list[dict[str, str]] | None = None,
        run_perf_start: float | None = None,
        run_clock_start: datetime | None = None,
        lock_delivery_outputs: bool = False,
        metadata_qc_result: object | None = None,
        video_spec_ok: bool = True,
        audio_spec_ok: bool = True,
        frame_alignment_ok: bool = True,
    ) -> None:
        media_rows = list(media_rows or [])
        unresolved_media_rows = list(unresolved_media_rows or [])
        media_stats = dict(media_stats or {})
        combined_iconik_rows = list(combined_iconik_rows or [])
        unmatched_mp4_rows = list(unmatched_mp4_rows or [])
        unmatched_csv_rows = list(unmatched_csv_rows or [])
        csv_merge_debug_rows = list(csv_merge_debug_rows or [])
        reports_dir = delivery_dir / "reports"
        try:
            reports_dir.mkdir(parents=True, exist_ok=True)
        except OSError:
            reports_dir = delivery_dir

        summary_path = delivery_dir / "delivery_summary.txt"
        ts_iso = datetime.now().isoformat(timespec="seconds")
        insp_total = len(inspired_rows)
        insp_warn = sum(1 for v in warnings_by_filename.values() if v)
        lines = [
            f"Timestamp: {ts_iso}",
            f"Selected file count: {len(self._mp4_paths)}",
            f"Completion state: {completion}",
            "",
            "Resolution counters:",
            f"  requested: {counters.requested}",
            f"  resolved: {counters.resolved}",
            f"  unresolved: {counters.unresolved}",
            f"  ambiguous: {counters.ambiguous}",
            f"  metadata_failures: {counters.metadata_failures}",
            f"  unmatched_mp4_vs_csv: {counters.unmatched_mp4}",
            f"  unmatched_csv_rows_vs_mp4: {counters.unmatched_csv}",
            "",
            "Inspired normalization:",
            f"  total_normalized_rows: {insp_total}",
            f"  rows_with_warnings: {insp_warn}",
            f"  duration_frames_populated: {duration_frames_populated}",
            "",
            "Media generation (Phase 5):",
            f"  video_ok: {media_stats.get('video_ok', 0)}",
            f"  video_fail: {media_stats.get('video_fail', 0)}",
            f"  com_wav_ok: {media_stats.get('com_ok', 0)}",
            f"  com_wav_fail: {media_stats.get('com_fail', 0)}",
            f"  cfx_wav_ok: {media_stats.get('cfx_ok', 0)}",
            f"  cfx_wav_fail: {media_stats.get('cfx_fail', 0)}",
            f"  cfx_source_missing: {media_stats.get('cfx_missing_source', 0)}",
            "",
            "Frame alignment QC (post-render):",
            f"  see reports/frame_alignment_qc_report.txt",
            "",
            "Team normalization (shared registry):",
            f"  unknown team values: {len(take_unknown_team_records())}",
            f"  see reports/unknown_teams.csv",
            "",
            "Outputs (delivery root):",
            "  delivery_log.txt",
            "  delivery_summary.txt",
            "  inspired_qc_summary.txt",
            "  resolved_assets.json",
            "  raw_search_result_debug.json (first resolved asset, diagnostics)",
            "  normalized_metadata_preview.json",
            "  inspired_metadata.xlsx",
            "  combined_iconik_metadata.csv",
            "  delivery_manifest.txt",
            "  media/ (packaged ENGP MP4 + COM/CFX WAV)",
            "",
            "Outputs (reports/):",
            "  delivery_validation_report.xlsx",
            "  media_generation_report.xlsx",
            "  unresolved_media_report.xlsx",
            "  unmatched_mp4_report.xlsx",
            "  unmatched_csv_rows_report.xlsx",
            "  ambiguous_metadata_matches.xlsx",
            "  csv_merge_debug_report.xlsx",
            "  cfx_index_debug.txt",
            "  frame_alignment_qc_report.txt",
            "  unknown_teams.csv",
            "  metadata_qc_report.xlsx",
            "  inspired_qc_summary.txt",
            "",
            "Optional post-run:",
            "  .delivery_locked (marker + read-only tree when “Lock delivery outputs” is enabled)",
        ]
        summary_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
        self._emit_status(delivery_log, "Wrote delivery_summary.txt.")

        from services.inspired_qc_summary import build_inspired_qc_summary_from_run, write_inspired_qc_summary

        qc_summary = build_inspired_qc_summary_from_run(
            files_processed=len(self._mp4_paths),
            media_rows=media_rows,
            metadata_qc_result=metadata_qc_result,
            video_spec_ok=video_spec_ok,
            audio_spec_ok=audio_spec_ok,
            frame_alignment_ok=frame_alignment_ok,
            warnings_by_filename=warnings_by_filename,
            unknown_team_count=len(take_unknown_team_records()),
        )
        for qc_path in (delivery_dir / "inspired_qc_summary.txt", reports_dir / "inspired_qc_summary.txt"):
            if write_inspired_qc_summary(qc_path, qc_summary):
                self._emit_status(delivery_log, f"Wrote {qc_path.relative_to(delivery_dir)!s}.")
            else:
                self._emit_status(delivery_log, f"[WARN] Could not write {qc_path.name}.")

        metadata_qc_rows: list[dict[str, str]] = []
        if metadata_qc_result is not None:
            for issue in getattr(metadata_qc_result, "issues", []) or []:
                metadata_qc_rows.append(
                    {
                        "filename": getattr(issue, "filename", ""),
                        "rule_id": getattr(issue, "rule_id", ""),
                        "severity": getattr(issue, "severity", ""),
                        "field": getattr(issue, "field", ""),
                        "value": getattr(issue, "value", ""),
                        "message": getattr(issue, "message", ""),
                    }
                )
        if write_metadata_qc_report_xlsx(reports_dir / "metadata_qc_report.xlsx", metadata_qc_rows):
            self._emit_status(
                delivery_log,
                f"Wrote reports/metadata_qc_report.xlsx ({len(metadata_qc_rows)} issue(s)).",
            )
        else:
            self._emit_status(delivery_log, "[WARN] openpyxl not available; skipped metadata_qc_report.xlsx.")

        resolved_path = delivery_dir / "resolved_assets.json"
        resolved_path.write_text(json.dumps(resolved_json, indent=2), encoding="utf-8")
        self._emit_status(delivery_log, "Wrote resolved_assets.json.")

        preview_doc = build_normalized_preview_document(
            rows=inspired_rows,
            warnings_by_filename=warnings_by_filename,
            completion=completion,
            duration_frames_populated=duration_frames_populated,
        )
        (delivery_dir / "normalized_metadata_preview.json").write_text(
            json.dumps(preview_doc, indent=2),
            encoding="utf-8",
        )
        self._emit_status(delivery_log, "Wrote normalized_metadata_preview.json.")

        combined_path = delivery_dir / "combined_iconik_metadata.csv"
        if write_combined_iconik_metadata_csv(combined_path, combined_iconik_rows):
            self._emit_status(
                delivery_log,
                f"Wrote combined_iconik_metadata.csv ({len(combined_iconik_rows)} row(s)).",
            )
        else:
            self._emit_status(delivery_log, "[WARN] Could not write combined_iconik_metadata.csv.")

        ambiguous_path = reports_dir / "ambiguous_metadata_matches.xlsx"
        if write_ambiguous_xlsx(ambiguous_path, ambiguous_xlsx_rows):
            self._emit_status(
                delivery_log,
                f"Wrote reports/ambiguous_metadata_matches.xlsx ({len(ambiguous_xlsx_rows)} row(s)).",
            )
        else:
            self._emit_status(
                delivery_log,
                "[WARN] openpyxl is not installed; cannot write reports/ambiguous_metadata_matches.xlsx.",
            )

        self._send({"type": "xlsx_progress", "value": 5})
        self._emit_status(delivery_log, "Writing inspired_metadata.xlsx…")
        inspired_path = delivery_dir / "inspired_metadata.xlsx"
        if write_inspired_metadata_xlsx(inspired_path, inspired_rows, warnings_by_filename=warnings_by_filename):
            self._send({"type": "xlsx_progress", "value": 100})
            self._emit_status(
                delivery_log,
                f"Wrote inspired_metadata.xlsx ({len(inspired_rows)} row(s), {len(INSPIRED_COLUMNS)} columns).",
            )
        else:
            self._send({"type": "xlsx_progress", "value": 0})
            self._emit_status(
                delivery_log,
                "[WARN] openpyxl is not installed; cannot write inspired_metadata.xlsx.",
            )

        media_rep = reports_dir / "media_generation_report.xlsx"
        if write_media_generation_report_xlsx(media_rep, media_rows):
            self._emit_status(
                delivery_log,
                f"Wrote reports/media_generation_report.xlsx ({len(media_rows)} row(s)).",
            )
        else:
            self._emit_status(
                delivery_log,
                "[WARN] openpyxl is not installed; cannot write reports/media_generation_report.xlsx.",
            )

        if write_unresolved_media_report_xlsx(reports_dir / "unresolved_media_report.xlsx", unresolved_media_rows):
            self._emit_status(delivery_log, "Wrote reports/unresolved_media_report.xlsx.")
        else:
            self._emit_status(
                delivery_log,
                "[WARN] openpyxl is not installed; cannot write reports/unresolved_media_report.xlsx.",
            )

        if write_unmatched_mp4_report_xlsx(reports_dir / "unmatched_mp4_report.xlsx", unmatched_mp4_rows):
            self._emit_status(delivery_log, "Wrote reports/unmatched_mp4_report.xlsx.")
        else:
            self._emit_status(
                delivery_log,
                "[WARN] openpyxl is not installed; cannot write reports/unmatched_mp4_report.xlsx.",
            )

        if write_unmatched_csv_rows_report_xlsx(reports_dir / "unmatched_csv_rows_report.xlsx", unmatched_csv_rows):
            self._emit_status(delivery_log, "Wrote reports/unmatched_csv_rows_report.xlsx.")
        else:
            self._emit_status(
                delivery_log,
                "[WARN] openpyxl is not installed; cannot write reports/unmatched_csv_rows_report.xlsx.",
            )

        if write_csv_merge_debug_report_xlsx(reports_dir / "csv_merge_debug_report.xlsx", csv_merge_debug_rows):
            self._emit_status(
                delivery_log,
                f"Wrote reports/csv_merge_debug_report.xlsx ({len(csv_merge_debug_rows)} source file(s)).",
            )
        else:
            self._emit_status(
                delivery_log,
                "[WARN] openpyxl is not installed; cannot write reports/csv_merge_debug_report.xlsx.",
            )

        unknown_team_records = take_unknown_team_records()
        unknown_teams_path = reports_dir / "unknown_teams.csv"
        if write_unknown_teams_csv(unknown_teams_path, unknown_team_records):
            self._emit_status(
                delivery_log,
                f"Wrote reports/unknown_teams.csv ({len(unknown_team_records)} unresolved value(s)).",
            )
        else:
            self._emit_status(delivery_log, "[WARN] Could not write reports/unknown_teams.csv.")

        run_t0 = float(run_perf_start) if run_perf_start is not None else time.perf_counter()
        run_duration = max(0.0, time.perf_counter() - run_t0)
        run_clock = run_clock_start or datetime.now()
        val_rows = build_delivery_validation_rows(
            self._mp4_paths,
            inspired_rows,
            media_rows,
            warnings_by_filename,
            completion=completion,
            ffmpeg_available=_ffmpeg_available(),
            csv_only_delivery=self._csv_only_delivery,
        )
        val_path = reports_dir / "delivery_validation_report.xlsx"
        if write_delivery_validation_report_xlsx(val_path, val_rows):
            self._emit_status(
                delivery_log,
                f"Wrote reports/delivery_validation_report.xlsx ({len(val_rows)} row(s)).",
            )
        else:
            self._emit_status(
                delivery_log,
                "[WARN] openpyxl is not installed; cannot write reports/delivery_validation_report.xlsx.",
            )

        log_path = delivery_dir / "delivery_log.txt"
        log_path.write_text("\n".join(delivery_log) + "\n", encoding="utf-8")
        self._send({"type": "status_line", "text": "Wrote delivery_log.txt."})

        manifest_path = delivery_dir / "delivery_manifest.txt"
        if write_delivery_manifest_txt(
            manifest_path,
            delivery_root=delivery_dir,
            completion=completion,
            counters=counters,
            media_stats=media_stats,
            n_mp4=len(self._mp4_paths),
            run_started_clock=run_clock,
            run_duration_sec=run_duration,
        ):
            self._emit_status(delivery_log, "Wrote delivery_manifest.txt.")
        else:
            self._emit_status(delivery_log, "[WARN] Could not write delivery_manifest.txt.")

        if lock_delivery_outputs:
            lock_delivery_outputs_tree(delivery_dir, emit=lambda m: self._emit_status(delivery_log, m))


@dataclass
class AppState:
    worker_thread: threading.Thread | None = None
    cancel_event: threading.Event = field(default_factory=threading.Event)
    msg_queue: queue.Queue = field(default_factory=queue.Queue)


class InspiredDeliveryGeneratorApp:
    def __init__(self, root: ctk.CTk) -> None:
        self.root = root
        root.title(APP_NAME)
        self.settings_path = _default_settings_path()
        self.settings = load_settings(self.settings_path)
        self.state = AppState()

        self._mp4_list: list[Path] = []
        self._metadata_csv_list: list[Path] = []
        self.lock_delivery_outputs_var = tk.BooleanVar(value=False)
        self.csv_only_delivery_var = tk.BooleanVar(value=False)
        self._last_delivery_dir: str | None = None
        self._last_completion: str = "READY"
        self._last_inspired_warnings: int = 0
        self._last_inspired_total_rows: int = 0
        self._last_media_counts: dict[str, int] = {}
        self._last_unmatched_mp4: int = 0
        self._last_unmatched_csv: int = 0
        self.show_diagnostics_var = tk.BooleanVar(value=False)

        self._build_ui()
        self._restore_settings()
        self._refresh_operator_checklist()
        self._set_workflow_state("READY")
        self._reset_counters_ui()
        self._refresh_results_summary()
        self._update_open_folder_buttons()
        self._poll_queue()

        root.protocol("WM_DELETE_WINDOW", self._on_close)

    def _configure_treeview_neon_style(self) -> None:
        col = self._COL
        style = ttk.Style(self.root)
        try:
            style.theme_use("clam")
        except tk.TclError:
            pass
        style.configure(
            "Neon.Treeview",
            background=col["input"],
            foreground=col["text"],
            fieldbackground=col["input"],
            rowheight=24,
            font=("SF Pro Text", 12) if sys.platform == "darwin" else ("Segoe UI", 12),
        )
        style.configure("Neon.Treeview.Heading", background=col["card_hl"], foreground=col["text"], relief="flat")
        style.map(
            "Neon.Treeview",
            background=[("selected", col["cyan"])],
            foreground=[("selected", "#0f172a")],
        )

    def _ctk_prog_set(self, w: object, ratio: float) -> None:
        r = max(0.0, min(1.0, float(ratio)))
        fn = getattr(w, "set", None)
        if callable(fn):
            fn(r)
        else:
            w["value"] = r * 100.0

    def _toggle_advanced_iconik(self) -> None:
        if self.show_advanced_iconik_var.get():
            self._advanced_iconik_frame.pack(fill=tk.X, pady=(10, 0))
        else:
            self._advanced_iconik_frame.pack_forget()

    def _build_ui(self) -> None:
        C = {
            "bg": "#0f172a",
            "sidebar": "#0b1220",
            "card": "#1e293b",
            "card_hl": "#24324a",
            "cyan": "#06b6d4",
            "magenta": "#ec4899",
            "text": "#e5e7eb",
            "muted": "#94a3b8",
            "warn": "#f59e0b",
            "ok": "#22c55e",
            "fail": "#ef4444",
            "input": "#111c2e",
        }
        self._COL = C
        self._main_wrap = 640
        self.root.configure(fg_color=C["bg"])

        self._f_title = ctk.CTkFont(size=20, weight="bold")
        self._f_section = ctk.CTkFont(size=16, weight="bold")
        self._f_body = ctk.CTkFont(size=13)
        self._f_muted = ctk.CTkFont(size=12)
        self._f_btn = ctk.CTkFont(size=12)
        self._f_mono = ctk.CTkFont(family="Menlo" if sys.platform == "darwin" else "Consolas", size=12)

        self._configure_treeview_neon_style()

        shell = ctk.CTkFrame(self.root, fg_color=C["bg"], corner_radius=0)
        shell.pack(fill=tk.BOTH, expand=True)

        body = ctk.CTkFrame(shell, fg_color=C["bg"], corner_radius=0)
        body.pack(side=tk.TOP, fill=tk.BOTH, expand=True)

        side = ctk.CTkFrame(body, fg_color=C["sidebar"], corner_radius=0, width=300)
        side.pack(side=tk.LEFT, fill=tk.Y, padx=(0, 4), pady=0)
        side.pack_propagate(False)
        self._sidebar_frame = side
        sp = ctk.CTkFrame(side, fg_color="transparent")
        sp.pack(fill=tk.BOTH, expand=True, padx=16, pady=18)

        ctk.CTkLabel(sp, text=APP_NAME, font=self._f_title, text_color=C["text"], anchor="w").pack(anchor="w")
        ctk.CTkLabel(
            sp,
            text="Dark-neon delivery dashboard for Inspired ENGP outputs.",
            font=self._f_muted,
            text_color=C["muted"],
            wraplength=260,
            justify="left",
            anchor="w",
        ).pack(anchor="w", pady=(6, 14))

        self.state_var = tk.StringVar(value="READY")
        self._state_badge = ctk.CTkLabel(
            sp,
            textvariable=self.state_var,
            font=self._f_section,
            corner_radius=14,
            fg_color=C["ok"],
            text_color="#052e16",
            anchor="w",
            height=44,
            padx=14,
            pady=8,
        )
        self._state_badge.pack(fill=tk.X, pady=(0, 12))
        self._apply_state_badge_style("READY")

        ctk.CTkFrame(sp, fg_color=C["card_hl"], corner_radius=2, height=2).pack(fill=tk.X, pady=(0, 12))

        ctk.CTkLabel(sp, text="Setup checklist", font=self._f_section, text_color=C["text"], anchor="w").pack(anchor="w")
        ctk.CTkLabel(
            sp,
            text="Required items must be OK to start. Recommended items warn only.",
            font=self._f_muted,
            text_color=C["muted"],
            wraplength=260,
            justify="left",
            anchor="w",
        ).pack(anchor="w", pady=(4, 10))

        self._checklist_status_vars: dict[str, tk.StringVar] = {}

        def chk_row(parent: ctk.CTkFrame, row: int, label: str, key: str) -> None:
            self._checklist_status_vars[key] = tk.StringVar(value="…")
            ctk.CTkLabel(parent, text=label, font=self._f_muted, text_color=C["muted"], anchor="w").grid(
                row=row, column=0, sticky="w", padx=(0, 8), pady=2
            )
            ctk.CTkLabel(parent, textvariable=self._checklist_status_vars[key], font=self._f_body, text_color=C["text"], width=110, anchor="w").grid(
                row=row, column=1, sticky="w", pady=2
            )

        req = ctk.CTkFrame(sp, fg_color=C["card"], corner_radius=12, border_width=1, border_color=C["card_hl"])
        req.pack(fill=tk.X, pady=(0, 8))
        r_in = ctk.CTkFrame(req, fg_color="transparent")
        r_in.pack(fill=tk.X, padx=10, pady=10)
        ctk.CTkLabel(r_in, text="Required", font=self._f_body, text_color=C["cyan"], anchor="w").grid(row=0, column=0, columnspan=2, sticky="w", pady=(0, 6))
        for i, (lab, k) in enumerate(
            [
                ("MP4 files", "req_mp4"),
                ("Iconik Base URL", "req_base"),
                ("App ID", "req_app"),
                ("Auth Token", "req_token"),
                ("Collection UUID", "req_coll"),
                ("Output folder", "req_out"),
                ("CFX source root", "req_cfx"),
            ]
        ):
            chk_row(r_in, i + 1, lab, k)
        r_in.grid_columnconfigure(1, weight=1)

        rec = ctk.CTkFrame(sp, fg_color=C["card"], corner_radius=12, border_width=1, border_color=C["card_hl"])
        rec.pack(fill=tk.X, pady=(0, 16))
        rec_in = ctk.CTkFrame(rec, fg_color="transparent")
        rec_in.pack(fill=tk.X, padx=10, pady=10)
        ctk.CTkLabel(rec_in, text="Recommended", font=self._f_body, text_color=C["magenta"], anchor="w").grid(
            row=0, column=0, columnspan=2, sticky="w", pady=(0, 6)
        )
        for i, (lab, k) in enumerate(
            [("CFX *_ENGP_CFX.mp4", "rec_cfx"), ("ffmpeg", "rec_ffmpeg"), ("ffprobe", "rec_ffprobe")]
        ):
            chk_row(rec_in, i + 1, lab, k)
        rec_in.grid_columnconfigure(1, weight=1)

        ctk.CTkFrame(sp, fg_color=C["card_hl"], corner_radius=2, height=2).pack(fill=tk.X, pady=(0, 12))

        self.generate_btn = ctk.CTkButton(
            sp,
            text="Generate Inspired Delivery",
            font=self._f_btn,
            height=40,
            corner_radius=22,
            fg_color=C["cyan"],
            text_color=C["bg"],
            hover_color="#22d3ee",
            command=self._generate,
        )
        self.generate_btn.pack(fill=tk.X, pady=(0, 8))
        self.cancel_btn = ctk.CTkButton(
            sp,
            text="Cancel Run",
            font=self._f_btn,
            height=36,
            corner_radius=20,
            fg_color="transparent",
            border_width=2,
            border_color=C["magenta"],
            text_color=C["magenta"],
            hover_color=C["card_hl"],
            command=self._cancel,
            state="disabled",
        )
        self.cancel_btn.pack(fill=tk.X, pady=(0, 10))

        self.open_delivery_btn = ctk.CTkButton(
            sp,
            text="Open Delivery Folder",
            font=self._f_btn,
            height=34,
            corner_radius=18,
            fg_color="transparent",
            border_width=2,
            border_color=C["cyan"],
            text_color=C["cyan"],
            hover_color=C["card_hl"],
            command=self._open_last_delivery_folder,
            state="disabled",
        )
        self.open_delivery_btn.pack(fill=tk.X, pady=(0, 6))
        self.open_reports_btn = ctk.CTkButton(
            sp,
            text="Open Reports Folder",
            font=self._f_btn,
            height=34,
            corner_radius=18,
            fg_color="transparent",
            border_width=2,
            border_color=C["cyan"],
            text_color=C["cyan"],
            hover_color=C["card_hl"],
            command=self._open_last_reports_folder,
            state="disabled",
        )
        self.open_reports_btn.pack(fill=tk.X)

        main = ctk.CTkFrame(body, fg_color=C["bg"], corner_radius=0)
        main.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(4, 0), pady=0)
        self._body_scroll_zone = main

        scroll = ctk.CTkScrollableFrame(
            main,
            fg_color=C["bg"],
            corner_radius=0,
            scrollbar_fg_color=C["card_hl"],
            scrollbar_button_color=C["cyan"],
            scrollbar_button_hover_color="#22d3ee",
        )
        scroll.pack(fill=tk.BOTH, expand=True, padx=6, pady=(0, 6))
        self._main_scroll = scroll

        def card(title: str, subtitle: str | None) -> ctk.CTkFrame:
            outer = ctk.CTkFrame(scroll, fg_color=C["card"], corner_radius=18, border_width=1, border_color=C["card_hl"])
            outer.pack(fill=tk.X, padx=8, pady=(0, 16))
            inner = ctk.CTkFrame(outer, fg_color="transparent")
            inner.pack(fill=tk.BOTH, expand=True, padx=18, pady=16)
            ctk.CTkLabel(inner, text=title, font=self._f_section, text_color=C["text"], anchor="w").pack(anchor="w")
            if subtitle:
                ctk.CTkLabel(
                    inner,
                    text=subtitle,
                    font=self._f_muted,
                    text_color=C["muted"],
                    wraplength=self._main_wrap,
                    justify="left",
                    anchor="w",
                ).pack(anchor="w", pady=(4, 12))
            body_fr = ctk.CTkFrame(inner, fg_color="transparent")
            body_fr.pack(fill=tk.BOTH, expand=True)
            return body_fr

        # --- 1 Input ---
        in_body = card(
            "Input files",
            "Add ENGP MP4 files for this delivery. Order does not change processing.",
        )
        self.drop_frame = tk.Frame(in_body, bg=C["input"], highlightthickness=1, highlightbackground=C["card_hl"])
        self.drop_frame.pack(fill=tk.BOTH, expand=True, pady=(0, 10))
        self.drop_label = tk.Label(
            self.drop_frame,
            text="",
            fg=C["muted"],
            bg=C["input"],
            justify=tk.CENTER,
            font=("SF Pro Text", 12),
        )
        self.drop_label.pack(expand=True, fill=tk.BOTH, padx=12, pady=14)

        dnd_active = False
        if _TKDND_AVAILABLE:
            try:
                self.drop_frame.drop_target_register(DND_FILES)
                self.drop_frame.dnd_bind("<<Drop>>", self._on_drop)
                dnd_active = True
            except tk.TclError:
                dnd_active = False

        if dnd_active:
            self.drop_label.configure(text="Drag and drop MP4 files here")
        else:
            self.drop_label.configure(text="Add MP4 files for this delivery.")

        self._input_dnd_note = ctk.CTkLabel(
            in_body,
            text="" if dnd_active else "Drag/drop unavailable — use Add MP4 Files.",
            font=self._f_muted,
            text_color=C["muted"],
            anchor="w",
            justify="left",
        )
        self._input_dnd_note.pack(fill=tk.X, pady=(2, 6))
        if dnd_active:
            self._input_dnd_note.pack_forget()

        br = ctk.CTkFrame(in_body, fg_color="transparent")
        br.pack(fill=tk.X)
        ctk.CTkButton(br, text="Add MP4 Files", font=self._f_btn, height=32, corner_radius=18, fg_color=C["cyan"], text_color=C["bg"], command=self._browse_mp4s).pack(side=tk.LEFT)
        ctk.CTkButton(br, text="Remove selected", font=self._f_btn, height=32, corner_radius=16, fg_color=C["card_hl"], text_color=C["text"], command=self._remove_selected).pack(side=tk.LEFT, padx=(8, 0))
        ctk.CTkButton(br, text="Clear Files", font=self._f_btn, height=32, corner_radius=16, fg_color=C["card_hl"], text_color=C["text"], command=self._clear_list).pack(side=tk.LEFT, padx=(8, 0))
        lf = tk.Frame(in_body, bg=C["card"], highlightthickness=0)
        lf.pack(fill=tk.BOTH, expand=True, pady=(10, 0))
        lb_sb = ttk.Scrollbar(lf)
        lb_sb.pack(side=tk.RIGHT, fill=tk.Y)
        self.file_listbox = tk.Listbox(
            lf,
            height=7,
            selectmode=tk.EXTENDED,
            yscrollcommand=lb_sb.set,
            bg=C["input"],
            fg=C["text"],
            highlightthickness=0,
            bd=0,
            font=("SF Pro Text", 12) if sys.platform == "darwin" else ("Segoe UI", 12),
            selectbackground=C["cyan"],
            selectforeground=C["bg"],
            activestyle="none",
        )
        self.file_listbox.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        lb_sb.config(command=self.file_listbox.yview)

        ctk.CTkLabel(in_body, text="Choose Metadata CSV(s)", font=self._f_body, text_color=C["cyan"], anchor="w").pack(
            anchor="w", pady=(14, 0)
        )
        ctk.CTkLabel(
            in_body,
            text="Explicit paths (optional). If empty, the app still looks for metadata.csv / combined_metadata.csv beside your MP4s.",
            font=self._f_muted,
            text_color=C["muted"],
            wraplength=self._main_wrap,
            anchor="w",
            justify="left",
        ).pack(anchor="w", pady=(0, 6))
        self.metadata_csv_summary_var = tk.StringVar(value="No CSV files selected.")
        ctk.CTkLabel(
            in_body,
            textvariable=self.metadata_csv_summary_var,
            font=self._f_muted,
            text_color=C["muted"],
            wraplength=self._main_wrap,
            anchor="w",
            justify="left",
        ).pack(anchor="w", pady=(0, 6))
        csv_btn_row = ctk.CTkFrame(in_body, fg_color="transparent")
        csv_btn_row.pack(fill=tk.X, pady=(0, 6))
        ctk.CTkButton(
            csv_btn_row,
            text="Add Metadata CSV(s)",
            font=self._f_btn,
            height=30,
            corner_radius=16,
            fg_color=C["card_hl"],
            text_color=C["text"],
            command=self._browse_metadata_csvs,
        ).pack(side=tk.LEFT)
        ctk.CTkButton(
            csv_btn_row,
            text="Clear CSV list",
            font=self._f_btn,
            height=30,
            corner_radius=16,
            fg_color=C["card_hl"],
            text_color=C["text"],
            command=self._clear_metadata_csvs,
        ).pack(side=tk.LEFT, padx=(8, 0))
        self.lock_delivery_chk = ctk.CTkCheckBox(
            in_body,
            text="Lock delivery outputs after completion (read-only + .delivery_locked marker)",
            variable=self.lock_delivery_outputs_var,
            font=self._f_muted,
            text_color=C["muted"],
        )
        self.lock_delivery_chk.pack(anchor="w", pady=(4, 0))
        self.csv_only_chk = ctk.CTkCheckBox(
            in_body,
            text="Local CSV only — do not call Iconik (metadata from CSV / per-file basename when needed)",
            variable=self.csv_only_delivery_var,
            font=self._f_muted,
            text_color=C["muted"],
        )
        self.csv_only_chk.pack(anchor="w", pady=(4, 0))

        # --- 2 Configuration ---
        cfg = card("Configuration", "Output location, CFX search root, and optional Iconik credentials.")
        self.output_var = tk.StringVar()
        ctk.CTkLabel(cfg, text="Output folder", font=self._f_body, text_color=C["cyan"], anchor="w").pack(anchor="w")
        ctk.CTkLabel(cfg, text="Each run creates a new subfolder here.", font=self._f_muted, text_color=C["muted"], anchor="w").pack(anchor="w", pady=(0, 6))
        orow = ctk.CTkFrame(cfg, fg_color="transparent")
        orow.pack(fill=tk.X)
        self.output_entry = ctk.CTkEntry(orow, textvariable=self.output_var, height=36, corner_radius=10, fg_color=C["input"], border_color=C["card_hl"])
        self.output_entry.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 8))
        ctk.CTkButton(orow, text="Choose Output Folder", font=self._f_btn, width=170, height=34, corner_radius=16, fg_color=C["card_hl"], command=self._browse_output).pack(side=tk.RIGHT)

        self.cfx_source_root_var = tk.StringVar()
        ctk.CTkLabel(cfg, text="CFX source root", font=self._f_body, text_color=C["cyan"], anchor="w").pack(anchor="w", pady=(14, 0))
        ctk.CTkLabel(
            cfg,
            text="Top folder for CFX MP4s (names ending in _ENGP_CFX.mp4).",
            font=self._f_muted,
            text_color=C["muted"],
            anchor="w",
        ).pack(anchor="w", pady=(0, 6))
        crow = ctk.CTkFrame(cfg, fg_color="transparent")
        crow.pack(fill=tk.X)
        self.cfx_source_root_entry = ctk.CTkEntry(crow, textvariable=self.cfx_source_root_var, height=36, corner_radius=10, fg_color=C["input"], border_color=C["card_hl"])
        self.cfx_source_root_entry.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 8))
        ctk.CTkButton(crow, text="Choose CFX Root", font=self._f_btn, width=150, height=34, corner_radius=16, fg_color=C["card_hl"], command=self._browse_cfx_source_root).pack(side=tk.RIGHT)

        self.open_output_btn = ctk.CTkButton(
            cfg,
            text="Open parent output folder",
            font=self._f_btn,
            height=30,
            corner_radius=14,
            fg_color="transparent",
            border_width=1,
            border_color=C["muted"],
            text_color=C["muted"],
            hover_color=C["card_hl"],
            command=self._open_output_folder,
        )
        self.open_output_btn.pack(anchor="w", pady=(12, 0))

        self.show_advanced_iconik_var = tk.BooleanVar(value=False)
        ctk.CTkCheckBox(
            cfg,
            text="Show Advanced / Iconik settings",
            variable=self.show_advanced_iconik_var,
            command=self._toggle_advanced_iconik,
            font=self._f_muted,
            text_color=C["muted"],
            fg_color=C["cyan"],
            hover_color=C["card_hl"],
        ).pack(anchor="w", pady=(14, 0))

        self._advanced_iconik_frame = ctk.CTkFrame(cfg, fg_color="transparent")
        self.iconik_base_var = tk.StringVar()
        self.iconik_app_var = tk.StringVar()
        self.iconik_token_var = tk.StringVar()
        self.iconik_collection_var = tk.StringVar()
        g = self._advanced_iconik_frame
        rows = [
            ("Base URL", self.iconik_base_var, False),
            ("App ID", self.iconik_app_var, False),
            ("Auth Token", self.iconik_token_var, True),
            ("Collection UUID", self.iconik_collection_var, False),
        ]
        for i, (lab, var, is_pw) in enumerate(rows):
            ctk.CTkLabel(g, text=lab, font=self._f_muted, text_color=C["muted"], anchor="w").grid(row=i, column=0, sticky="w", padx=(0, 10), pady=6)
            kw = dict(textvariable=var, height=32, corner_radius=8, fg_color=C["input"], border_color=C["card_hl"])
            if is_pw:
                kw["show"] = "•"
            ctk.CTkEntry(g, **kw).grid(row=i, column=1, sticky="ew", pady=6)
        g.grid_columnconfigure(1, weight=1)

        # --- 3 Progress ---
        prog = card("Progress", "Library matching, overall progress, frame timing, and spreadsheet export.")
        ctk.CTkLabel(prog, text="Library matching", font=self._f_body, text_color=C["cyan"], anchor="w").pack(anchor="w")
        cnt = ctk.CTkFrame(prog, fg_color="transparent")
        cnt.pack(fill=tk.X, pady=(0, 10))
        self._cnt_labels = {}
        names = [
            ("requested", "Requested"),
            ("resolved", "Resolved"),
            ("unresolved", "Unresolved"),
            ("ambiguous", "Ambiguous"),
            ("metadata_failures", "Field issues"),
        ]
        for col, (key, lab) in enumerate(names):
            ctk.CTkLabel(cnt, text=f"{lab}:", font=self._f_muted, text_color=C["muted"]).grid(row=0, column=col * 2, padx=(0, 4), pady=2)
            v = tk.StringVar(value="0")
            self._cnt_labels[key] = v
            ctk.CTkLabel(cnt, textvariable=v, font=self._f_body, text_color=C["text"], width=36).grid(row=0, column=col * 2 + 1, padx=(0, 10), pady=2)

        ctk.CTkLabel(prog, text="Overall progress", font=self._f_body, text_color=C["cyan"], anchor="w").pack(anchor="w", pady=(6, 4))
        self.progress = ctk.CTkProgressBar(prog, height=14, corner_radius=8, progress_color=C["cyan"], fg_color=C["card_hl"])
        self.progress.pack(fill=tk.X, pady=(0, 8))
        self.progress.set(0)

        ctk.CTkLabel(prog, text="Frame timing (Duration / Frames)", font=self._f_body, text_color=C["cyan"], anchor="w").pack(anchor="w")
        self.ffprobe_progress_var = tk.StringVar(value="Frame timing lookup: idle")
        self.duration_frames_pop_var = tk.StringVar(value="Duration (Frames) filled: 0")
        ctk.CTkLabel(prog, textvariable=self.ffprobe_progress_var, font=self._f_muted, text_color=C["muted"], anchor="w").pack(anchor="w")
        ctk.CTkLabel(prog, textvariable=self.duration_frames_pop_var, font=self._f_muted, text_color=C["muted"], anchor="w").pack(anchor="w", pady=(0, 8))

        ctk.CTkLabel(prog, text="Spreadsheet export", font=self._f_body, text_color=C["cyan"], anchor="w").pack(anchor="w")
        self.xlsx_progress_var = tk.StringVar(value="Spreadsheet export: idle")
        ctk.CTkLabel(prog, textvariable=self.xlsx_progress_var, font=self._f_muted, text_color=C["muted"], anchor="w").pack(anchor="w")
        self.xlsx_bar = ctk.CTkProgressBar(prog, height=12, corner_radius=6, progress_color=C["magenta"], fg_color=C["card_hl"])
        self.xlsx_bar.pack(fill=tk.X, pady=(4, 0))
        self.xlsx_bar.set(0)

        # --- 4 Metadata preview ---
        meta = card("Metadata preview", "Resolved assets and normalized spreadsheet rows.")
        ctk.CTkLabel(meta, text="Resolved assets", font=self._f_body, text_color=C["cyan"], anchor="w").pack(anchor="w")
        ph = tk.Frame(meta, bg=C["card_hl"], highlightthickness=0)
        ph.pack(fill=tk.BOTH, expand=True, pady=(6, 10))
        pscroll_y = ttk.Scrollbar(ph)
        pscroll_y.pack(side=tk.RIGHT, fill=tk.Y)
        pscroll_x = ttk.Scrollbar(ph, orient=tk.HORIZONTAL)
        pscroll_x.pack(side=tk.BOTTOM, fill=tk.X)
        cols = ("filename", "asset_id", "title", "metadata_ok", "metadata_preview")
        self.preview_tree = ttk.Treeview(
            ph,
            columns=cols,
            show="headings",
            yscrollcommand=pscroll_y.set,
            xscrollcommand=pscroll_x.set,
            height=5,
            style="Neon.Treeview",
        )
        self.preview_tree.heading("filename", text="Filename")
        self.preview_tree.heading("asset_id", text="Asset ID")
        self.preview_tree.heading("title", text="Title")
        self.preview_tree.heading("metadata_ok", text="OK")
        self.preview_tree.heading("metadata_preview", text="Metadata preview")
        self.preview_tree.column("filename", width=110, stretch=False)
        self.preview_tree.column("asset_id", width=180, stretch=False)
        self.preview_tree.column("title", width=130, stretch=True)
        self.preview_tree.column("metadata_ok", width=44, stretch=False)
        self.preview_tree.column("metadata_preview", width=240, stretch=True)
        self.preview_tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        pscroll_y.config(command=self.preview_tree.yview)
        pscroll_x.config(command=self.preview_tree.xview)

        ctk.CTkLabel(meta, text="Normalized rows", font=self._f_body, text_color=C["cyan"], anchor="w").pack(anchor="w", pady=(4, 0))
        ins_top = ctk.CTkFrame(meta, fg_color="transparent")
        ins_top.pack(fill=tk.X, pady=(4, 4))
        self.inspired_total_var = tk.StringVar(value="Total normalized rows: 0")
        self.inspired_warn_var = tk.StringVar(value="Rows with warnings: 0")
        ctk.CTkLabel(ins_top, textvariable=self.inspired_total_var, font=self._f_muted, text_color=C["muted"]).pack(side=tk.LEFT, padx=(0, 16))
        ctk.CTkLabel(ins_top, textvariable=self.inspired_warn_var, font=self._f_muted, text_color=C["muted"]).pack(side=tk.LEFT)
        ih = tk.Frame(meta, bg=C["card_hl"], highlightthickness=0)
        ih.pack(fill=tk.BOTH, expand=True)
        iscroll_y = ttk.Scrollbar(ih)
        iscroll_y.pack(side=tk.RIGHT, fill=tk.Y)
        iscroll_x = ttk.Scrollbar(ih, orient=tk.HORIZONTAL)
        iscroll_x.pack(side=tk.BOTTOM, fill=tk.X)
        icols = (
            "filename",
            "resolved_title",
            "country",
            "match_date",
            "home_team",
            "away_team",
            "goal_scored",
            "goal_time_frames",
            "normalization",
        )
        self.inspired_tree = ttk.Treeview(
            ih,
            columns=icols,
            show="headings",
            yscrollcommand=iscroll_y.set,
            xscrollcommand=iscroll_x.set,
            height=5,
            style="Neon.Treeview",
        )
        for c, t in zip(
            icols,
            [
                "Filename",
                "Resolved title",
                "Country",
                "Match date",
                "Home",
                "Away",
                "Goal",
                "Goal time (fr)",
                "Normalization",
            ],
        ):
            self.inspired_tree.heading(c, text=t)
        self.inspired_tree.column("filename", width=100, stretch=False)
        self.inspired_tree.column("resolved_title", width=130, stretch=True)
        self.inspired_tree.column("country", width=56, stretch=False)
        self.inspired_tree.column("match_date", width=86, stretch=False)
        self.inspired_tree.column("home_team", width=80, stretch=False)
        self.inspired_tree.column("away_team", width=80, stretch=False)
        self.inspired_tree.column("goal_scored", width=64, stretch=False)
        self.inspired_tree.column("goal_time_frames", width=88, stretch=False)
        self.inspired_tree.column("normalization", width=180, stretch=True)
        self.inspired_tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        iscroll_y.config(command=self.inspired_tree.yview)
        iscroll_x.config(command=self.inspired_tree.xview)

        # --- 5 Media generation ---
        med = card("Media generation", "Video repackage and COM / CFX WAV outputs per file.")
        self.media_video_var = tk.StringVar(value="Video packaging: idle")
        self.media_com_var = tk.StringVar(value="COM audio (WAV): idle")
        self.media_cfx_var = tk.StringVar(value="CFX audio (WAV): idle")
        self.media_counts_var = tk.StringVar(value="Media summary: —")
        ctk.CTkLabel(med, textvariable=self.media_video_var, font=self._f_muted, text_color=C["muted"], anchor="w").pack(anchor="w")
        self.media_video_pb = ctk.CTkProgressBar(med, height=10, corner_radius=5, progress_color=C["cyan"], fg_color=C["card_hl"])
        self.media_video_pb.pack(fill=tk.X, pady=(2, 8))
        self.media_video_pb.set(0)
        ctk.CTkLabel(med, textvariable=self.media_com_var, font=self._f_muted, text_color=C["muted"], anchor="w").pack(anchor="w")
        self.media_com_pb = ctk.CTkProgressBar(med, height=10, corner_radius=5, progress_color=C["magenta"], fg_color=C["card_hl"])
        self.media_com_pb.pack(fill=tk.X, pady=(2, 8))
        self.media_com_pb.set(0)
        ctk.CTkLabel(med, textvariable=self.media_cfx_var, font=self._f_muted, text_color=C["muted"], anchor="w").pack(anchor="w")
        self.media_cfx_pb = ctk.CTkProgressBar(med, height=10, corner_radius=5, progress_color=C["cyan"], fg_color=C["card_hl"])
        self.media_cfx_pb.pack(fill=tk.X, pady=(2, 8))
        self.media_cfx_pb.set(0)
        ctk.CTkLabel(med, textvariable=self.media_counts_var, font=self._f_muted, text_color=C["muted"], anchor="w").pack(anchor="w")

        # --- 6 Results (last) ---
        res = card("Results summary", "Attention items, delivery folder, files, metrics, and report names.")
        self.results_attention_var = tk.StringVar(value="Items needing attention: —")
        self.results_delivery_var = tk.StringVar(value="Last delivery folder: —")
        self.results_files_var = tk.StringVar(value="Files generated: —")
        self.results_reports_var = tk.StringVar(
            value="Delivery root: inspired_metadata.xlsx, combined_iconik_metadata.csv, delivery_summary.txt, "
            "delivery_log.txt, delivery_manifest.txt, JSON previews, media/. Reports: delivery_validation_report.xlsx, "
            "csv_merge_debug_report.xlsx, cfx_index_debug.txt, media_generation_report.xlsx, unresolved_media_report.xlsx, unmatched_mp4_report.xlsx, "
            "unmatched_csv_rows_report.xlsx, ambiguous_metadata_matches.xlsx. Optional: .delivery_locked when locking is on."
        )
        ctk.CTkLabel(res, textvariable=self.results_attention_var, font=self._f_body, text_color=C["text"], wraplength=self._main_wrap, anchor="w", justify="left").pack(anchor="w", pady=(0, 8))
        ctk.CTkLabel(res, textvariable=self.results_delivery_var, font=self._f_muted, text_color=C["muted"], wraplength=self._main_wrap, anchor="w", justify="left").pack(anchor="w", pady=(0, 8))
        ctk.CTkLabel(res, textvariable=self.results_files_var, font=self._f_muted, text_color=C["muted"], wraplength=self._main_wrap, anchor="w", justify="left").pack(anchor="w", pady=(0, 8))
        self.results_metrics_var = tk.StringVar(value="Last run summary: —")
        ctk.CTkLabel(res, textvariable=self.results_metrics_var, font=self._f_body, text_color=C["text"], wraplength=self._main_wrap, anchor="w", justify="left").pack(anchor="w", pady=(0, 10))
        ctk.CTkLabel(res, textvariable=self.results_reports_var, font=self._f_muted, text_color=C["muted"], wraplength=self._main_wrap, anchor="w", justify="left").pack(anchor="w")

        # --- Bottom technical log ---
        bottom = ctk.CTkFrame(shell, fg_color=C["bg"], corner_radius=0)
        bottom.pack(side=tk.BOTTOM, fill=tk.X)
        ctk.CTkCheckBox(
            bottom,
            text="Show technical log (support — hidden by default)",
            variable=self.show_diagnostics_var,
            command=self._toggle_diagnostics,
            font=self._f_muted,
            text_color=C["muted"],
            fg_color=C["cyan"],
            hover_color=C["card_hl"],
        ).pack(anchor="w", padx=16, pady=(8, 4))

        self._diagnostics_outer = ctk.CTkFrame(bottom, fg_color=C["card"], corner_radius=12, border_width=1, border_color=C["card_hl"])
        self.diagnostics_text = ctk.CTkTextbox(
            self._diagnostics_outer,
            height=140,
            font=self._f_mono,
            fg_color=C["input"],
            text_color=C["text"],
            border_color=C["card_hl"],
            border_width=1,
            corner_radius=10,
            wrap="word",
            state="disabled",
        )
        self.diagnostics_text.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)

        self._wire_checklist_traces()
        self.file_listbox.bind("<<ListboxSelect>>", lambda _e: self._refresh_operator_checklist())

        self.root.minsize(1020, 640)
        self.root.geometry("1120x980")

    def _reset_counters_ui(self) -> None:
        for var in self._cnt_labels.values():
            var.set("0")

    def _reset_enrichment_ui(self) -> None:
        self.ffprobe_progress_var.set("Frame timing lookup: idle")
        self.xlsx_progress_var.set("Spreadsheet export: idle")
        self.duration_frames_pop_var.set("Duration (Frames) filled: 0")
        self.xlsx_bar.set(0)
        self.media_video_var.set("Video packaging: idle")
        self.media_com_var.set("COM audio (WAV): idle")
        self.media_cfx_var.set("CFX audio (WAV): idle")
        self.media_counts_var.set("Media summary: —")
        self.media_video_pb.set(0)
        self.media_com_pb.set(0)
        self.media_cfx_pb.set(0)

    def _restore_settings(self) -> None:
        if self.settings.output_folder:
            self.output_var.set(self.settings.output_folder)
        self.iconik_base_var.set(self.settings.iconik_base_url)
        self.iconik_app_var.set(self.settings.iconik_app_id)
        self.iconik_token_var.set(self.settings.iconik_auth_token)
        self.iconik_collection_var.set(self.settings.iconik_collection_uuid)
        self.cfx_source_root_var.set(clean_cfx_source_root_path(self.settings.cfx_source_root_folder))
        self._metadata_csv_list = [Path(p) for p in self.settings.metadata_csv_paths if str(p).strip()]
        self.lock_delivery_outputs_var.set(bool(self.settings.lock_delivery_outputs))
        self.csv_only_delivery_var.set(bool(self.settings.csv_only_delivery))
        self._refresh_metadata_csv_summary()
        if self.settings.window_geometry:
            try:
                self.root.geometry(self.settings.window_geometry)
            except tk.TclError:
                pass
        else:
            self.root.geometry("1100x980")

    def _persist_iconik_to_settings(self) -> None:
        self.settings.iconik_base_url = self.iconik_base_var.get().strip()
        self.settings.iconik_app_id = self.iconik_app_var.get().strip()
        self.settings.iconik_auth_token = self.iconik_token_var.get().strip()
        self.settings.iconik_collection_uuid = self.iconik_collection_var.get().strip()
        cfx_clean = clean_cfx_source_root_path(self.cfx_source_root_var.get())
        self.cfx_source_root_var.set(cfx_clean)
        self.settings.cfx_source_root_folder = cfx_clean

    def _append_status(self, line: str) -> None:
        self.diagnostics_text.configure(state="normal")
        self.diagnostics_text.insert("end", line + "\n")
        self.diagnostics_text.see("end")
        self.diagnostics_text.configure(state="disabled")

    def _set_workflow_state(self, state: str) -> None:
        self.state_var.set(state)
        self._apply_state_badge_style(state)

    def _apply_state_badge_style(self, state: str) -> None:
        C = getattr(self, "_COL", {})
        cyan = C.get("cyan", "#06b6d4")
        ok = C.get("ok", "#22c55e")
        fail = C.get("fail", "#ef4444")
        warn = C.get("warn", "#f59e0b")
        muted = C.get("muted", "#94a3b8")
        card = C.get("card", "#1e293b")
        styles: dict[str, tuple[str, str]] = {
            "READY": ("#0f172a", ok),
            "RUNNING": ("#0f172a", cyan),
            "COMPLETED": ("#0f172a", ok),
            "COMPLETED WITH WARNINGS": ("#0f172a", warn),
            "FAILED": ("#fef2f2", fail),
            "CANCELLED": ("#e5e7eb", muted),
        }
        tc, fc = styles.get(state, ("#e5e7eb", card))
        self._state_badge.configure(text_color=tc, fg_color=fc)

    def _toggle_diagnostics(self) -> None:
        if self.show_diagnostics_var.get():
            self._diagnostics_outer.pack(fill=tk.BOTH, expand=False, padx=16, pady=(0, 12))
        else:
            self._diagnostics_outer.pack_forget()

    def _open_path_in_os(self, path: Path | str | None) -> None:
        if not path:
            messagebox.showinfo(APP_NAME, "No folder is available yet.")
            return
        p = Path(str(path).strip())
        if not p.is_dir():
            messagebox.showwarning(APP_NAME, f"Folder does not exist or is not a directory:\n{p}")
            return
        try:
            if sys.platform == "darwin":
                subprocess.run(["open", str(p)], check=False)
            elif os.name == "nt":
                os.startfile(str(p))  # type: ignore[attr-defined]
            else:
                subprocess.run(["xdg-open", str(p)], check=False)
        except OSError as exc:
            messagebox.showerror(APP_NAME, f"Could not open folder:\n{exc}")

    def _open_output_folder(self) -> None:
        out = self.output_var.get().strip()
        if not out:
            messagebox.showwarning(APP_NAME, "Choose an output folder under Configuration first.")
            return
        self._open_path_in_os(out)

    def _open_last_delivery_folder(self) -> None:
        self._open_path_in_os(self._last_delivery_dir)

    def _open_last_reports_folder(self) -> None:
        base = self._last_delivery_dir
        if not base:
            messagebox.showinfo(APP_NAME, "No folder is available yet.")
            return
        p = Path(base.strip()) / "reports"
        self._open_path_in_os(p if p.is_dir() else Path(base.strip()))

    def _run_has_warnings(self) -> bool:
        if self._last_inspired_warnings > 0:
            return True
        c = self._cnt_labels
        try:
            if int(c["unresolved"].get() or 0) > 0:
                return True
            if int(c["ambiguous"].get() or 0) > 0:
                return True
            if int(c["metadata_failures"].get() or 0) > 0:
                return True
        except (ValueError, tk.TclError):
            pass
        if int(getattr(self, "_last_unmatched_mp4", 0) or 0) > 0:
            return True
        if int(getattr(self, "_last_unmatched_csv", 0) or 0) > 0:
            return True
        m = self._last_media_counts
        for k in ("video_fail", "com_fail", "cfx_fail", "cfx_missing_source"):
            try:
                if int(m.get(k, 0) or 0) > 0:
                    return True
            except (TypeError, ValueError):
                pass
        return False

    def _format_last_run_metrics(self) -> str:
        m = self._last_media_counts

        def g(k: str, default: int = 0) -> int:
            try:
                return int(m.get(k, default) or 0)
            except (TypeError, ValueError):
                return default

        c = self._cnt_labels
        try:
            u = int(c["unresolved"].get() or 0)
            a = int(c["ambiguous"].get() or 0)
            mf = int(c["metadata_failures"].get() or 0)
        except (ValueError, tk.TclError):
            u = a = mf = 0
        um = int(getattr(self, "_last_unmatched_mp4", 0) or 0)
        uc = int(getattr(self, "_last_unmatched_csv", 0) or 0)
        warn = self._last_inspired_warnings + u + a + mf + um + uc
        miss = g("cfx_missing_source")
        fail_media = g("video_fail") + g("com_fail") + g("cfx_fail")
        meta_rows = self._last_inspired_total_rows
        return (
            "Last run summary:\n"
            f"  • Metadata rows exported (normalized): {meta_rows}\n"
            f"  • Video files packaged (successful): {g('video_ok')}\n"
            f"  • COM WAV files generated (successful): {g('com_ok')}\n"
            f"  • CFX WAV files generated (successful): {g('cfx_ok')}\n"
            f"  • Warning / attention count (inspired warnings + unresolved + ambiguous + metadata failures + "
            f"CSV↔MP4 unmatched): {warn}\n"
            f"  • MP4 files with no matching CSV row (when a sidecar CSV is present): {um}\n"
            f"  • CSV data rows with no matching MP4 (when a sidecar CSV is present): {uc}\n"
            f"  • Missing CFX source count: {miss}\n"
            f"  • Failed media operations (video/COM/CFX): {fail_media}"
        )

    def _refresh_results_summary(self) -> None:
        if self._last_completion == "READY":
            self.results_attention_var.set(
                "Add MP4 files in Input files, complete Configuration (and Advanced settings if needed), then start from the sidebar."
            )
            if self._last_delivery_dir:
                self.results_delivery_var.set(f"Last delivery folder:\n{self._last_delivery_dir}")
            else:
                self.results_delivery_var.set("Last delivery folder: — (after a successful run you can open it here)")
            self.results_files_var.set("Files generated: —")
            self.results_metrics_var.set("Last run summary: —")
            return

        if self._last_completion == "RUNNING":
            self.results_attention_var.set("Run in progress — follow Progress and Media generation in the main panel.")
            self.results_delivery_var.set("When the run completes, the delivery folder path will appear here.")
            self.results_files_var.set("Files generated: —")
            self.results_metrics_var.set("Last run summary: (updates when the run finishes)")
            return

        c = self._cnt_labels
        parts: list[str] = []
        try:
            u = int(c["unresolved"].get() or 0)
            a = int(c["ambiguous"].get() or 0)
            mf = int(c["metadata_failures"].get() or 0)
        except (ValueError, tk.TclError):
            u = a = mf = 0
        if u:
            parts.append(f"{u} unresolved asset(s)")
        if a:
            parts.append(f"{a} ambiguous match(es)")
        if mf:
            parts.append(f"{mf} spreadsheet field issue(s)")
        um = int(getattr(self, "_last_unmatched_mp4", 0) or 0)
        uc = int(getattr(self, "_last_unmatched_csv", 0) or 0)
        if um:
            parts.append(f"{um} MP4(s) with no CSV row match")
        if uc:
            parts.append(f"{uc} CSV row(s) with no MP4 match")
        if self._last_inspired_warnings > 0:
            parts.append(f"{self._last_inspired_warnings} inspired row(s) with warnings")
        m = self._last_media_counts
        try:
            if int(m.get("video_fail", 0) or 0) > 0:
                parts.append("video encoding issue(s)")
            if int(m.get("com_fail", 0) or 0) > 0:
                parts.append("COM WAV issue(s)")
            if int(m.get("cfx_fail", 0) or 0) > 0:
                parts.append("CFX WAV issue(s)")
            if int(m.get("cfx_missing_source", 0) or 0) > 0:
                parts.append("missing CFX source(s)")
        except (TypeError, ValueError):
            pass
        if parts:
            self.results_attention_var.set("Items needing attention: " + "; ".join(parts))
        else:
            self.results_attention_var.set("Items needing attention: none detected (expand the technical log only if you need detail).")

        if self._last_delivery_dir:
            self.results_delivery_var.set(f"Last delivery folder:\n{self._last_delivery_dir}")
        else:
            self.results_delivery_var.set("Last delivery folder: — (run a delivery first)")

        n = len(self._mp4_list)
        if self._last_completion in ("COMPLETED", "COMPLETED WITH WARNINGS") and self._last_delivery_dir:
            self.results_files_var.set(
                f"Files generated: ENGP inputs processed ({n}); packaged media, JSON, logs, and spreadsheets "
                f"inside the delivery folder above."
            )
        else:
            self.results_files_var.set("Files generated: — (complete a run to see a summary here)")

        self.results_metrics_var.set(self._format_last_run_metrics())

    def _update_open_folder_buttons(self) -> None:
        ok = bool(self._last_delivery_dir and Path(self._last_delivery_dir).is_dir())
        self.open_delivery_btn.configure(state="normal" if ok else "disabled")
        self.open_reports_btn.configure(state="normal" if ok else "disabled")

    def _set_busy(self, busy: bool) -> None:
        self.generate_btn.configure(state="disabled" if busy else "normal")
        self.cancel_btn.configure(state="normal" if busy else "disabled")
        if busy:
            self.open_delivery_btn.configure(state="disabled")
            self.open_reports_btn.configure(state="disabled")
        else:
            self._update_open_folder_buttons()

    def _on_drop(self, event) -> None:
        raw = getattr(event, "data", "") or ""
        paths = self._parse_drop_paths(raw)
        self._add_mp4_paths(paths)

    @staticmethod
    def _parse_drop_paths(data: str) -> list[Path]:
        parts: list[str] = []
        cur = ""
        i = 0
        while i < len(data):
            ch = data[i]
            if ch == "{":
                end = data.find("}", i + 1)
                if end == -1:
                    cur += ch
                    i += 1
                    continue
                inner = data[i + 1 : end]
                parts.append(inner)
                cur = ""
                i = end + 1
                while i < len(data) and data[i] == " ":
                    i += 1
                continue
            if ch == " ":
                if cur:
                    parts.append(cur)
                    cur = ""
                i += 1
                continue
            cur += ch
            i += 1
        if cur:
            parts.append(cur)
        out: list[Path] = []
        for p in parts:
            p = p.strip()
            if not p:
                continue
            out.append(Path(p))
        return out

    def _add_mp4_paths(self, paths: list[Path]) -> None:
        added = 0
        for p in paths:
            if p.is_file() and _is_mp4(p):
                if p not in self._mp4_list:
                    self._mp4_list.append(p)
                    self.file_listbox.insert(tk.END, str(p))
                    added += 1
        if added == 0 and paths:
            messagebox.showinfo(APP_NAME, "No MP4 files were added (only .mp4 files are accepted).")
        self._refresh_operator_checklist()

    def _browse_mp4s(self) -> None:
        files = filedialog.askopenfilenames(
            parent=self.root,
            title="Choose MP4 files",
            filetypes=[("MP4 video", "*.mp4"), ("All files", "*.*")],
        )
        if files:
            self._add_mp4_paths([Path(f) for f in files])

    def _refresh_metadata_csv_summary(self) -> None:
        if not self._metadata_csv_list:
            self.metadata_csv_summary_var.set("No CSV files selected (sidecar auto-discovery still applies).")
            return
        parts = [p.name for p in self._metadata_csv_list[:8]]
        tail = f" (+{len(self._metadata_csv_list) - 8} more)" if len(self._metadata_csv_list) > 8 else ""
        self.metadata_csv_summary_var.set(f"{len(self._metadata_csv_list)} file(s): {', '.join(parts)}{tail}")

    def _browse_metadata_csvs(self) -> None:
        files = filedialog.askopenfilenames(
            parent=self.root,
            title="Choose metadata CSV file(s)",
            filetypes=[("CSV", "*.csv"), ("All files", "*.*")],
        )
        if not files:
            return
        seen: set[str] = set()
        for s in self._metadata_csv_list:
            try:
                seen.add(str(s.resolve()))
            except OSError:
                seen.add(str(s))
        added = 0
        for f in files:
            p = Path(f)
            try:
                key = str(p.resolve()) if p.is_file() else str(p)
            except OSError:
                key = str(p)
            if not p.is_file() or key in seen:
                continue
            seen.add(key)
            self._metadata_csv_list.append(p)
            added += 1
        if added == 0:
            messagebox.showinfo(APP_NAME, "No new CSV files were added (duplicates and non-files are skipped).")
        self.settings.metadata_csv_paths = [str(p) for p in self._metadata_csv_list]
        save_settings(self.settings_path, self.settings)
        self._refresh_metadata_csv_summary()

    def _clear_metadata_csvs(self) -> None:
        self._metadata_csv_list.clear()
        self.settings.metadata_csv_paths = []
        save_settings(self.settings_path, self.settings)
        self._refresh_metadata_csv_summary()

    def _browse_output(self) -> None:
        initial = self.output_var.get().strip() or str(Path.home())
        d = filedialog.askdirectory(parent=self.root, title="Choose output folder", initialdir=initial)
        if d:
            self.output_var.set(d)
            self.settings.output_folder = d
            self._persist_iconik_to_settings()
            save_settings(self.settings_path, self.settings)
            self._refresh_operator_checklist()

    def _browse_cfx_source_root(self) -> None:
        initial = self.cfx_source_root_var.get().strip() or str(Path.home())
        d = filedialog.askdirectory(parent=self.root, title="Choose CFX root folder", initialdir=initial)
        if d:
            cleaned = clean_cfx_source_root_path(d)
            self.cfx_source_root_var.set(cleaned)
            self.settings.cfx_source_root_folder = cleaned
            self._persist_iconik_to_settings()
            save_settings(self.settings_path, self.settings)
            self._refresh_operator_checklist()

    def _remove_selected(self) -> None:
        sel = list(self.file_listbox.curselection())
        if not sel:
            return
        for idx in reversed(sel):
            self.file_listbox.delete(idx)
            del self._mp4_list[idx]
        self._refresh_operator_checklist()

    def _clear_list(self) -> None:
        self.file_listbox.delete(0, tk.END)
        self._mp4_list.clear()
        self._refresh_operator_checklist()

    def _clear_preview(self) -> None:
        for iid in self.preview_tree.get_children():
            self.preview_tree.delete(iid)

    def _clear_inspired_preview(self) -> None:
        for iid in self.inspired_tree.get_children():
            self.inspired_tree.delete(iid)
        self.inspired_total_var.set("Total normalized rows: 0")
        self.inspired_warn_var.set("Rows with warnings: 0")

    def _wire_checklist_traces(self) -> None:
        def _sched(*_a: object) -> None:
            self.root.after_idle(self._refresh_operator_checklist)

        for v in (
            self.iconik_base_var,
            self.iconik_app_var,
            self.iconik_token_var,
            self.iconik_collection_var,
            self.output_var,
            self.cfx_source_root_var,
        ):
            v.trace_add("write", _sched)

    def _chk_set(self, key: str, ok: bool | None) -> None:
        var = self._checklist_status_vars.get(key)
        if not var:
            return
        if ok is True:
            var.set("OK")
        elif ok is False:
            var.set("Missing / invalid")
        else:
            var.set("—")

    def _gui_output_dir_valid(self) -> bool:
        x = self.output_var.get().strip()
        if not x:
            return False
        p = os.path.expanduser(x)
        try:
            return bool(os.path.isdir(p) and Path(p).is_dir())
        except OSError:
            return False

    def _refresh_operator_checklist(self) -> None:
        if not getattr(self, "_checklist_status_vars", None):
            return
        self._chk_set("req_mp4", bool(self._mp4_list))
        self._chk_set("req_base", bool(self.iconik_base_var.get().strip()))
        self._chk_set("req_app", bool(self.iconik_app_var.get().strip()))
        self._chk_set("req_token", bool(self.iconik_token_var.get().strip()))
        self._chk_set("req_coll", bool(self.iconik_collection_var.get().strip()))
        self._chk_set("req_out", self._gui_output_dir_valid())
        cfx_clean = clean_cfx_source_root_path(self.cfx_source_root_var.get())
        cfx_ok = bool(cfx_clean) and cfx_source_root_usable_for_scan(cfx_clean)
        self._chk_set("req_cfx", cfx_ok if cfx_clean else False)

        if not cfx_ok:
            self._chk_set("rec_cfx", None)
        else:
            n = count_engp_cfx_mp4s_under_root(cfx_clean)
            self._chk_set("rec_cfx", n > 0)
        self._chk_set("rec_ffmpeg", _ffmpeg_available())
        self._chk_set("rec_ffprobe", _ffprobe_available())

    def _validate_required_for_run(self) -> tuple[bool, str]:
        missing: list[str] = []
        if not self._mp4_list:
            missing.append("• Add at least one ENGP MP4 file in Input files.")
        if not bool(self.csv_only_delivery_var.get()):
            if not self.iconik_base_var.get().strip():
                missing.append("• Iconik Base URL is empty (open Advanced / Iconik settings).")
            if not self.iconik_app_var.get().strip():
                missing.append("• App ID is empty (open Advanced / Iconik settings).")
            if not self.iconik_token_var.get().strip():
                missing.append("• Auth Token is empty (open Advanced / Iconik settings).")
            if not self.iconik_collection_var.get().strip():
                missing.append("• Collection UUID is empty (open Advanced / Iconik settings).")
        if not self.output_var.get().strip():
            missing.append("• Output folder is not set.")
        elif not self._gui_output_dir_valid():
            missing.append("• Output folder is missing or not a valid directory.")
        cfx_clean = clean_cfx_source_root_path(self.cfx_source_root_var.get())
        if not cfx_clean:
            missing.append("• CFX source root is required — choose a folder under Configuration.")
        elif not cfx_source_root_usable_for_scan(cfx_clean):
            missing.append("• CFX source root is not a valid directory (see the checklist in the sidebar).")
        if missing:
            return False, "Cannot start — please fix the following:\n\n" + "\n".join(missing)
        return True, ""

    def _collect_recommended_warnings(self) -> list[str]:
        warns: list[str] = []
        cfx_clean = clean_cfx_source_root_path(self.cfx_source_root_var.get())
        if cfx_clean and cfx_source_root_usable_for_scan(cfx_clean):
            if count_engp_cfx_mp4s_under_root(cfx_clean) == 0:
                warns.append(
                    "No files matching *_ENGP_CFX.mp4 were found under the CFX root (media may still use siblings)."
                )
        if not _ffmpeg_available():
            warns.append("ffmpeg was not found on PATH — video and WAV packaging will be skipped.")
        if not _ffprobe_available():
            warns.append("ffprobe was not found on PATH — Duration (Frames) may stay incomplete.")
        return warns

    def _show_start_delivery_confirmation(self, n_files: int, out: str, cfx_display: str) -> bool:
        confirmed: list[bool] = [False]
        top = tk.Toplevel(self.root)
        top.title("Confirm delivery")
        top.transient(self.root)
        top.grab_set()
        top.resizable(True, True)
        fr = ttk.Frame(top, padding=14)
        fr.pack(fill=tk.BOTH, expand=True)
        body = (
            f"Requested files: {n_files}\n\n"
            f"Output folder:\n{out}\n\n"
            f"CFX root:\n{cfx_display}\n\n"
            "Expected outputs (in the new delivery folder):\n"
            "  • inspired_metadata.xlsx\n"
            "  • Video-only ENGP MP4 (per input)\n"
            "  • *_ENGP_COM.wav\n"
            "  • *_ENGP_CFX.wav\n"
            "  • plus logs, JSON previews, and other reports\n"
        )
        ttk.Label(fr, text=body, justify=tk.LEFT).pack(anchor=tk.W, pady=(0, 12))
        btn_row = ttk.Frame(fr)
        btn_row.pack(fill=tk.X)

        def _start() -> None:
            confirmed[0] = True
            top.destroy()

        def _cancel() -> None:
            top.destroy()

        ttk.Button(btn_row, text="Start Delivery", command=_start).pack(side=tk.LEFT, padx=(0, 8))
        ttk.Button(btn_row, text="Cancel", command=_cancel).pack(side=tk.LEFT)
        top.protocol("WM_DELETE_WINDOW", _cancel)
        self.root.wait_window(top)
        return confirmed[0]

    def _generate(self) -> None:
        if self.state.worker_thread and self.state.worker_thread.is_alive():
            return
        self._refresh_operator_checklist()
        ok, err = self._validate_required_for_run()
        if not ok:
            messagebox.showerror("Cannot start delivery", err)
            return
        rec = self._collect_recommended_warnings()
        if rec:
            detail = "\n\n".join(rec)
            if not messagebox.askyesno(
                "Recommended checks",
                detail + "\n\nThese issues may reduce output quality. Continue anyway?",
            ):
                return
        out = self.output_var.get().strip()
        raw_cfx_gui = self.cfx_source_root_var.get()
        cleaned_cfx = clean_cfx_source_root_path(raw_cfx_gui)
        cfx_show = cleaned_cfx if cleaned_cfx else "(not set)"
        if not self._show_start_delivery_confirmation(len(self._mp4_list), out, cfx_show):
            return
        self._execute_generate_delivery(out, raw_cfx_gui, cleaned_cfx)

    def _execute_generate_delivery(self, out: str, raw_cfx_gui: str, cleaned_cfx: str) -> None:
        out_path = Path(out)
        self.cfx_source_root_var.set(cleaned_cfx)

        self._persist_iconik_to_settings()
        self.settings.output_folder = out
        self.settings.metadata_csv_paths = [str(p) for p in self._metadata_csv_list]
        self.settings.lock_delivery_outputs = bool(self.lock_delivery_outputs_var.get())
        self.settings.csv_only_delivery = bool(self.csv_only_delivery_var.get())
        save_settings(self.settings_path, self.settings)

        self.state.cancel_event = threading.Event()
        self.state.msg_queue = queue.Queue()
        self.state.worker_thread = DeliveryWorker(
            mp4_paths=list(self._mp4_list),
            output_parent=out_path,
            message_queue=self.state.msg_queue,
            cancel_event=self.state.cancel_event,
            iconik_base_url=self.settings.iconik_base_url,
            iconik_app_id=self.settings.iconik_app_id,
            iconik_auth_token=self.settings.iconik_auth_token,
            iconik_collection_uuid=self.settings.iconik_collection_uuid,
            cfx_source_root_folder=cleaned_cfx,
            cfx_source_root_folder_raw=raw_cfx_gui,
            metadata_csv_paths=list(self._metadata_csv_list),
            lock_delivery_outputs=bool(self.lock_delivery_outputs_var.get()),
            csv_only_delivery=bool(self.csv_only_delivery_var.get()),
        )

        self.progress.set(0)
        self.diagnostics_text.configure(state="normal")
        self.diagnostics_text.delete("0.0", "end")
        self.diagnostics_text.configure(state="disabled")
        self._last_media_counts = {}
        self._last_unmatched_mp4 = 0
        self._last_unmatched_csv = 0
        self._last_inspired_warnings = 0
        self._last_inspired_total_rows = 0
        self._reset_counters_ui()
        self._reset_enrichment_ui()
        self._clear_preview()
        self._clear_inspired_preview()
        self._set_workflow_state("RUNNING")
        self._set_busy(True)
        self._last_completion = "RUNNING"
        self._refresh_results_summary()
        self.state.worker_thread.start()

    def _cancel(self) -> None:
        self.state.cancel_event.set()
        self._append_status("Cancel requested…")

    def _poll_queue(self) -> None:
        try:
            while True:
                msg = self.state.msg_queue.get_nowait()
                mtype = msg.get("type")
                if mtype == "progress":
                    self._ctk_prog_set(self.progress, float(msg.get("value", 0)) / 100.0)
                elif mtype == "status_line":
                    self._append_status(str(msg.get("text", "")))
                elif mtype == "state":
                    self._set_workflow_state(str(msg.get("state", "RUNNING")))
                elif mtype == "counters":
                    for k in ("requested", "resolved", "unresolved", "ambiguous", "metadata_failures"):
                        if k in self._cnt_labels:
                            self._cnt_labels[k].set(str(msg.get(k, 0)))
                    self._last_unmatched_mp4 = int(msg.get("unmatched_mp4", 0) or 0)
                    self._last_unmatched_csv = int(msg.get("unmatched_csv", 0) or 0)
                elif mtype == "preview_clear":
                    self._clear_preview()
                    self._clear_inspired_preview()
                    self._reset_enrichment_ui()
                elif mtype == "preview_row":
                    self.preview_tree.insert(
                        "",
                        tk.END,
                        values=(
                            msg.get("filename", ""),
                            msg.get("asset_id", ""),
                            msg.get("title", ""),
                            msg.get("metadata_ok", ""),
                            msg.get("metadata_preview", ""),
                        ),
                    )
                elif mtype == "inspired_preview_row":
                    self.inspired_tree.insert(
                        "",
                        tk.END,
                        values=(
                            msg.get("filename", ""),
                            msg.get("resolved_title", ""),
                            msg.get("country", ""),
                            msg.get("match_date", ""),
                            msg.get("home_team", ""),
                            msg.get("away_team", ""),
                            msg.get("goal_scored", ""),
                            msg.get("goal_time_frames", ""),
                            msg.get("normalization", ""),
                        ),
                    )
                elif mtype == "inspired_stats":
                    self._last_inspired_warnings = int(msg.get("with_warnings", 0) or 0)
                    self._last_inspired_total_rows = int(msg.get("total", 0) or 0)
                    self.inspired_total_var.set(f"Total normalized rows: {msg.get('total', 0)}")
                    self.inspired_warn_var.set(f"Rows with warnings: {msg.get('with_warnings', 0)}")
                elif mtype == "ffprobe_progress":
                    cur = int(msg.get("current", 0))
                    tot = int(msg.get("total", 1))
                    pop = int(msg.get("populated", 0))
                    self.ffprobe_progress_var.set(
                        f"Frame timing: {cur} / {tot} files (duration fields filled: {pop})"
                    )
                elif mtype == "xlsx_progress":
                    v = float(msg.get("value", 0))
                    self.xlsx_bar.set(max(0.0, min(1.0, v / 100.0)))
                    self.xlsx_progress_var.set(f"Spreadsheet export: {int(v)}%")
                elif mtype == "duration_frames_count":
                    self.duration_frames_pop_var.set(f"Duration (Frames) filled: {int(msg.get('count', 0))}")
                elif mtype == "media_phase_begin":
                    tot = max(int(msg.get("total", 1)), 1)
                    for pb in (self.media_video_pb, self.media_com_pb, self.media_cfx_pb):
                        pb.set(0)
                    self.media_video_var.set(f"Video packaging: 0 / {tot}")
                    self.media_com_var.set(f"COM audio (WAV): 0 / {tot}")
                    self.media_cfx_var.set(f"CFX audio (WAV): 0 / {tot}")
                elif mtype == "media_video_progress":
                    cur = int(msg.get("current", 0))
                    tot = max(int(msg.get("total", 1)), 1)
                    fn = str(msg.get("filename", ""))
                    self.media_video_pb.set(cur / tot)
                    self.media_video_var.set(f"Video packaging: {cur} / {tot} — {fn}")
                elif mtype == "media_com_progress":
                    cur = int(msg.get("current", 0))
                    tot = max(int(msg.get("total", 1)), 1)
                    fn = str(msg.get("filename", ""))
                    self.media_com_pb.set(cur / tot)
                    self.media_com_var.set(f"COM audio (WAV): {cur} / {tot} — {fn}")
                elif mtype == "media_cfx_progress":
                    cur = int(msg.get("current", 0))
                    tot = max(int(msg.get("total", 1)), 1)
                    fn = str(msg.get("filename", ""))
                    self.media_cfx_pb.set(cur / tot)
                    self.media_cfx_var.set(f"CFX audio (WAV): {cur} / {tot} — {fn}")
                elif mtype == "media_counts":
                    keys = ("video_ok", "video_fail", "com_ok", "com_fail", "cfx_ok", "cfx_fail", "cfx_missing_source")
                    self._last_media_counts = {k: int(msg.get(k, 0) or 0) for k in keys}
                    self.media_counts_var.set(
                        "Media summary: "
                        f"video OK / issue {msg.get('video_ok', 0)}/{msg.get('video_fail', 0)} · "
                        f"COM OK / issue {msg.get('com_ok', 0)}/{msg.get('com_fail', 0)} · "
                        f"CFX OK / issue {msg.get('cfx_ok', 0)}/{msg.get('cfx_fail', 0)} · "
                        f"CFX source missing {msg.get('cfx_missing_source', 0)}"
                    )
                elif mtype == "finished":
                    completion = str(msg.get("completion", "FAILED"))
                    dd = msg.get("delivery_dir")
                    if isinstance(dd, str) and dd.strip():
                        self._last_delivery_dir = dd.strip()

                    display = completion
                    if completion == "COMPLETED" and self._run_has_warnings():
                        display = "COMPLETED WITH WARNINGS"
                    self._last_completion = display

                    self._set_workflow_state(display)
                    if completion == "COMPLETED":
                        self.progress.set(1.0)
                    self._set_busy(False)
                    self.state.worker_thread = None
                    self._refresh_results_summary()
                    self._update_open_folder_buttons()

                    if completion == "COMPLETED" or display == "COMPLETED WITH WARNINGS":
                        if self._last_delivery_dir:
                            self._append_status(f"Done. Delivery directory: {self._last_delivery_dir}")
                    elif completion == "CANCELLED":
                        self._append_status("Run cancelled.")
                    else:
                        self._append_status("Run finished with errors.")
        except queue.Empty:
            pass
        self.root.after(80, self._poll_queue)

    def _on_close(self) -> None:
        if self.state.worker_thread and self.state.worker_thread.is_alive():
            self.state.cancel_event.set()
            self.state.worker_thread.join(timeout=2.0)
        self.settings.output_folder = self.output_var.get().strip()
        self.settings.metadata_csv_paths = [str(p) for p in self._metadata_csv_list]
        self.settings.lock_delivery_outputs = bool(self.lock_delivery_outputs_var.get())
        self.settings.csv_only_delivery = bool(self.csv_only_delivery_var.get())
        self._persist_iconik_to_settings()
        self.settings.window_geometry = self.root.geometry()
        save_settings(self.settings_path, self.settings)
        self.root.destroy()


def main() -> None:
    ctk.set_appearance_mode("dark")
    root = ctk.CTk()
    root.title(APP_NAME)
    InspiredDeliveryGeneratorApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
